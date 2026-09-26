"""Evidence-first Gecko V2 resume pipeline. Job Scout is intentionally separate."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import subprocess
import sys
from collections import Counter
from pathlib import Path

import fitz
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from silent_subprocess import windows_creationflags

ROOT = Path(__file__).resolve().parents[1]
MASTER = ROOT / "input/master-resume/Dave-Call-resume-9-23-26.docx"
SECTIONS = ("Professional Summary", "Core Competencies & Technical Skills", "Professional Experience", "Education & Certifications")
TERMS = (
    "SEO", "technical SEO", "paid search", "Google Ads", "Bing Ads", "Microsoft Ads", "Meta Ads", "Google Analytics",
    "GA4", "Google Tag Manager", "GTM", "Looker Studio", "Tableau", "SQL", "Python", "JavaScript",
    "HTML5", "CSS3", "Magento", "Shopify", "WordPress", "Amazon", "HubSpot", "CRM", "email",
    "automation", "attribution", "CRO", "conversion rate optimization", "keyword research", "content",
    "schema markup", "metadata", "taxonomy", "link building", "Semrush", "Search Console", "e-commerce",
    "B2B", "B2C", "budget", "forecasting", "reporting", "analytics", "leadership", "management",
    "social media", "customer acquisition", "lifecycle", "AI", "machine learning", "Adobe Analytics",
    "Excel", "Power BI", "Salesforce", "Marketo", "Klaviyo", "Mailchimp", "Snowflake",
    "Core Web Vitals", "canonicalization", "XML sitemaps", "robots directives", "faceted navigation",
    "AI search", "AI Overviews", "backlinks", "schema", "structured data", "catalog", "merchandising",
    "Codex", "ChatGPT", "Claude", "Perplexity", "Cursor", "AntiGravity",
)
TOOLS = set("Google Ads|Bing Ads|Microsoft Ads|Meta Ads|Google Analytics|GA4|Google Tag Manager|GTM|Looker Studio|Tableau|SQL|Python|JavaScript|HTML5|CSS3|Magento|Shopify|WordPress|Amazon|HubSpot|Semrush|Search Console|Adobe Analytics|Excel|Power BI|Salesforce|Marketo|Klaviyo|Mailchimp|Snowflake|Codex|ChatGPT|Claude|Perplexity|Cursor|AntiGravity".split("|"))
STOP = set("the and for with from into your their our you will are this that have has must should ability using use work role team experience years strong plus including across based more about through".split())


def normalized(text: str) -> str:
    return re.sub(r"\s+", " ", text.replace("\ufffd", " ")).strip()


def key(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", text.casefold()).strip()


def contains_term(haystack: str, term: str) -> bool:
    return bool(re.search(r"(?:^| )" + re.escape(key(term)) + r"(?: |$)", key(haystack)))


def words(text: str) -> set[str]:
    return {w for w in key(text).split() if len(w) > 2 and w not in STOP}


def source_text() -> str:
    """Read the current DOCX on every call, including table-based skills."""
    doc = Document(MASTER)
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    lines.extend(cell.text for table in doc.tables for row in table.rows for cell in row.cells if cell.text.strip())
    return "\n".join(lines)


def source_hashes() -> dict[str, str]:
    return {str(MASTER.relative_to(ROOT)): hashlib.sha256(MASTER.read_bytes()).hexdigest()}


def extract_evidence(source: str) -> list[dict]:
    """Take results and work-history bullets from the canonical DOCX only."""
    lines = [p.text.strip() for p in Document(MASTER).paragraphs]
    starts = [next(i for i, line in enumerate(lines) if line == f"{job['company']} - {job['location']}")
              for job in master_jobs()]
    end = lines.index("PAID SEARCH LEADERSHIP & CLIENT PARTNERSHIP")
    evidence = []
    for job_index, start in enumerate(starts):
        stop = starts[job_index + 1] if job_index + 1 < len(starts) else end
        for line in lines[start + 1:stop]:
            if not re.match(r"^[\u2022\ufffd]\s+", line):
                continue
            quote = normalized(re.sub(r"^[\u2022\ufffd]\s+", "", line))
            if len(quote) >= 35:
                evidence.append({"id": f"E{len(evidence)+1:03d}", "job": job_index,
                                 "source": str(MASTER.relative_to(ROOT)), "quote": quote})
    results = lines[lines.index("SELECTED RESULTS") + 1:lines.index("PROFESSIONAL EXPERIENCE")]
    for job_index, line in zip((0, 3, 1, 4), results):
        quote = normalized(re.sub(r"^[\u2022\ufffd]\s+", "", line))
        evidence.append({"id": f"E{len(evidence)+1:03d}", "job": job_index,
                         "source": str(MASTER.relative_to(ROOT)), "quote": quote})
    return evidence


def listing_metadata(listing: str, path: Path) -> dict[str, str]:
    def field(name: str) -> str:
        match = re.search(rf"(?mi)^\s*-\s*\*\*{re.escape(name)}[^*]*\*\*\s*:?\s*(.*?)\s*$", listing)
        return match.group(1).strip(" `") if match else ""
    headline = next((line.lstrip("# ").strip() for line in listing.splitlines() if line.startswith("# ")), "")
    company = field("Company") or (headline.split(" — ")[-1] if " — " in headline else "")
    title = headline.split(" — ")[0] if headline else ""
    job_number = field("Job Key (Indeed jk)") or field("Job Number") or path.stem.split("+")[-1]
    if not company or not title or not job_number:
        raise ValueError("Listing needs a title, company, and job number; add metadata to the archived description.")
    safe_company = re.sub(r'[\\/:*?"<>|]', "", company).replace(" ", "-").rstrip(" .")
    return {"company": company, "safe_company": safe_company, "title": title, "job_number": job_number}


def resume_filename(company: str, title: str, job_number: str) -> str:
    """Build a Word-safe name while retaining the job number as the final field."""
    def safe_part(value: str) -> str:
        value = re.sub(r'[\\/:*?"<>|+]+', "-", normalized(value))
        value = re.sub(r"\s+", "-", value)
        return re.sub(r"-{2,}", "-", value).strip(" .-")

    company_part, title_part, number_part = map(safe_part, (company, title, job_number))
    if not all((company_part, title_part, number_part)):
        raise ValueError("Resume filename needs a company, job title, and job number")
    # Keep the complete path comfortable for Word on Windows; only long titles are shortened.
    available = 180 - len(f"Dave-Call+{company_part}++{number_part}.docx")
    if available < 20:
        raise ValueError("Company and job number leave too little room for the job title")
    return f"Dave-Call+{company_part}+{title_part[:available].rstrip(' .-')}+{number_part}.docx"


REQUIREMENT_STARTERS = (
    "ability to", "accredited", "analyze", "bachelor", "build", "collect", "collaborate",
    "comfortable", "communicate", "conduct", "coordinate", "create", "define", "develop",
    "drive", "ensure", "evangelize", "execute", "experience", "identify", "implement",
    "lead", "leverage", "maintain", "manage", "minimum", "monitor", "optimize", "oversee",
    "own", "partner", "present", "prior experience", "product knowledge", "proven", "provide",
    "represent", "report", "strong", "support", "understand", "when required", "work with",
)


def _requirement_candidates(listing: str) -> list[tuple[str, str]]:
    """Extract bullets and aggregator-flattened requirement clauses with section context."""
    text = listing.replace("\\n", "\n")
    # Aggregators frequently flatten section headings and every bullet into one paragraph.
    inline_headings = (
        "Professional Experience/Background to be successful in this role",
        "Competencies (Attributes needed to be successful in this role)",
        "Expected Outcomes in 3, 6, or 12 months",
        "Responsibilities", "Qualifications", "Required Qualifications",
        "Preferred Qualifications", "Requirements", "What You'll Do", "What You Will Do",
    )
    for label in inline_headings:
        text = re.sub(rf"\s*{re.escape(label)}\s*:\s*", f"\n## {label}\n", text, flags=re.I)
    starter_pattern = "|".join(re.escape(value) for value in sorted(REQUIREMENT_STARTERS, key=len, reverse=True))
    heading = ""
    candidates: list[tuple[str, str]] = []
    for raw in text.splitlines():
        stripped = raw.strip()
        if stripped.startswith("#"):
            heading = key(stripped.lstrip("# "))
            continue
        if not stripped or stripped.startswith("- **") or stripped.startswith("http"):
            continue
        is_bullet = stripped.startswith(("-", "*", "•", "ï‚·"))
        clean = re.sub(r"^(?:[-*•]|ï‚·)\s+", "", stripped).strip()
        pieces = [clean] if is_bullet else re.split(
            rf"(?<=[.!?;])\s+|\s+(?=(?:{starter_pattern})\b)", clean, flags=re.I
        )
        for piece in pieces:
            candidate = normalized(piece).strip(" -•")
            if len(candidate) < 18 or len(candidate) > 500:
                continue
            contextual = bool(re.search(
                r"responsibil|qualification|requirement|experience|background|competenc|outcome|what you",
                heading, re.I,
            ))
            starts_like_requirement = bool(re.match(rf"(?:{starter_pattern})\b", candidate, re.I))
            if is_bullet or contextual or starts_like_requirement:
                candidates.append((heading, candidate))
    return candidates


def requirements(listing: str, source: str, evidence: list[dict]) -> dict[str, list[dict]]:
    result = {name: [] for name in ("required_skills", "preferred_skills", "responsibilities", "tools", "seniority_signals", "industry_terminology", "ats_keywords")}
    for heading, candidate in _requirement_candidates(listing):
        preferred = bool(re.search(r"preferred|nice to have|bonus|strong plus", heading + " " + candidate, re.I))
        required = bool(re.search(r"required|qualification|must have", heading + " " + candidate, re.I))
        category = "preferred_skills" if preferred else "required_skills" if required else "responsibilities"
        overlaps = [(len(words(candidate) & words(item["quote"])), item["id"]) for item in evidence]
        best = max(overlaps, default=(0, ""))
        term_hits = [term for term in TERMS if contains_term(candidate, term)]
        supported_terms = [term for term in term_hits if contains_term(source, term)]
        unsupported_terms = [term for term in term_hits if term not in supported_terms]
        status = "gap" if unsupported_terms else "supported" if supported_terms and best[0] >= 2 else "review" if best[0] >= 3 else "gap"
        result[category].append({"text": candidate, "status": status,
                                 "evidence_ids": [best[1]] if status != "gap" else [],
                                 "matched_terms": supported_terms, "unsupported_terms": unsupported_terms})
        for term in term_hits:
            group = "tools" if term in TOOLS else "ats_keywords"
            if not any(x["text"].casefold() == term.casefold() for x in result[group]):
                result[group].append({"text": term, "status": "supported" if term in supported_terms else "gap",
                                      "evidence_ids": [best[1]] if term in supported_terms and best[1] else []})
        if re.search(r"lead|manag|director|own|mentor|budget|executive|senior", candidate, re.I):
            result["seniority_signals"].append({"text": candidate, "status": status, "evidence_ids": [best[1]] if best[1] else []})
        if re.search(r"healthcare|medical|e-commerce|retail|b2b|b2c|saas|agency|catalog", candidate, re.I):
            result["industry_terminology"].append({"text": candidate, "status": status, "evidence_ids": [best[1]] if best[1] else []})
    return result


def relevant_metric_ids(listing: str, evidence: list[dict]) -> list[str]:
    """Protect proven scale/outcome bullets when the target work makes them useful."""
    metric_rules = (
        (0, r"\$30 million per month", r"paid|budget|marketing manager|leadership|director"),
        (1, r"\$200K to \$800K|fourfold", r"e-commerce|ecommerce|growth|revenue|paid"),
        (2, r"650 B2B retail partners", r"b2b|retail|e-commerce|ecommerce|catalog"),
        (3, r"75\+ Google Ads accounts", r"paid|search|agency|account|sem"),
        (4, r"\$9 million in revenue", r"revenue|e-commerce|ecommerce|marketing|growth"),
    )
    result = []
    for job, metric, trigger in metric_rules:
        if re.search(trigger, listing, re.I):
            match = next((item for item in evidence if item["job"] == job and re.search(metric, item["quote"], re.I)), None)
            if match:
                result.append(match["id"])
    return result


def polish_quote(quote: str) -> str:
    """Apply small grammar edits without changing facts or product names."""
    for old, new in (("Setup ", "Set up "), ("Created a inventory", "Created an inventory"),
                     ("I helped convert", "Helped convert")):
        quote = quote.replace(old, new)
    return quote


def master_identity() -> dict[str, str]:
    doc = Document(MASTER)
    lines = [p.text.strip() for p in doc.paragraphs]
    return {"name": lines[0], "headline": lines[1], "contact": lines[2], "summary": lines[4],
            "education": next(line for line in lines if line.startswith("Bachelor of Science")),
            "certifications": next(line for line in lines if line.startswith("Certifications:"))}


def master_jobs() -> list[dict[str, str]]:
    doc = Document(MASTER)
    lines = [p.text.strip() for p in doc.paragraphs]
    experience = lines[lines.index("PROFESSIONAL EXPERIENCE") + 1:lines.index("PAID SEARCH LEADERSHIP & CLIENT PARTNERSHIP")]
    employers = [line.rsplit(" - ", 1) for line in experience
                 if line and not re.match(r"^[\u2022\ufffd]\s+", line) and " - " in line]
    titles = [table.cell(0, 0).text.strip() for table in doc.tables[1:1 + len(employers)]]
    if len(titles) != len(employers):
        raise ValueError("Master DOCX job titles and employers do not align.")
    return [{"title": title, "company": company, "location": location}
            for title, (company, location) in zip(titles, employers)]


def create_plan(listing_path: Path) -> dict:
    listing = listing_path.read_text(encoding="utf-8-sig")
    source = source_text()
    evidence = extract_evidence(source)
    meta = listing_metadata(listing, listing_path)
    reqs = requirements(listing, source, evidence)
    target = words(listing)
    mandatory_metrics = relevant_metric_ids(listing, evidence)
    selected = []
    for index in range(len(master_jobs())):
        pool = [item for item in evidence if item["job"] == index]
        concise = [item for item in pool if len(item["quote"].split()) <= 55]
        if len(concise) >= 4:
            pool = concise
        ranked = sorted(pool, key=lambda item: (
            len(words(item["quote"]) & target) + 3 * bool(re.search(r"\$[\d,]+|\b\d+[%+]", item["quote"]))
            - 4 * bool(re.search(r"\b(?:I|my|we|our)\b", item["quote"])),
            ), reverse=True)
        chosen = [item["id"] for item in ranked[:4]]
        for eid in mandatory_metrics:
            if next((item for item in evidence if item["id"] == eid), None)["job"] == index and eid not in chosen:
                chosen[-1] = eid
        selected.extend(chosen)
    skill_terms = [term for term in TERMS if contains_term(listing, term) and contains_term(source, term)]
    return {"version": 2, "job": meta, "listing": str(listing_path.resolve()),
            "listing_sha256": hashlib.sha256(listing_path.read_bytes()).hexdigest(),
            "source_sha256": source_hashes(), "requirements": reqs, "evidence": evidence,
            "resume": {**master_identity(), "jobs": master_jobs(),
                       "skills": skill_terms[:24], "selected_evidence_ids": selected,
                       "relevant_metric_ids": mandatory_metrics},
            "review_notes": ["Review requirement classifications and evidence links before generation.",
                             "All factual content comes from the current master DOCX; other resumes are format references only."]}


def verify_plan(plan: dict) -> None:
    if plan.get("version") != 2 or plan.get("source_sha256") != source_hashes():
        raise ValueError("Plan source has changed; rebuild the tailoring plan.")
    listing_path = Path(plan["listing"])
    if hashlib.sha256(listing_path.read_bytes()).hexdigest() != plan["listing_sha256"]:
        raise ValueError("Listing has changed; rebuild the tailoring plan.")
    current_source = source_text()
    actual = {item["id"]: item for item in extract_evidence(current_source)}
    if plan["requirements"] != requirements(listing_path.read_text(encoding="utf-8-sig"), current_source, list(actual.values())):
        raise ValueError("Requirement classifications differ from the listing; rebuild the plan.")
    for item in plan["evidence"]:
        if actual.get(item["id"]) != item:
            raise ValueError(f"Unsupported or altered evidence: {item['id']}")
    for eid in plan["resume"]["selected_evidence_ids"]:
        if eid not in actual:
            raise ValueError(f"Unknown evidence ID: {eid}")
    if plan["resume"]["relevant_metric_ids"] != relevant_metric_ids(listing_path.read_text(encoding="utf-8-sig"), list(actual.values())):
        raise ValueError("Relevant accomplishment anchors differ from the source and listing.")
    source = source_text()
    for term in plan["resume"]["skills"]:
        if not contains_term(source, term):
            raise ValueError(f"Unsupported skill: {term}")
    if any(plan["resume"].get(field) != value for field, value in master_identity().items()):
        raise ValueError("Identity, summary, education, or certifications differ from the current master DOCX.")
    if plan["resume"].get("jobs") != master_jobs():
        raise ValueError("Job history differs from the current master DOCX.")


def make_resume(plan: dict, path: Path) -> None:
    verify_plan(plan)
    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(.48)
        section.left_margin = section.right_margin = Inches(.5)
        section.page_width, section.page_height = Inches(8.5), Inches(11)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(0)

    def para(text: str, *, bold=False, center=False, after=0, before=0, color=None, style=None):
        p = doc.add_paragraph(style=style)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(text)
        r.bold = bold
        r.font.name, r.font.size = "Calibri", Pt(11)
        if color:
            r.font.color.rgb = RGBColor(*color)
        return p

    def section(title):
        p = para(title.upper(), bold=True, before=5, after=2, color=(16, 44, 87))
        p.paragraph_format.keep_with_next = True
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for attr, value in (("val", "single"), ("sz", "6"), ("color", "102C57")):
            bottom.set(qn("w:" + attr), value)
        border.append(bottom)
        p._p.get_or_add_pPr().append(border)

    para(plan["resume"]["name"], bold=True, center=True, after=1, color=(16, 44, 87)).runs[0].font.size = Pt(19)
    para(plan["resume"]["headline"].upper(), bold=True, center=True, after=2, color=(40, 70, 110))
    para(plan["resume"]["contact"], center=True, after=3).runs[0].font.size = Pt(9.5)
    section(SECTIONS[0])
    para(plan["resume"]["summary"], after=2)
    section(SECTIONS[1])
    para(" • ".join(plan["resume"]["skills"]), after=2)
    section(SECTIONS[2])
    evidence = {item["id"]: item for item in plan["evidence"]}
    for index, job in enumerate(plan["resume"]["jobs"]):
        title, company, location = job["title"], job["company"], job["location"]
        selected = [evidence[eid] for eid in plan["resume"]["selected_evidence_ids"] if evidence[eid]["job"] == index]
        if not selected:
            continue
        table = doc.add_table(rows=1, cols=2)
        table.autofit = False
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.columns[0].width, table.columns[1].width = Inches(5.7), Inches(1.8)
        table.cell(0, 0).width, table.cell(0, 1).width = Inches(5.7), Inches(1.8)
        table.cell(0, 0).paragraphs[0].add_run(f"{title} | {company} — {location}").bold = True
        table.cell(0, 1).text = ""  # Required right-side date cell.
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            node = OxmlElement("w:" + edge)
            node.set(qn("w:val"), "none")
            borders.append(node)
        table._tbl.tblPr.append(borders)
        for item in selected:
            p = para(polish_quote(item["quote"]), style="List Bullet", after=1)
            p.paragraph_format.left_indent = Inches(.18)
    section(SECTIONS[3])
    para(plan["resume"]["education"], after=1)
    para(plan["resume"]["certifications"])
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def inspect_docx(plan: dict, docx_path: Path) -> list[str]:
    issues = []
    if plan.get("source_sha256") != source_hashes():
        issues.append("Master DOCX changed after planning; rebuild the plan and resume.")
    doc = Document(docx_path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    evidence = {item["id"]: item for item in plan["evidence"]}
    allowed = {polish_quote(evidence[eid]["quote"]) for eid in plan["resume"]["selected_evidence_ids"]}
    actual_bullets = [p.text.strip() for p in doc.paragraphs if p.style.name == "List Bullet"]
    if set(actual_bullets) != allowed or len(actual_bullets) != len(allowed):
        issues.append("Resume bullets differ from the source-backed tailoring plan.")
    for eid in plan["resume"]["relevant_metric_ids"]:
        if polish_quote(evidence[eid]["quote"]) not in actual_bullets:
            issues.append(f"Relevant quantified accomplishment omitted: {eid}")
    fixed = {plan["resume"]["name"], plan["resume"]["headline"].upper(), plan["resume"]["summary"],
             plan["resume"]["contact"], " • ".join(plan["resume"]["skills"]),
             plan["resume"]["education"], plan["resume"]["certifications"]}
    fixed.update(s.upper() for s in SECTIONS)
    if any(p not in fixed and p not in allowed for p in paragraphs):
        issues.append("DOCX contains text outside the approved source-backed plan.")
    if len(doc.tables) != sum(any(evidence[eid]["job"] == i for eid in plan["resume"]["selected_evidence_ids"]) for i in range(len(plan["resume"]["jobs"]))):
        issues.append("Job header/date-cell count differs from the plan.")
    expected_headers = [f"{job['title']} | {job['company']} — {job['location']}"
                        for i, job in enumerate(plan["resume"]["jobs"])
                        if any(evidence[eid]["job"] == i for eid in plan["resume"]["selected_evidence_ids"])]
    if [table.cell(0, 0).text.strip() for table in doc.tables] != expected_headers:
        issues.append("Job headers differ from the current master-backed plan.")
    if any(cell.text.strip() for table in doc.tables for cell in [table.cell(0, 1)]):
        issues.append("A date cell contains visible text.")
    if any(node.get(qn("w:val")) != "none" for table in doc.tables for node in table._tbl.tblPr.xpath(".//w:tblBorders/*")):
        issues.append("A job table has visible borders.")
    if [p for p in paragraphs if p in [s.upper() for s in SECTIONS]] != [s.upper() for s in SECTIONS]:
        issues.append("Required section order is missing or changed.")
    if not paragraphs[2].endswith("linkedin.com/in/mdavidcall") or "Spanish" in paragraphs[2]:
        issues.append("Header contact format is incorrect.")
    if abs(doc.styles["Normal"].font.size.pt - 11) > .01 or abs(doc.styles["Normal"].paragraph_format.line_spacing - 1.15) > .01:
        issues.append("Body font or line spacing differs from Gecko rules.")
    for p in doc.paragraphs[3:]:
        if p.paragraph_format.line_spacing not in (None, 1.15):
            issues.append("A body paragraph has incorrect line spacing.")
            break
        if any(run.font.size is not None and abs(run.font.size.pt - 11) > .01 for run in p.runs):
            issues.append("A body paragraph has incorrect font size.")
            break
    if len(doc.sections) != 1 or doc.sections[0].page_width != Inches(8.5) or doc.sections[0].page_height != Inches(11):
        issues.append("Page size/layout differs from Gecko rules.")
    if any(getattr(doc.sections[0], edge) < Inches(.4) for edge in ("top_margin", "bottom_margin", "left_margin", "right_margin")):
        issues.append("Page margins are smaller than Gecko's safety margin.")
    if doc.element.body.xpath(".//w:shd"):
        issues.append("Background shading violates the white-page layout.")
    approved_terms = source_text()
    for term in plan["resume"]["skills"]:
        if not contains_term(approved_terms, term):
            issues.append(f"Unsupported skill: {term}")
    for token, count in Counter(re.findall(r"\b[A-Za-z][A-Za-z0-9+]{2,}\b", " ".join(paragraphs).casefold())).items():
        if token not in STOP and count >= 18 and token in words(Path(plan["listing"]).read_text(encoding="utf-8-sig")):
            issues.append(f"Possible keyword stuffing: {token} appears {count} times.")
    return issues


def native_qa(plan: dict, docx_path: Path, scratch: Path) -> dict:
    scratch.mkdir(parents=True, exist_ok=True)
    pdf = scratch / "word-export.pdf"
    status = scratch / "validation-status.json"
    cmd = ["powershell.exe", "-NoProfile", "-NonInteractive", "-WindowStyle", "Hidden",
           "-ExecutionPolicy", "Bypass", "-File", str(ROOT / "scripts/validate_word_native.ps1"),
           "-DocxPath", str(docx_path), "-PdfPath", str(pdf), "-ResultPath", str(status)]
    result = subprocess.run(cmd, cwd=ROOT, capture_output=True, text=True,
                            creationflags=windows_creationflags())
    issues = inspect_docx(plan, docx_path)
    native = json.loads(status.read_text(encoding="utf-8-sig")) if status.exists() else {}
    if result.returncode or native.get("status") != "native-valid" or native.get("word_pages") != 2 or native.get("pdf_pages") != 2:
        issues.append("Native Word pagination did not confirm exactly two pages. " + (result.stderr.strip() or result.stdout.strip())[-500:])
    if pdf.exists():
        with fitz.open(pdf) as rendered:
            if len(rendered) == 2:
                for number, page in enumerate(rendered, 1):
                    blocks = [fitz.Rect(b[:4]) for b in page.get_text("blocks") if normalized(str(b[4]))]
                    if not blocks or max(block.y1 for block in blocks) > page.rect.height - 18:
                        issues.append(f"Page {number} has possible bottom overflow or no extractable text.")
                    if any(b.x0 < 12 or b.x1 > page.rect.width - 12 for b in blocks):
                        issues.append(f"Page {number} has possible horizontal overflow.")
    reqs = plan["requirements"]
    weaknesses = [r["text"] for group in ("required_skills", "preferred_skills", "responsibilities") for r in reqs[group] if r["status"] != "supported"]
    report = {"status": "pass" if not issues else "fail", "word_pages": native.get("word_pages"),
              "pdf_pages": native.get("pdf_pages"), "issues": issues, "remaining_weaknesses": weaknesses,
              "native_validation": str(status), "word_output": result.stdout.strip()}
    (scratch / "v2-qa.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
    return report


def write_match_report(plan: dict, qa: dict) -> Path:
    if qa["status"] != "pass":
        raise ValueError("A match report is final only after V2 QA passes.")
    reqs = plan["requirements"]
    core = reqs["required_skills"] + reqs["responsibilities"]
    supported = [r["text"] for r in core if r["status"] == "supported"]
    gaps = qa["remaining_weaknesses"]
    keywords = [r["text"] for r in reqs["ats_keywords"] if r["status"] == "supported"]
    job = plan["job"]
    def bullets(items: list[str], limit=8) -> str:
        return "\n".join(f"- {item}" for item in items[:limit]) or "- None confirmed from the approved sources."
    body = (f"# {job['title']} — {job['company']}\n\n"
            "Review the evidence links and gaps in the tailoring plan before application.\n\n"
            f"## Strongest alignment areas\n\n{bullets(supported)}\n\n"
            f"## Weaknesses or missing requirements\n\n{bullets(gaps, 15)}\n\n"
            f"## ATS keyword alignment\n\n{bullets(keywords)}\n\n"
            "## Recommended resume emphasis\n\n"
            "Emphasize the source-backed bullets selected in the tailoring plan; keep uncertain requirements as gaps.\n\n"
            "## Interview/application considerations\n\n"
            "Prepare concrete examples for supported requirements and address the listed gaps honestly.\n\n"
            f"QA: Word {qa['word_pages']} pages; exported PDF {qa['pdf_pages']} pages; no automated QA issues.\n")
    path = ROOT / "output/match-reports" / f"Dave-Call+{job['safe_company']}+{job['job_number']}.md"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(body, encoding="utf-8")
    return path


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    p = sub.add_parser("plan", help="Extract job requirements and source-backed tailoring plan")
    p.add_argument("listing", type=Path)
    p = sub.add_parser("generate", help="Generate DOCX from an existing plan")
    p.add_argument("plan", type=Path)
    p = sub.add_parser("qa", help="Check DOCX and run authoritative Word pagination")
    p.add_argument("plan", type=Path)
    args = parser.parse_args()
    if args.command == "plan":
        plan = create_plan(args.listing.resolve())
        name = f"{plan['job']['safe_company']}+{plan['job']['job_number']}"
        scratch = ROOT / "scratch" / name
        scratch.mkdir(parents=True, exist_ok=True)
        path = scratch / "tailoring-plan.json"
        path.write_text(json.dumps(plan, indent=2, ensure_ascii=False), encoding="utf-8")
        print(path)
        return 0
    plan = json.loads(args.plan.read_text(encoding="utf-8"))
    verify_plan(plan)
    name = f"{plan['job']['safe_company']}+{plan['job']['job_number']}"
    output = ROOT / "output/resumes" / resume_filename(
        plan["job"]["company"], plan["job"]["title"], plan["job"]["job_number"])
    scratch = ROOT / "scratch" / name
    if args.command == "generate":
        make_resume(plan, output)
        report = native_qa(plan, output, scratch)
        match = str(write_match_report(plan, report)) if report["status"] == "pass" else None
        print(json.dumps({"docx": str(output), "match_report": match, **report}, indent=2))
        return 0 if report["status"] == "pass" else 1
    if not output.exists():
        old_output = ROOT / "output/resumes" / f"Dave-Call+{name}.docx"
        if old_output.exists():
            output = old_output
    if not output.exists():
        raise FileNotFoundError(output)
    report = native_qa(plan, output, scratch)
    match = str(write_match_report(plan, report)) if report["status"] == "pass" else None
    print(json.dumps({"match_report": match, **report}, indent=2))
    return 0 if report["status"] == "pass" else 1


if __name__ == "__main__":
    sys.exit(main())
