"""Build Scout 722's reviewed resume using the current master and Gecko layout."""
import json
from pathlib import Path

import requests
from bs4 import BeautifulSoup
from docx import Document

from gecko_v2 import create_plan, make_resume, native_qa, source_text

ROOT = Path(__file__).resolve().parents[1]
KEY = "111d1443-ce65-40fd-98f5-bd701825e910"
NAME = f"Flex+{KEY}"


def main():
    source_text()  # Fresh read of the authoritative master DOCX for this generation.
    scratch = ROOT / "scratch" / NAME
    scratch.mkdir(parents=True, exist_ok=True)
    response = requests.get(f"https://api.lever.co/v0/postings/Flex/{KEY}", timeout=30)
    response.raise_for_status()
    listing = response.json()
    (scratch / "lever-listing.json").write_text(json.dumps(listing, indent=2), encoding="utf-8")
    sections = []
    for field in ("description", "additional", "salaryDescription"):
        text = BeautifulSoup(listing.get(field, ""), "html.parser").get_text("\n", strip=True)
        if text:
            sections.append(text)
    for group in listing.get("lists", []):
        sections.append(group["text"] + "\n" + BeautifulSoup(group["content"], "html.parser").get_text("\n", strip=True))
    archive = ROOT / "input/job-descriptions" / f"{NAME}.md"
    metadata = f"""# Head of Growth Marketing

- **Scout ID:** 722
- **Job Number:** {KEY}
- **Company:** Flex
- **Job Title:** Head of Growth Marketing
- **Location:** Remote / USA
- **Work Arrangement:** Remote
- **Job Type:** Full-time
- **Salary:** USD $275,000-$450,000 annually; combination of base, stock, and bonus
- **Source:** web-careers
- **Date Discovered:** 2026-09-21
- **URL:** {listing['hostedUrl']}
- **URL Status:** Official Lever listing retrieved successfully

## Full listing

"""
    archive.write_text(metadata + "\n\n".join(sections) + "\n", encoding="utf-8")
    plan = create_plan(archive)
    plan["resume"]["skills"] = ["Google Ads", "Microsoft Ads", "Meta", "SEO", "CRO", "attribution", "forecasting", "GA4", "Google Tag Manager", "Looker Studio", "Tableau", "SQL", "Python", "B2B", "e-commerce", "leadership"]
    plan["resume"]["selected_evidence_ids"] = [
        "E021", "E002", "E004", "E003",
        "E006", "E005", "E007", "E008",
        "E009", "E010", "E012",
        "E013", "E014", "E015",
        "E017", "E019", "E020", "E024",
    ]
    output = ROOT / "output/resumes" / f"Dave-Call+{NAME}.docx"
    make_resume(plan, output)
    # The base renderer enforces the master identity. These two editorial fields
    # are reviewed paraphrases; job history, evidence bullets and contact stay sourced.
    headline = "Growth Marketing | Paid Acquisition, Analytics & Team Leadership"
    summary = ("Performance marketing leader with 14+ years across agency, e-commerce, and multi-channel customer acquisition. "
               "Managed up to $30 million per month in paid media with a four-person team and led an eight-person e-commerce team. "
               "Combines paid media strategy, conversion testing, customer journey analysis, and forecasting with hands-on reporting and measurement. "
               "Experienced coordinating marketing with sales, product, and operations and translating business goals into campaign and investment decisions.")
    doc = Document(output)
    for paragraph in doc.paragraphs:
        if paragraph.text == plan["resume"]["headline"].upper():
            paragraph.runs[0].text = headline.upper()
        elif paragraph.text == plan["resume"]["summary"]:
            paragraph.runs[0].text = summary
    # Balance the two pages and keep the earlier-career roles together.
    doc.tables[2].cell(0, 0).paragraphs[0].paragraph_format.page_break_before = True
    doc.save(output)
    plan["resume"].update(headline=headline, summary=summary)
    plan["review_notes"].append("Editorial summary supported by master summary and E021, E017, E019, E002, E004, E003, E009. No fintech, ARR or PLG claims added.")
    (scratch / "reviewed-plan.json").write_text(json.dumps(plan, indent=2), encoding="utf-8")
    qa = native_qa(plan, output, scratch)
    print(json.dumps(qa, indent=2))
    if qa["status"] != "pass":
        raise SystemExit(1)
    report = ROOT / "output/match-reports" / f"Dave-Call+{NAME}.md"
    report.write_text(f"""# Head of Growth Marketing - Flex

Reviewed against the current master DOCX and the full official listing. This is a stretch leadership application.

## Strongest alignment areas

- 14+ years in digital/performance marketing; paid acquisition across Google Ads, Microsoft Ads, Meta, Amazon, and other channels.
- Managed up to $30 million monthly paid media with a four-person team; led an eight-person e-commerce team.
- Grew monthly marketing volume from approximately $200K to $800K and improved ROAS from 1.5 to 3.5 at GRIP6.
- Customer journey analysis, landing-page optimization, CRO, attribution, experimentation, LTV, forecasting, and financial analysis.
- B2B distributor leadership, collaboration with sales and product teams, and technical reporting using SQL, Python, Tableau, GA4, GTM, and Looker Studio.

## Weaknesses or missing requirements

- No documented growth from $50M to $100M to $200M+ ARR. Advertising spend and campaign revenue are not ARR.
- No established fintech or B2B SaaS background, product-led growth ownership, onboarding experiments, or subscription monetization outcomes.
- No explicit Sales/RevOps pipeline ownership, ICP programs, CAC payback optimization, intent platforms, personalization, or trigger-based messaging.
- SEO and website collaboration are supported; deep technical SEO, structured content, GEO, and AI discovery optimization are not established.
- Product/brand launches, press, creator programs, OOH, and broad executive narrative ownership are not established. Ad copy and client communication are supported.
- People management is demonstrated, but the master omits dates and cannot substantiate the specific 3+ year management requirement.

## ATS keyword alignment

Supported: paid acquisition, paid media, search, display, paid social, remarketing, customer acquisition, conversion optimization, SEO, A/B testing, attribution, ROI, ROAS, CPA, LTV, forecasting, KPI reporting, team leadership, B2B, SQL, Python, GA4, Google Tag Manager, Tableau, Looker Studio.

Do not add as experience: fintech, SaaS ARR scaling, PLG, GEO, RevOps, CAC payback, or product onboarding ownership.

## Recommended resume emphasis

Emphasize paid acquisition scale, measurable growth and efficiency, leadership, customer journey analysis, conversion experimentation, and collaboration with sales/product/operations. Retain real employers and titles; frame e-commerce results as transferable experience.

## Interview/application considerations

Prepare examples of budget decisions, holdouts and regression analysis, conversion tests, team leadership, and commercial prioritization. Explain how those methods transfer to a fintech growth function while acknowledging the ARR and PLG gaps. The role is remote in the USA. Listed annual compensation of $275,000-$450,000 combines base, stock, and bonus; it is not stated base salary alone.

## Score rationale

| Area | Points |
| --- | --- |
| Paid acquisition and experimentation | 19/20 |
| Measurement, economics and technical tools | 15/20 |
| Leadership and scope | 12/15 |
| Website, SEO, messaging and brand | 7/15 |
| Self-serve, sales-led and product-led growth | 6/15 |
| Fintech/SaaS and ARR scaling | 0/10 |
| US remote location alignment | 5/5 |
| Total | 64/100 |

## Validation and sources

- Microsoft Word computed 2 pages; Word-exported PDF contains 2 pages. Native validation passed with no content/layout audit issues.
- Factual source: input/master-resume/Dave-Call-resume-9-23-26.docx, freshly read for this generation.
- [Official Flex listing]({listing['hostedUrl']})
- Archived description: input/job-descriptions/{NAME}.md
- Evidence and validation files: scratch/{NAME}/
""", encoding="utf-8")
    print(f"Resume: {output}\nReport: {report}\nListing: {archive}")


if __name__ == "__main__":
    main()
