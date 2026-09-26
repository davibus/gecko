"""Generate source-backed Gecko deliverables for a reviewed Scout ID batch.

Usage: python scripts/generate_requested_scout_batch.py <scout-id>
Tracker updates are deliberately separate and only follow native Word validation.
"""

from __future__ import annotations

import copy
import json
from pathlib import Path
import re
import subprocess
import sys

from silent_subprocess import windows_creationflags

import fitz
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "job-scout"))
from handoff import archive_listing, job_number, safe_name  # noqa: E402
from storage import JobStore  # noqa: E402
from generate_apply_selected_resumes import CONFIG, build_resume  # noqa: E402
from gecko_v2 import source_text, source_hashes, resume_filename  # noqa: E402

FAMILIES = {
    "paid": (37, "PAID SEARCH | STRATEGY, OPTIMIZATION & ANALYTICS",
        "Paid search and performance marketing leader with 14+ years managing Google Ads and Microsoft Ads across agency, ecommerce, and multi-channel programs. Managed 75+ Google Ads accounts with more than $276K in monthly spend, while advising clients on campaign structure, budgets, and growth. Combines hands-on bidding, keyword and audience strategy with attribution, conversion measurement, testing, and clear performance reporting."),
    "performance": (724, "PERFORMANCE MARKETING | GROWTH & ANALYTICS",
        "Performance marketing leader with 14+ years across paid media, ecommerce, agency programs, and customer acquisition. Managed up to $30 million per month in paid media with a four-person team and led an eight-person ecommerce group. Uses attribution, customer journey analysis, testing, and forecasting to guide budget decisions and communicate practical recommendations."),
    "growth": (721, "GROWTH MARKETING | ACQUISITION, TESTING & ANALYTICS",
        "Digital marketing professional with 14+ years growing acquisition across agency, ecommerce, and multi-channel programs. Advised clients on paid media, SEO, email, social, and website opportunities; developed audience and messaging recommendations; and measured conversion and revenue outcomes. Led an eight-person ecommerce team and used dashboards, forecasts, and structured testing to guide growth decisions."),
    "social": (37, "PAID SOCIAL & PERFORMANCE MARKETING | AUDIENCES & MEASUREMENT",
        "Performance marketing professional with 14+ years across paid search, social, ecommerce, and acquisition. Managed Meta, Instagram, and TikTok alongside Google, Microsoft, and Amazon advertising and used audience analysis, bidding tests, attribution, and conversion measurement to improve results. Experienced with budget allocation, reporting automation, and cross-functional campaign decisions."),
    "analytics": (710, "MARKETING ANALYTICS | DATA VISUALIZATION & BUSINESS INSIGHTS",
        "Marketing analytics and digital marketing professional with 14+ years turning campaign, customer, and financial data into practical decisions. Built Tableau and Looker Studio dashboards and automated reporting across more than 10 channels with Python, SQL, Databricks, Excel, and Funnel.io. Used attribution, regression models, and forecasts to guide investment and present KPI recommendations to leaders and clients."),
    "digital": (704, "DIGITAL MARKETING | CLIENT STRATEGY & CAMPAIGN MEASUREMENT",
        "Digital marketing leader with 14+ years across agency client strategy, multi-channel acquisition, ecommerce, and performance measurement. Consulted with clients on business goals, audiences, budget allocation, SEO, website, email, social, and paid media opportunities. Led an eight-person ecommerce team and built reporting that translates campaign data into clear recommendations."),
    "content": (698, "SEARCH MARKETING | SEO COLLABORATION & ANALYTICS",
        "Search and digital marketing professional with 14+ years in paid search, ecommerce, agency consulting, and performance analysis. Conducted keyword and competitive research, advised clients on SEO and website opportunities, and coordinated channel strategy with web and creative teams. Configured GA4 and Google Tag Manager measurement and built dashboards to guide growth decisions."),
    "sales": (45, "AGENCY CLIENT STRATEGY | PAID MEDIA & PERFORMANCE ANALYTICS",
        "Digital marketing leader with 14+ years in paid media, agency account strategy, and client-facing performance analysis. Managed 75+ Google Ads accounts and more than $276K in monthly spend while advising clients on business goals, channel allocation, and growth priorities. Presented data-backed recommendations, managed teams and large budgets, and coordinated with business stakeholders."),
    "ecommerce": (724, "ECOMMERCE GROWTH | PAID MEDIA & CHANNEL ANALYTICS",
        "Ecommerce and digital marketing leader with 14+ years managing acquisition, paid media, and marketplace-related business decisions. Led an eight-person ecommerce team, expanded a B2B distributor into multiple B2C channels, and worked with more than 650 retail partners. Used attribution, conversion analysis, forecasting, and cross-functional reporting to guide budgets and profitable growth."),
}

# Gaps were reviewed against the current master DOCX and available Scout listings.
JOBS = {
    398: ("growth", ["Direct ownership of a full B2B demand-generation funnel, MQL/SQL pipeline, and lifecycle programs is not documented.", "CRM, ABM, marketing automation, and applied AI campaign workflows are unsupported."], "The Scout listing is truncated; verify the full Pattern requirements before applying."),
    707: ("content", ["The master does not document ownership of editorial content strategy, blogs, thought leadership, or AI-search content.", "Creative briefs, video production, and content-team delivery are not established."], "The supplied Wpromote description mixes content marketing with performance-creative duties; clarify the actual scope."),
    716: ("growth", ["The listing requires 15+ years of digital/demand marketing and 8+ years leading strategy and teams; the master supports 14+ years overall and documented teams of four and eight.", "ABM, lead scoring, nurture programs, Salesforce, Marketo, and marketing-sourced pipeline ownership are unsupported."], "Ask whether Pattern will accept adjacent paid-media leadership in place of its demand-generation and ABM requirements."),
    39: ("performance", ["Franchise portfolio marketing and enterprise-wide brand governance are not documented.", "The Scout excerpt does not contain the complete qualifications."], "This appears to be the same Five Star Franchising role as Scout ID 73; verify one active application destination."),
    779: ("social", ["A dedicated paid-social growth engine, Meta CAPI/server-side signal architecture, and B2B pipeline accountability are not documented.", "A fuller Brex posting describes a New York City hybrid role, while Scout lists Salt Lake City; confirm location eligibility."], "Clarify the work location and exact paid-social specialization before applying."),
    71: ("paid", ["The master documents search and shopping work, but not this employer's multi-brand portfolio or specific outdoor-retail business.", "Scout provides only an excerpt; confirm current platform and budget requirements."], "Show hands-on Google/Microsoft Ads, Shopping, Performance Max, ROAS, budgets, and testing."),
    245: ("performance", ["The master does not establish dedicated B2C paid-social creative ownership or Amazon-agency client specialization.", "Full qualifications are unavailable in Scout's short excerpt."], "Emphasize verified Meta/TikTok, Amazon, ecommerce, and campaign measurement work."),
    88: ("digital", ["The listing excerpt does not identify complete platform or responsibility requirements.", "Do not infer unsupported CRM, content, or lifecycle ownership from the Digital Marketing Manager title."], "Verify the full job description and work arrangement before applying."),
    571: ("sales", ["The master does not document recruiting and leading a performance-marketing sales team or a sales quota.", "Sales-call coaching, proposal selling, and new-business closure are unsupported."], "This is primarily a sales-leadership role; ask whether media account strategy qualifies."),
    561: ("growth", ["The master does not document full demand-generation engine ownership or real-estate SaaS pipeline programs.", "CRM lifecycle, sales handoffs, and marketing-sourced pipeline accountability are unsupported."], "Clarify the on-site expectation and the scope of demand-generation ownership."),
    78: ("paid", ["Pattern-specific ecommerce marketplace operations and its exact campaign tools are not documented.", "Scout has only a short excerpt, so confirm the full scope of the role."], "Lead with paid search/display execution, account optimization, and ecommerce results."),
    50: ("analytics", ["The master does not document pest-control lead operations or company-wide marketing analytics leadership.", "Complete platform and reporting requirements are absent from the Scout excerpt."], "Emphasize Tableau, Looker Studio, Python, SQL, Databricks, attribution, and forecasting."),
    4: ("paid", ["The Scout description is incomplete, and role seniority and account scope need confirmation.", "The role appears narrower than the master resume's leadership scope."], "Lead with hands-on SEM/PPC work and ask about account ownership and compensation."),
    568: ("sales", ["The master does not document a B2B sales-closing record, new-business quota, or SEO-service deal ownership.", "Consultative media strategy is not evidence of high-percentage sales closing."], "Confirm this is a sales role rather than a marketing delivery role."),
    569: ("sales", ["The master does not document a B2B sales-closing record, new-business quota, or SEO-service deal ownership.", "Consultative media strategy is not evidence of high-percentage sales closing."], "This appears to duplicate Scout ID 568 under a separate listing key; confirm one active destination."),
    83: ("digital", ["The master documents SEO collaboration, but not end-to-end SEO or content-marketing ownership.", "Scout provides only a short excerpt, limiting requirement certainty."], "Present agency client strategy, paid search/social, reporting, and cross-channel coordination."),
    85: ("digital", ["Website, SEO, and email-channel experience is documented, but full ownership of those programs is not.", "The Scout excerpt omits complete requirements."], "Emphasize paid media, local/search work, analytics setup, and client collaboration."),
    487: ("performance", ["The master does not establish programmatic platform ownership or lifecycle email program management.", "The available Jooble excerpt omits complete qualifications."], "This appears to be the same Brady role as Scout ID 74; verify the employer posting."),
    74: ("performance", ["The master does not establish programmatic platform ownership or lifecycle email program management.", "The available Jooble excerpt omits complete qualifications."], "This appears to be the same Brady role as Scout ID 487; verify the employer posting."),
    107: ("paid", ["Franchise multi-location and automotive-sector account experience is not documented.", "The Jooble excerpt omits the full client and platform requirements."], "Lead with agency account ownership, client communication, and paid-search optimization."),
    31: ("growth", ["Public-sector marketing, government procurement, and channel-partner programs are not documented.", "The Scout description is incomplete, so the exact campaign and compliance requirements need confirmation."], "Position B2B cross-functional and analytics experience as transferable, not public-sector tenure."),
    157: ("paid", ["B2B SaaS client specialization is not documented in the master.", "The Jooble excerpt omits detailed platform and reporting requirements."], "Lead with 75+ agency Google Ads accounts, testing, attribution, and client strategy."),
    344: ("growth", ["The excerpt describes an EMEA-focused growth role; US remote eligibility is unclear despite Scout's location field.", "Testing/certification industry demand generation and lifecycle ownership are not documented."], "Confirm geography and employment eligibility before applying."),
    238: ("performance", ["The excerpt does not establish complete channel, budget, or industry requirements.", "Dedicated ownership of any unspecified lifecycle or CRM systems is unsupported."], "Emphasize multi-channel acquisition, ROAS gains, testing, and reporting."),
    259: ("growth", ["The master does not establish full lifecycle or CRM-led growth program ownership.", "The Jooble listing appears old and provides only a short excerpt."], "Confirm this Boostability opening remains active before applying."),
    73: ("performance", ["Franchise portfolio marketing and enterprise-wide brand governance are not documented.", "The Jooble excerpt does not contain the complete qualifications."], "This appears to be the same Five Star Franchising role as Scout ID 39; verify one active application destination."),
    70: ("ecommerce", ["Vision-care or healthcare channel experience is not documented.", "The Scout excerpt omits the detailed ecommerce-channel requirements."], "Lead with B2B/B2C ecommerce, marketplace decisions, paid media, and cross-functional reporting."),
}

STRENGTHS = {
    "paid": ["14+ years managing Google and Microsoft Ads across agency and ecommerce programs.", "Direct evidence of account scale, search and shopping optimization, attribution, budgets, testing, and client reporting."],
    "performance": ["Large paid-media scope, multi-channel acquisition, ROAS improvement, and team leadership.", "Attribution, forecasting, budget allocation, and reporting automation support performance decisions."],
    "growth": ["14+ years in acquisition, ecommerce, and agency marketing with measurable channel results.", "Audience research, conversion testing, KPI reporting, and cross-functional growth planning are documented."],
    "social": ["Managed Meta, Instagram, and TikTok along with search and ecommerce advertising.", "Audience analysis, attribution, budget allocation, and campaign testing are documented."],
    "analytics": ["Built Tableau and Looker Studio dashboards and automated reporting across 10+ channels.", "Python, SQL, Databricks, attribution, regression, forecasting, and KPI presentations are documented."],
    "digital": ["Agency client consulting and coordinated paid, social, SEO, email, and website recommendations.", "Campaign measurement, team leadership, and cross-functional delivery are documented."],
    "content": ["Keyword and competitive research, SEO collaboration, and website recommendations are documented.", "GA4, Google Tag Manager, dashboards, and conversion analysis support content-adjacent measurement."],
    "sales": ["Agency client consultation, account strategy, performance presentations, and budget recommendations are documented.", "Paid-media and ecommerce knowledge can support informed sales conversations."],
    "ecommerce": ["Led an eight-person ecommerce team and expanded a B2B distributor into B2C channels.", "Marketplace-related decisions, paid media, channel analytics, and financial reporting are documented."],
}


def prepare(scout_id: int) -> dict:
    family, gaps, note = JOBS[scout_id]
    with JobStore() as store:
        job = store.get(scout_id)
    if job is None:
        raise ValueError(f"Scout ID {scout_id} was not found")
    listing = archive_listing(job, ROOT)
    # The short Adzuna/Jooble records often carry derived salary estimates.
    if job.source.casefold() in {"adzuna", "jooble"}:
        content = listing.read_text(encoding="utf-8")
        content = re.sub(r"(?m)^- \*\*Salary:\*\*.*$", "- **Salary:**", content)
        listing.write_text(content, encoding="utf-8")
    company, key = safe_name(job.company), job_number(job)
    template_id, headline, summary = FAMILIES[family]
    cfg = copy.deepcopy(CONFIG[template_id])
    cfg.update(company=company, key=key, title=job.title, headline=headline, summary=summary)
    CONFIG[scout_id] = cfg
    build_resume(scout_id)  # Re-reads the current master DOCX for this job.
    scratch = ROOT / "scratch" / f"{company}+{key}"
    resume = ROOT / "output/resumes" / resume_filename(company, job.title, key)
    report = ROOT / "output/match-reports" / f"Dave-Call+{company}+{key}.md"
    pdf = scratch / "word-export.pdf"
    status_path = scratch / "validation-status.json"
    result = subprocess.run(["powershell.exe", "-NoProfile", "-NonInteractive", "-WindowStyle", "Hidden",
                             "-ExecutionPolicy", "Bypass", "-File",
                             str(ROOT / "scripts/validate_word_native.ps1"), "-DocxPath", str(resume),
                             "-PdfPath", str(pdf), "-ResultPath", str(status_path)],
                            cwd=ROOT, capture_output=True, text=True,
                            creationflags=windows_creationflags())
    native = json.loads(status_path.read_text(encoding="utf-8-sig")) if status_path.exists() else {}
    if result.returncode or native.get("status") != "native-valid" or native.get("word_pages") != 2 or native.get("pdf_pages") != 2:
        raise RuntimeError(f"Native Word validation failed for {scout_id}: {result.stdout} {result.stderr} {native}")
    doc = Document(resume)
    if len(doc.tables) != 5 or doc.styles["Normal"].font.size.pt != 11 or doc.styles["Normal"].paragraph_format.line_spacing != 1.15:
        raise RuntimeError(f"Structural DOCX validation failed for {scout_id}")
    if not doc.paragraphs[2].text.endswith("linkedin.com/in/mdavidcall"):
        raise RuntimeError(f"Header validation failed for {scout_id}")
    with fitz.open(pdf) as pages:
        if len(pages) != 2 or any(max((b[3] for b in page.get_text("blocks")), default=0) > page.rect.height - 18 for page in pages):
            raise RuntimeError(f"PDF layout validation failed for {scout_id}")
    # Re-read the master at evaluation/report time, independent of the generator call.
    master = source_text()
    if "Built Tableau dashboards" not in master or "75+ Google Ads accounts" not in master:
        raise RuntimeError("Current master DOCX did not contain expected source evidence")
    brief = len(job.description) < 1000
    strengths = "\n".join(f"- {x}" for x in STRENGTHS[family])
    weakness = "\n".join(f"- {x}" for x in gaps)
    ats = "; ".join(body.rstrip(".") for _lead, body in cfg["skills"][:3])
    source_notice = ("Scout holds only a short excerpt. Confirm the full employer listing before applying. "
                     if brief else "")
    body = (f"# {job.title} â€” {job.company}\n\n"

            f"**Scout ID:** {scout_id}  \n**Job Number:** `{key}`  \n**Company:** {job.company}  \n"
            f"**Job Title:** {job.title}  \n**Job Link:** {job.canonical_url or job.url}  \n"
            f"**Generated Resume:** `output/resumes/{resume.name}`\n\n"
            f"## Strongest alignment areas\n\n{strengths}\n\n"
            f"## Weaknesses or missing requirements\n\n{weakness}\n\n"
            f"## ATS keyword alignment\n\nSupported master-resume terms: {ats}.\n\n"
            f"Do not add tools, tenure, industry experience, or outcomes that are absent from the current master DOCX.\n\n"
            f"## Recommended resume emphasis\n\nLead with the source-backed experience and skills highlighted above; "
            f"keep unsupported role requirements as gaps.\n\n"
            f"## Interview/application considerations\n\n{source_notice}{note}\n\n"
            f"**Validation:** Microsoft Word computed 2 pages; Word-exported PDF has 2 pages. "
            f"Master SHA-256: `{source_hashes()[str((ROOT / 'input/master-resume/Dave-Call-resume-9-23-26.docx').relative_to(ROOT))]}`.\n")
    report.parent.mkdir(parents=True, exist_ok=True)
    report.write_text(body, encoding="utf-8")
    return {"scout_id": scout_id, "company": job.company, "job_number": key,
            "resume": str(resume), "report": str(report), "listing": str(listing),
            "word_pages": native["word_pages"], "pdf_pages": native["pdf_pages"]}


if __name__ == "__main__":
    if len(sys.argv) != 2 or not sys.argv[1].isdigit():
        raise SystemExit("Usage: python scripts/generate_requested_scout_batch.py <scout-id>")
    print(json.dumps(prepare(int(sys.argv[1])), indent=2))

