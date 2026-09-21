"""Build and validate the Gecko resume for Trove Brands job 5813466800."""

from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.shared import RGBColor

from generate_trove_brands_resume import build_resume as build_content_resume
from generate_vitality_medical_resume import export_and_validate, set_run


ROOT = Path(__file__).resolve().parents[1]
JOB_KEY = "5813466800"
COMPANY = "Trove-Brands"
OUTPUT = ROOT / "output" / "resumes" / f"Dave-Call+{COMPANY}+{JOB_KEY}.docx"
SCRATCH = ROOT / "scratch" / f"{COMPANY}+{JOB_KEY}"
PDF = SCRATCH / "resume-preview.pdf"
CONTENT_TEMPLATE = SCRATCH / "base-template.docx"


def clear_content(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_plain(paragraph, text, *, size=11, bold=False, color=None):
    clear_content(paragraph)
    set_run(paragraph.add_run(text), size=size, bold=bold, color=color)


def build_resume(path: Path):
    SCRATCH.mkdir(parents=True, exist_ok=True)
    build_content_resume(CONTENT_TEMPLATE, SCRATCH)
    doc = docx.Document(CONTENT_TEMPLATE)
    paragraphs = doc.paragraphs

    set_plain(
        paragraphs[1],
        "ASSOCIATE DIGITAL MARKETING DIRECTOR | E-COMMERCE GROWTH",
        size=11.5,
        bold=True,
        color=RGBColor(40, 70, 110),
    )
    set_plain(
        paragraphs[4],
        "Director-level digital marketing and ecommerce leader with 14+ years driving measurable acquisition, revenue, "
        "and channel growth for consumer and B2B businesses. Leads teams, budgets, forecasting, and full-funnel programs "
        "across paid media, website performance, SEO, email, conversion optimization, and analytics. Proven at scaling "
        "monthly volume from $200K to $800K while improving ROAS from 1.5 to 3.5, managing programs producing more than "
        "$9M in revenue, and translating attribution, customer value, and performance data into decisions for senior "
        "leaders and cross-functional partners.",
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


if __name__ == "__main__":
    build_resume(OUTPUT)
    export_and_validate(OUTPUT, PDF, SCRATCH)
