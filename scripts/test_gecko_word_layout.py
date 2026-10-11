"""Word-rendered regression for the shared employment-heading layout."""

import tempfile
import unittest
from pathlib import Path

from docx import Document

import gecko_v2 as gecko


@unittest.skipUnless(__import__("sys").platform == "win32", "Native Word validation requires Windows")
class GeckoWordLayoutTests(unittest.TestCase):
    def test_sample_has_two_pages_and_every_employment_heading_is_one_line(self):
        listing = gecko.ROOT / "input/job-descriptions/Vitality-Medical+b1e13fda0b6ba9cb.md"
        plan = gecko.create_plan(listing)
        with tempfile.TemporaryDirectory(prefix="heading-layout-", dir=gecko.ROOT / "scratch") as temp:
            scratch = Path(temp)
            docx = scratch / "sample-resume.docx"
            report = gecko.generate_two_page_resume(plan, docx, scratch)
            self.assertEqual(report["status"], "pass", report["issues"])
            self.assertEqual((report["word_pages"], report["pdf_pages"]), (2, 2))
            document = Document(docx)
            self.assertEqual(len(document.tables) - 1, len(gecko.master_jobs()))
            for table, job in zip(document.tables[1:], gecko.master_jobs()):
                self.assertEqual(table.cell(0, 0).text, gecko.employment_heading(job))
                self.assertEqual(table.cell(0, 1).text, gecko.employment_dates(job))


if __name__ == "__main__":
    unittest.main()
