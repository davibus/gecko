"""Build and validate the Gecko resume for Extra Space Storage job 5830372710."""

from pathlib import Path

import docx
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor

from generate_vitality_medical_resume import (
    BODY_SIZE, FONT, LINE_SPACING, MUTED, NAVY, TEXT, export_and_validate,
    keep_table_together, remove_table_borders, set_cell_margins, set_run,
)


ROOT = Path(__file__).resolve().parents[1]
JOB_KEY = "5830372710"
COMPANY = "Extra-Space-Storage"
OUTPUT = ROOT / "output" / "resumes" / f"Dave-Call+{COMPANY}+{JOB_KEY}.docx"
SCRATCH = ROOT / "scratch" / f"{COMPANY}+{JOB_KEY}"
PDF = SCRATCH / "resume-preview.pdf"


def build_resume(path: Path):
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(0.36)
        section.bottom_margin = Inches(0.36)
        section.left_margin = Inches(0.48)
        section.right_margin = Inches(0.48)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY_SIZE)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.line_spacing = LINE_SPACING
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    doc.styles["List Bullet"].font.name = FONT
    doc.styles["List Bullet"].font.size = Pt(BODY_SIZE)

    def paragraph(after=0, before=0, keep=False, align=None):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = LINE_SPACING
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.keep_with_next = keep
        if align is not None:
            p.alignment = align
        return p

    def section_header(title, before=3.0):
        p = paragraph(after=1.5, before=before, keep=True)
        set_run(p.add_run(title.upper()), size=11, bold=True, color=NAVY)
        p._p.get_or_add_pPr().append(parse_xml(
            f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" '
            'w:sz="6" w:space="1" w:color="102C57"/></w:pBdr>'
        ))

    def job_header(title, company, location, before=1.0):
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        remove_table_borders(table)
        keep_table_together(table)
        for index, width in enumerate((Inches(5.95), Inches(1.55))):
            cell = table.cell(0, index)
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
        left = table.cell(0, 0).paragraphs[0]
        left.paragraph_format.space_before = Pt(before)
        left.paragraph_format.space_after = Pt(0.5)
        left.paragraph_format.line_spacing = LINE_SPACING
        left.paragraph_format.keep_with_next = True
        set_run(left.add_run(title), bold=True, color=NAVY)
        set_run(left.add_run(" | "), color=MUTED)
        set_run(left.add_run(f"{company} — {location}"), bold=True)
        right = table.cell(0, 1).paragraphs[0]
        right.paragraph_format.space_before = Pt(before)
        right.paragraph_format.space_after = Pt(0.5)
        right.paragraph_format.keep_with_next = True
        # Intentionally blank: preserve Gecko's right-side date cell without visible dates.

    def bullet(lead, body, after=0.2):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.first_line_indent = Inches(-0.02)
        p.paragraph_format.line_spacing = LINE_SPACING
        p.paragraph_format.space_after = Pt(after)
        set_run(p.add_run(lead), bold=True)
        set_run(p.add_run(body))

    def skill(lead, body, after=0.8):
        p = paragraph(after=after)
        p.paragraph_format.left_indent = Inches(0.10)
        set_run(p.add_run(lead), bold=True, color=NAVY)
        set_run(p.add_run(body))

    p = paragraph(after=0.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_run(p.add_run("DAVE CALL"), size=19, bold=True, color=NAVY)
    p = paragraph(after=1.2, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_run(
        p.add_run("PAID SEARCH ANALYST | PPC & PERFORMANCE MARKETING"),
        size=11.5, bold=True, color=RGBColor(40, 70, 110),
    )
    p = paragraph(after=3.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_run(
        p.add_run(
            "Lehi, UT 84043   •   (530) 507-8269   •   mdavidcall@gmail.com   •   "
            "linkedin.com/in/mdavidcall"
        ),
        size=9.5, color=MUTED,
    )

    section_header("Professional Summary", before=2.0)
    p = paragraph(after=2.0)
    set_run(p.add_run(
        "Hands-on paid-search and performance-marketing specialist with 14+ years of experience building, "
        "monitoring, testing, and improving high-volume campaigns. Deep background in Google Ads, Bing Ads, "
        "desktop editors, keyword and search-query analysis, ad copy, bid testing, budget pacing, account health, "
        "and KPI reporting. Combines careful execution with advanced Excel and analytics skills, clear communication, "
        "and a consistent focus on process quality and measurable business results."
    ))

    section_header("Core Competencies & Technical Skills")
    skill("Paid Search Execution: ", "Google Ads, Bing Ads, Google Ads Editor, Bing Editor, campaign and ad-group builds, keyword research, search-query analysis, negative keywords, ad copy, bidding, and account optimization.")
    skill("Testing & Account Health: ", "Bid experiments, ad and landing-page testing, A/B and multivariate testing, pacing, anomaly monitoring, attribution, conversion paths, CRO, CPA, ROAS, and LTV.")
    skill("Reporting & Data Analysis: ", "Advanced Excel, pivot tables, GA4, Google Tag Manager, Tableau, Looker Studio, Funnel.io, Adobe Analytics, Python, SQL, Databricks, and automated reporting.")
    skill("Additional Media & Platforms: ", "Google Display Network, shopping, Demand Gen-related campaign experience, YouTube, video, mobile, remarketing, Meta, Amazon, LinkedIn, Salesforce, HubSpot, and WordPress.", after=1.2)

    section_header("Professional Experience")
    job_header("Marketing Strategist / Analyst", "1-800 Contacts", "Draper, UT")
    bullet("Paid Search Analysis: ", "Evaluated max CPC versus smart bidding, brand and non-brand CPA, Performance Max and Shopping mix, audiences, attribution, and incremental CPA across scale changes.")
    bullet("Testing & Forecasting: ", "Analyzed holdouts, geo promotions, bidding, reach-to-conversion funnels, lifetime value, and regression scenarios to guide investment decisions.")
    bullet("Automated Reporting: ", "Used Python, SQL, Databricks, Excel, Funnel.io, and Tableau to automate accurate weekly reporting across more than 10 marketing platforms.")

    job_header("Director of Digital Marketing", "GRIP6", "Midvale, UT")
    bullet("Hands-On Platform Management: ", "Managed Google Ads, Bing Ads, Amazon, Facebook, Instagram, and TikTok across a $400K monthly budget.")
    bullet("Performance Improvement: ", "Increased monthly volume from $200K to $800K while improving ROAS from 1.5 to 3.5.")
    bullet("Measurement & Reporting: ", "Implemented custom GA4 and GTM reporting and built Looker Studio KPIs for campaign, website, email, inventory, and budget decisions.")

    doc.add_page_break()

    job_header("SEM Manager / Project Manager", "The Infinite Agency", "Dallas, TX", before=0)
    bullet("Account Builds & Scale: ", "Built and managed 75+ Google Ads accounts totaling more than $276K per month across search, display, mobile, remarketing, YouTube, and social.")
    bullet("Monitoring & Quality: ", "Created scripts for continuous campaign monitoring and a standardized account-quality system adopted across the agency.")
    bullet("Optimization & Testing: ", "Used analytical tools to make data-driven recommendations and created conversion models, attribution models, and multivariate tests.")
    bullet("Client Reporting: ", "Consolidated performance data with Funnel.io and presented actionable results through Looker Studio reporting.")

    job_header("E-Commerce Manager", "LifeSpan Fitness", "Salt Lake City, UT")
    bullet("Campaign Operations: ", "Managed 150+ campaigns and a $400K budget across Google Ads, Bing Ads, Gemini, and Amazon, producing more than $9M in revenue—35% of company revenue.")
    bullet("Search Optimization: ", "Performed keyword and competitive research, search-query and channel analysis, landing-page optimization, bid strategy development, ad-copy creation, and bid experiments.")
    bullet("Editors & Analysis: ", "Used Excel, Google Ads Editor, Bing Editor, Google Analytics, Magento, ACCTivate, and QuickBooks for campaign, performance, and business reporting.")
    bullet("Process Automation: ", "Implemented human-assisted JavaScript automation based on promotions and business goals while maintaining hands-on oversight.")

    job_header("Director of E-Commerce & Marketing", "Intercon Inc.", "Salt Lake City, UT")
    bullet("Data & Process Improvement: ", "Developed and maintained cross-department databases and reporting while implementing systems that increased productivity, transparency, and collaboration.")
    bullet("Cross-Functional Execution: ", "Worked with IT, operations, finance, product development, HR, and sales while managing changing priorities and project-level cost/benefit decisions.")

    section_header("Education, Languages & Certifications", before=1.5)
    p = paragraph(after=0.6)
    set_run(p.add_run("Brigham Young University — Bachelor of Science in Business Management (2008)"), bold=True)
    set_run(p.add_run("  |  Entrepreneur of the Year team; Regional Business Plan Competition winner, two consecutive years"), color=MUTED)
    p = paragraph(after=0.6)
    set_run(p.add_run("Languages: "), bold=True, color=NAVY)
    set_run(p.add_run("English (Native)  •  Spanish (Fluent)"))
    p = paragraph(after=0)
    set_run(p.add_run("Certifications: "), bold=True, color=NAVY)
    set_run(p.add_run("Google Ads  •  Google Analytics  •  Facebook Blueprint  •  Bing Ads  •  Amazon Marketing Services"))

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


if __name__ == "__main__":
    build_resume(OUTPUT)
    export_and_validate(OUTPUT, PDF, SCRATCH)
