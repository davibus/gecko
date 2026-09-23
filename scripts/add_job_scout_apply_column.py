"""One-time, in-place Job Scout Apply? column migration for XLSX and Google Sheets."""

from __future__ import annotations

from copy import copy
from pathlib import Path
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "job-scout"))
from google_sheets_sync import GoogleSheetsConfig, build_sheets_service, load_environment_files  # noqa: E402
from tracker_sync import HEADERS, PRE_APPLY_HEADERS, _save_atomic  # noqa: E402

WORKBOOK = ROOT / "output" / "job-tracker.xlsx"
SHEET = "Job Scout"


def migrate_workbook() -> str:
    workbook = load_workbook(WORKBOOK)
    sheet = workbook[SHEET]
    headers = [cell.value for cell in sheet[1]]
    if headers == HEADERS:
        return "already present"
    if headers != PRE_APPLY_HEADERS:
        raise ValueError("Unexpected Job Scout XLSX headers; refusing to move user columns")
    if sheet.tables or sheet.data_validations.dataValidation or sheet.merged_cells.ranges:
        raise ValueError("Unexpected Excel structure; refusing automatic column insertion")

    old_widths = {key: value.width for key, value in sheet.column_dimensions.items()}
    old_filter = sheet.auto_filter.ref
    old_rules = [(str(area.sqref), [copy(rule) for rule in sheet.conditional_formatting[area]])
                 for area in sheet.conditional_formatting]
    sheet.insert_cols(6)
    for column in range(len(PRE_APPLY_HEADERS), 5, -1):
        source = get_column_letter(column)
        target = get_column_letter(column + 1)
        if source in old_widths:
            sheet.column_dimensions[target].width = old_widths[source]
    sheet.column_dimensions["F"].width = 12
    for row in range(1, sheet.max_row + 1):
        reference = sheet.cell(row, 5)
        target = sheet.cell(row, 6)
        if reference.has_style:
            target._style = copy(reference._style)
        target.number_format = "General"
    sheet["F1"] = "Apply?"
    if old_filter:
        min_col, min_row, max_col, max_row = range_boundaries(old_filter)
        sheet.auto_filter.ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col + 1)}{max_row}"
    sheet.conditional_formatting._cf_rules.clear()
    for area, rules in old_rules:
        min_col, min_row, max_col, max_row = range_boundaries(area)
        new_area = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col + 1)}{max_row}"
        for rule in rules:
            if rule.formula:
                rule.formula = [formula.replace("$G", "$H").replace("$S", "$T").replace("$U", "$V")
                                for formula in rule.formula]
            sheet.conditional_formatting.add(new_area, rule)
    assert [cell.value for cell in sheet[1]] == HEADERS
    _save_atomic(workbook, WORKBOOK)
    return "inserted F"


def migrate_google(service, spreadsheet_id: str) -> str:
    metadata = service.spreadsheets().get(spreadsheetId=spreadsheet_id,
        fields="sheets(properties(sheetId,title,gridProperties),basicFilter,conditionalFormats,tables)").execute()
    matches = [sheet for sheet in metadata["sheets"] if sheet["properties"]["title"] == SHEET]
    if len(matches) != 1:
        raise ValueError("Expected exactly one Job Scout Google worksheet")
    sheet = matches[0]
    sheet_id = sheet["properties"]["sheetId"]
    values = service.spreadsheets().values().get(spreadsheetId=spreadsheet_id,
        range="'Job Scout'!A1:X1").execute().get("values", [[]])[0]
    if "Apply?" in values and values.index("Apply?") == 5:
        return "already present"
    if len(values) != len(PRE_APPLY_HEADERS) or set(values) != set(PRE_APPLY_HEADERS) or values[4] != "Gecko Status":
        raise ValueError(f"Unexpected Google Job Scout headers; refusing to move user columns: {values!r}")
    requests = [
        {"insertDimension": {"range": {"sheetId": sheet_id, "dimension": "COLUMNS",
                                        "startIndex": 5, "endIndex": 6}, "inheritFromBefore": True}},
        {"updateCells": {"range": {"sheetId": sheet_id, "startRowIndex": 0,
                                     "endRowIndex": 1, "startColumnIndex": 5, "endColumnIndex": 6},
                         "rows": [{"values": [{"userEnteredValue": {"stringValue": "Apply?"}}]}],
                         "fields": "userEnteredValue"}},
    ]
    service.spreadsheets().batchUpdate(spreadsheetId=spreadsheet_id,
        body={"requests": requests}).execute()
    return "inserted F"


def main() -> None:
    load_environment_files(ROOT)
    config = GoogleSheetsConfig.from_environment()
    service = build_sheets_service(config)
    print("Google Sheet:", migrate_google(service, config.spreadsheet_id))
    print("XLSX:", migrate_workbook())
    print("Google URL:", f"https://docs.google.com/spreadsheets/d/{config.spreadsheet_id}/edit")


if __name__ == "__main__":
    main()
