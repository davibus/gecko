"""Create and Word-validate the Studio McGee Gecko deliverables."""

from __future__ import annotations

import copy
import json
from pathlib import Path
import subprocess

from silent_subprocess import windows_creationflags

import fitz
from docx import Document

from generate_apply_selected_resumes import CONFIG, build_resume
from gecko_v2 import ROOT, resume_filename, source_hashes, source_text


JOB_ID = 990001
COMPANY = "Studio-McGee"
TITLE = "Digital Performance Marketing Manager"
KEY = "4ead51309acecc1f"
LISTING = ROOT / "input/job-descriptions" / f"{COMPANY}+{KEY}.md"
SCRATCH = ROOT / "scratch" / f"{COMPANY}+{KEY}"
RESUME = ROOT / "output/resumes" / resume_filename(COMPANY, TITLE, KEY)
REPORT = ROOT / "output/match-reports" / f"Dave-Call+{COMPANY}+{KEY}.md"


def main() -> None:
    if not LISTING.exists():
        raise FileNotFoundError(LISTING)
    # Re-read the current master for this assessment and again in the generator.
    master = source_text()
    for phrase in ("75+ Google Ads accounts", "affiliate marketing", "Built Tableau dashboards"):
        if phrase not in master:
            raise RuntimeError(f"The current master does not support expected evidence: {phrase}")
    cfg = copy.deepcopy(CONFIG[37])
    cfg.update(
        company=COMPANY,
        key=KEY,
        title=TITLE,
        headline="DIGITAL PERFORMANCE MARKETING | PAID MEDIA & ECOMMERCE GROWTH",
        summary=(
            "Performance marketing leader with 14+ years across paid search, social, ecommerce, "
            "and multi-channel acquisition. Managed Google Ads, Microsoft Ads, Meta, TikTok, "
            "and Amazon advertising while guiding large budgets, audience strategy, and "
            "conversion testing. Combines SEO collaboration and affiliate marketing experience "
            "with attribution, forecasting, and clear GA4, Looker Studio, and Tableau reporting "
            "to guide profitable investment decisions."
        ),
        skills=[
            ["Paid Media & Search: ", "Google Ads, Microsoft Ads, Meta, Instagram, TikTok, Amazon, Search, Shopping, Display, YouTube, remarketing, keyword strategy, and paid social."],
            ["Channel Coordination: ", "SEO collaboration, affiliate marketing experience, website and creative partnership, audience research, messaging, and cross-channel campaign planning."],
            ["Budget & Economics: ", "Budget allocation and pacing, CPA, ROAS, ROI, LTV, forecasting, financial reporting, holdouts, regression models, and investment scenarios."],
            ["Measurement & Testing: ", "GA4, Google Tag Manager, Looker Studio, Tableau, Funnel.io, attribution, conversion tracking, A/B and multivariate testing."],
            ["Leadership & Delivery: ", "Four-person paid media team management, eight-person ecommerce team leadership, agency client strategy, stakeholder presentations, and cross-functional coordination."],
        ],
    )
    CONFIG[JOB_ID] = cfg
    build_resume(JOB_ID)
    if not RESUME.exists():
        raise FileNotFoundError(RESUME)
    SCRATCH.mkdir(parents=True, exist_ok=True)
    pdf = SCRATCH / "word-export.pdf"
    status_path = SCRATCH / "validation-status.json"
    proc = subprocess.run(
        ["powershell.exe", "-NoProfile", "-NonInteractive", "-WindowStyle", "Hidden",
         "-ExecutionPolicy", "Bypass", "-File",
         str(ROOT / "scripts/validate_word_native.ps1"), "-DocxPath", str(RESUME),
         "-PdfPath", str(pdf), "-ResultPath", str(status_path)],
        cwd=ROOT, capture_output=True, text=True, errors="replace",
        creationflags=windows_creationflags(),
    )
    status = json.loads(status_path.read_text(encoding="utf-8-sig")) if status_path.exists() else {}
    if proc.returncode or status.get("status") != "native-valid" or status.get("word_pages") != 2 or status.get("pdf_pages") != 2:
        raise RuntimeError(f"Native Word validation failed: {proc.stdout} {proc.stderr} {status}")
    doc = Document(RESUME)
    if len(doc.tables) != 5 or doc.styles["Normal"].font.size.pt != 11 or doc.styles["Normal"].paragraph_format.line_spacing != 1.15:
        raise RuntimeError("Final DOCX structure does not meet Gecko requirements")
    if not doc.paragraphs[2].text.endswith("linkedin.com/in/mdavidcall"):
        raise RuntimeError("Header contact line is incorrect")
    with fitz.open(pdf) as pages:
        if len(pages) != 2 or any(max((b[3] for b in page.get_text("blocks")), default=0) > page.rect.height - 18 for page in pages):
            raise RuntimeError("Word PDF shows a possible bottom-page overflow")
    # Re-read the master for the final evaluation rather than using generated text as evidence.
    master = source_text()
    if "affiliate marketing" not in master or "Meta, Instagram, and TikTok" not in master:
        raise RuntimeError("Current master evidence is missing during report evaluation")
    body = f"""# {TITLE} - Studio McGee

**Company:** Studio McGee  
**Job Title:** {TITLE}  
**Job Key (Indeed jk):** `{KEY}`  
**Job Link:** https://www.indeed.com/viewjob?jk={KEY}  
**Generated Resume:** `output/resumes/{RESUME.name}`

## Strongest alignment areas

- The master documents 14+ years across paid search, social, ecommerce, and multi-channel acquisition. It includes Google Ads, Microsoft Ads, Meta, Instagram, TikTok, Display, YouTube, Amazon, and remarketing.
- Paid-media budget and performance evidence is substantial: up to $30 million per month with a four-person team, approximately $400K monthly at GRIP6, and improved ROAS from 1.5 to 3.5 while scaling volume fourfold.
- Keyword and audience strategy, campaign structure, bids, ad copy, competitive research, paid/organic coordination, and creative/website collaboration map closely to the search and agency-stewardship work.
- GA4, Google Tag Manager, Looker Studio, Tableau, attribution, CPA, ROAS, LTV, forecasting, holdouts, and A/B and multivariate testing support the reporting and investment requirements.
- The master lists affiliate marketing experience and documents ecommerce leadership, financial reporting, and cross-functional work with marketing, operations, finance, IT, and product teams.

## Weaknesses or missing requirements

- The master lists affiliate marketing but does not establish ownership of an affiliate program, partner recruitment and activation, affiliate platforms, or incremental-revenue measurement.
- Catalog and retargeting direct mail planning or oversight is not documented.
- Programmatic DSP ownership, TikTok Shop, Google Search Console, SEMRush, and BrightEdge are not documented. The resume does not claim them.
- SEO collaboration is documented; direct ownership of SEO briefs, organic-search optimization, or agency-side SEO direction is not.
- The master establishes agency client strategy but does not clearly establish managing an external paid-media agency from the advertiser side.
- The posting requests 4-6 years of experience; the master shows 14+ years, so role scope and compensation are worth checking.

## ATS keyword alignment

**Supported:** Google Ads, Microsoft Ads, Meta, TikTok, paid search, paid social, SEM, Display, remarketing, affiliate marketing experience, SEO collaboration, GA4, Looker Studio, Tableau, attribution, CPA, ROAS, LTV, budget pacing, forecasting, A/B testing, multivariate testing, ecommerce, cross-functional coordination.

**Unverified or unsupported:** affiliate program management, direct mail, programmatic DSPs, TikTok Shop, Google Search Console, SEMRush, BrightEdge, SEO brief ownership.

## Recommended resume emphasis

Lead with multi-channel paid-media scale, budget stewardship, ecommerce ROAS growth, search/social execution, and analytics. Show the source-backed affiliate and SEO exposure accurately without describing either as a fully owned program.

## Interview/application considerations

Clarify how much of the role is spent on affiliate recruitment, direct mail, and DSP management versus paid-media strategy and agency oversight. Ask about the agency decision rights, investment mix, remote/hybrid expectations, and manager-level scope. Prepare examples of cross-channel budget decisions, test design, and translating performance analysis into recommendations.

**Validation:** Microsoft Word computed 2 pages; Word-exported PDF has 2 pages. Master SHA-256: `{source_hashes()[str((ROOT / 'input/master-resume/Dave-Call-resume-9-23-26.docx').relative_to(ROOT))]}`.
"""
    REPORT.parent.mkdir(parents=True, exist_ok=True)
    REPORT.write_text(body, encoding="utf-8")
    print(json.dumps({"resume": str(RESUME), "report": str(REPORT), "listing": str(LISTING),
                      "word_pages": status["word_pages"], "pdf_pages": status["pdf_pages"]}, indent=2))


if __name__ == "__main__":
    main()
