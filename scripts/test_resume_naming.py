"""Filename limits, collision prevention, and lossless job matching."""
import json
import sys
import tempfile
from pathlib import Path
import unittest
from docx import Document
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'job-scout'))
import resume_naming as naming

class NamingTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.folder=Path(self.temp.name)/'resumes'

    def test_preview_makes_no_files_or_directories(self):
        self.assertEqual(naming.filename('Array','Digital Marketing Manager','6210668004',self.folder),
                         'Dave-Call+Array+Digital-Marketing-Manager+00001.docx')
        self.assertFalse(self.folder.exists())

    def test_long_numeric_uuid_and_matching_suffixes_get_distinct_short_ids(self):
        numbers=['6210668004','7bbbdcf6-7301-4440-8a19-60d9d633d1ea','99999668004']
        names=[naming.filename('Example','Marketing Manager',n,self.folder,reserve=True) for n in numbers]
        self.assertEqual(len(set(names)),3)
        for name,number in zip(names,numbers):
            self.assertRegex(name,r'\+\d{5}\.docx$')
            self.assertEqual(naming.job_number(self.folder/name),number)
        self.assertEqual(naming.filename('Example','Marketing Manager',numbers[0],self.folder,reserve=True),names[0])

    def test_revision_gets_new_short_identifier(self):
        one=naming.filename('Acme','Role','long-id',self.folder,reserve=True)
        two=naming.filename('Acme','Role','long-id',self.folder,reserve=True,revision=True)
        self.assertNotEqual(one,two);self.assertNotIn('+v',two)
        self.assertEqual(naming.job_number(self.folder/two),'long-id')

    def test_existing_short_suffix_is_not_reused(self):
        self.folder.mkdir();(self.folder/'Dave-Call+Old+Role+00001.docx').write_bytes(b'old')
        self.assertTrue(naming.filename('New','Role','new-job',self.folder,reserve=True).endswith('+00002.docx'))

    def test_legacy_long_names_still_resolve(self):
        self.assertEqual(naming.job_number(self.folder/'Dave-Call+Acme+Role+abc123xyz.docx'),'abc123xyz')

    def test_title_is_readable_and_filename_is_bounded(self):
        name=naming.filename('Acme / West','Director of Marketing Strategy '*30,'a'*100,self.folder)
        self.assertLessEqual(len(name),150);self.assertIn('Dave-Call+Acme-West+Director-of-Marketing',name)
        self.assertNotIn('a'*100,name)

    def test_exported_docx_preserves_identity_without_registry(self):
        name=naming.filename('Acme','Role','uuid-full-job-number',self.folder,reserve=True)
        doc=Document();doc.core_properties.identifier='gecko-job:uuid-full-job-number';doc.add_paragraph('Resume')
        exported=Path(self.temp.name)/name;doc.save(exported)
        self.assertEqual(naming.job_number(exported),'uuid-full-job-number')

    def test_short_name_supports_tracker_package_and_gmail_matching(self):
        sys.path.insert(0,str(Path(__file__).resolve().parent))
        from manage_job_tracker import filename_job_number
        from prepare_application_package import validate_resume_identity
        from gmail_response_tracker import _job_number
        name=naming.filename('Acme','Product Manager','source-uuid-123456',self.folder,reserve=True)
        path=self.folder/name
        doc=Document();doc.core_properties.identifier='gecko-job:source-uuid-123456';doc.save(path)
        self.assertEqual(filename_job_number(path),'source-uuid-123456')
        validate_resume_identity(path,{'Company':'Acme','Job Title':'Product Manager'},'source-uuid-123456')
        self.assertEqual(_job_number(path.as_uri()),'source-uuid-123456')

    def test_canonical_queue_publishes_short_name_with_full_job_mapping(self):
        from unittest.mock import patch
        import generate_apply_queue as queue
        root=Path(self.temp.name)
        listing=root/'listing.md';listing.write_text('Complete archived posting',encoding='utf-8')
        number='7bbbdcf6-7301-4440-8a19-60d9d633d1ea'
        plan={'job':{'company':'Acme','safe_company':'Acme','title':'Product Manager','job_number':number}}
        item=queue.QueueRow(2,42,'Acme','Product Manager')
        def generate(_plan,path,folder):
            doc=Document();doc.core_properties.identifier='gecko-job:'+number
            doc.add_paragraph('Professional Summary');doc.add_paragraph('Source-backed marketing experience.')
            doc.add_paragraph('Core Strengths');doc.save(path)
            return {'status':'pass'}
        with patch.object(queue,'ROOT',root),patch.object(queue,'description_is_sufficient',return_value=True), \
             patch.object(queue.gecko_v2,'create_plan',return_value=plan), \
             patch.object(queue.gecko_v2,'generate_two_page_resume',side_effect=generate), \
             patch('resume_visual_review.review',return_value=None):
            result=queue._finish_generation(item,listing,())
        self.assertEqual(result.resume.name,'Dave-Call+Acme+Product-Manager+00001.docx')
        self.assertEqual(naming.job_number(result.resume),number)
        self.assertTrue(result.resume.is_file())

if __name__=='__main__':unittest.main()
