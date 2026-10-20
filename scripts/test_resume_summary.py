"""Summary evidence, relevance, and saved-content regression checks."""
import tempfile
import unittest
from unittest.mock import patch
from pathlib import Path
from docx import Document
import resume_summary as summary

MASTER = Path(__file__).resolve().parents[1] / "input/master-resume/master-resume-v4.docx"


class SummaryTests(unittest.TestCase):
    def test_codex_plan_does_not_require_unused_preset_summary(self):
        import gecko_v2
        with tempfile.TemporaryDirectory() as folder:
            listing = Path(folder)/'example.md'
            listing.write_text('# Senior Product Marketing Manager\n\n- **Company:** Example\n- **Job Number:** example\n- **URL:** https://example.com/job/example\n\n## Job description\n\nOwn marketing strategy and paid search campaigns.',encoding='utf-8')
            with patch.object(summary,'choose',side_effect=ValueError('No substantively distinct summary remains')) as choose:
                plan = gecko_v2.create_plan(listing,defer_summary=True)
                choose.assert_not_called()
                self.assertEqual(plan['resume']['summary'],'')
                with self.assertRaisesRegex(ValueError,'distinct summary'):
                    gecko_v2.create_plan(listing)

    def setUp(self):
        from resume_source import read_master_text
        self.source = read_master_text(MASTER)
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)

    def save(self, text, name):
        doc = Document()
        doc.add_paragraph("Professional Summary")
        doc.add_paragraph(text)
        doc.add_paragraph("Core Strengths")
        doc.add_paragraph("Marketing Analytics")
        path = self.folder / name
        doc.save(path)
        return path

    def test_all_catalog_claims_have_current_source_support(self):
        for key in summary.UNITS:
            if key != "lead":
                summary.render(["lead", key, "crm" if key != "crm" else "clients"], self.source)

    def test_new_crm_work_selected_for_operations_job(self):
        text, selection = summary.choose("crm operations automation reporting integration enablement", self.source, self.folder)
        self.assertIn("crm", summary.TAILORED[selection[1]][1])
        self.assertIn("Go High Level", text)

    def test_missing_source_evidence_stops_generation(self):
        with self.assertRaisesRegex(ValueError, "evidence missing"):
            summary.render(["lead", "crm", "landing"], "Unrelated source")

    def test_composed_paragraphs_have_support_and_different_openings(self):
        openings = set()
        for key in summary.TAILORED:
            text = summary.render(["tailored", key], self.source)
            opening = text.split(".")[0]
            self.assertNotIn(opening, openings)
            openings.add(opening)

    def test_role_family_prevents_ecommerce_focus_for_operations(self):
        text, selection = summary.choose("ecommerce client content crm", self.source, self.folder, "Marketing Operations Specialist")
        self.assertEqual(selection, ["tailored", "operations_data"])

    def test_future_postings_continue_with_verified_focused_composition(self):
        first, _ = summary.choose("crm operations automation reporting integration", self.source, self.folder, "Marketing Operations Specialist")
        self.save(first, "previous.docx")
        text, selection = summary.choose("crm operations automation reporting integration", self.source, self.folder, "Marketing Operations Specialist")
        self.assertEqual(selection[0], "focused")
        self.assertNotEqual(text, first)
        self.assertNotIn(summary.UNITS['lead'][0], text)

    def test_small_summary_change_cannot_hide_same_document(self):
        text = " ".join("source supported marketing accomplishment number " + str(i) for i in range(100))
        self.save("Original summary. " + text, "first.docx")
        candidate = self.save("Different summary. " + text, "second.docx")
        with self.assertRaisesRegex(ValueError, "too similar"):
            summary.assert_unique(candidate, self.folder)

    def test_twenty_similar_postings_get_distinct_evidence(self):
        listing = "Marketing account director agency client strategy leadership team digital performance analytics data reporting growth content ai ecommerce crm operations"
        selections = set()
        for index in range(20):
            text, selection = summary.choose(listing, self.source, self.folder)
            self.assertNotIn(frozenset(selection), selections)
            selections.add(frozenset(selection))
            self.save(text, f"{index}.docx")

    def test_duplicate_content_ignores_document_metadata(self):
        one = self.save("A distinct professional summary.", "first.docx")
        two = self.save("A distinct professional summary.", "second.docx")
        doc = Document(two)
        doc.core_properties.title = "Different filename and metadata are not different content"
        doc.save(two)
        with self.assertRaisesRegex(ValueError, "Duplicate"):
            summary.assert_unique(two, self.folder)

    def test_inspection_failure_is_not_silently_skipped(self):
        (self.folder / "broken.docx").write_text("broken")
        with self.assertRaisesRegex(ValueError, "Cannot check"):
            summary.choose("marketing operations", self.source, self.folder)


if __name__ == "__main__":
    unittest.main()
