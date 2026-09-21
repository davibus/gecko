"""Build and validate the Gecko resume for Digital Marketing job 5330785039."""

from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.shared import RGBColor

from generate_marksmen_growth_resume import build_resume as build_base_resume
from generate_vitality_medical_resume import NAVY, export_and_validate, set_run


ROOT = Path(__file__).resolve().parents[1]
JOB_KEY = "5330785039"
COMPANY = "Digital-Marketing"
OUTPUT = ROOT / "output" / "resumes" / f"Dave-Call+{COMPANY}+{JOB_KEY}.docx"
SCRATCH = ROOT / "scratch" / f"{COMPANY}+{JOB_KEY}"
TEMPLATE = SCRATCH / "base-template.docx"
PDF = SCRATCH / "resume-preview.pdf"


def clear_content(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_plain(paragraph, text, *, size=11, bold=False, color=None):
    clear_content(paragraph)
    set_run(paragraph.add_run(text), size=size, bold=bold, color=color)


def set_labeled(paragraph, lead, body, *, lead_color=None):
    clear_content(paragraph)
    set_run(paragraph.add_run(lead), bold=True, color=lead_color)
    set_run(paragraph.add_run(body))


def build_resume(path: Path):
    SCRATCH.mkdir(parents=True, exist_ok=True)
    build_base_resume(TEMPLATE)
    doc = docx.Document(TEMPLATE)
    paragraphs = doc.paragraphs

    set_plain(
        paragraphs[1],
        "SENIOR SEO & DIGITAL MARKETING ANALYST",
        size=11.5,
        bold=True,
        color=RGBColor(40, 70, 110),
    )
    set_plain(
        paragraphs[4],
        "SEO and digital marketing strategist with 14+ years spanning agency client work, ecommerce, content, website "
        "optimization, analytics, and conversion. Develops practical search strategies through keyword and competitor "
        "research, technical and on-page SEO, content planning, site structure, link building, testing, and performance "
        "measurement. Experienced managing 75+ client accounts, coordinating cross-functional priorities, and translating "
        "traffic, conversion, attribution, and ROI data into clear recommendations. Combines analytical rigor with hands-on "
        "HTML, CSS, JavaScript, Python, SQL, automation, and AI-assisted research workflows.",
    )

    skills = {
        6: ("SEO Strategy & Execution: ", "Technical and on-page SEO, SEO audits, keyword and competitor research, metadata, taxonomy, schema markup, internal content promotion, link building, off-page SEO, local/mobile SEO, Search Console, Semrush, Moz, and SimilarWeb."),
        7: ("Content, Traffic & Conversion: ", "Content strategy and development, copywriting, landing-page optimization, website management, conversion paths, CRO, A/B and multivariate testing, customer journeys, lead generation, and campaign planning."),
        8: ("Analytics & ROI: ", "GA4 and Universal Analytics, Google Tag Manager, Adobe Analytics, Tableau, Looker Studio, Funnel.io, Excel, attribution, LTV, CAC, ROAS, regression analysis, KPI reporting, and forecasting."),
        9: ("Agency & Project Leadership: ", "Multi-client strategy, project management, budget alignment, stakeholder recommendations, sales collaboration, social and content coordination, account management, deadline management, and cross-functional prioritization."),
        10: ("Web, Data & Automation: ", "WordPress, Shopify, Magento, HTML5, CSS3, JavaScript, Python, SQL, VBA, APIs, Git/GitHub, ChatGPT, Claude, Perplexity, Codex, Cursor, and AntiGravity."),
    }
    for index, (lead, body) in skills.items():
        set_labeled(paragraphs[index], lead, body, lead_color=NAVY)

    bullets = {
        12: ("Marketing Analysis: ", "Evaluated attribution, lifetime value, keyword and audience performance, brand/non-brand CPA, bidding approaches, and five-year projections using predictive and regression models."),
        13: ("Full-Funnel Testing: ", "Led reach, comparison, conversion, and lifetime-value testing and assessed Click-to-Call and iSpot attribution systems."),
        14: ("Reporting Automation: ", "Used Python, SQL, Databricks, Excel, Funnel.io, and Tableau to automate accurate weekly reporting across more than 10 marketing platforms."),
        15: ("Investment Guidance: ", "Managed $30 million per month with a team of four, using performance analysis to guide channel investment, testing, and acquisition decisions."),
        16: ("SEO & Website Optimization: ", "Guided on-page SEO and site optimization while implementing custom GA4 reporting and Google Tag Manager configurations."),
        17: ("Digital Growth: ", "Managed Google, Bing, Amazon, Facebook, Instagram, and TikTok; increased monthly volume from $200K to $800K while improving ROAS from 1.5 to 3.5."),
        18: ("Cross-Functional Reporting: ", "Built Looker Studio KPIs supporting inventory, engineering, management, web development, email, and cross-platform budget decisions."),
        19: ("Planning & Prioritization: ", "Created an inventory-planning methodology used across departments to connect product availability, operating needs, and marketing decisions."),
        21: ("Digital & Commerce Transformation: ", "Led ecommerce and marketing for a large distributor, building a functional digital environment and adding multiple B2C platforms while supporting 650+ retail partners."),
        22: ("Information Architecture: ", "Established information architecture and ERP workflows and maintained databases used for accurate financial, catalog, and marketing reporting."),
        23: ("Website & Marketplace Expansion: ", "Developed ecommerce-specific products and expanded Amazon operations into CastleGate, Target.com, Costco.com, and other channels."),
        24: ("Cross-Functional Delivery: ", "Worked with IT, operations, finance, product development, HR, and sales leaders while managing changing priorities and project-level cost/benefit decisions."),
        25: ("Multi-Client Strategy: ", "Advised clients on websites, SEO, email, paid search, display, mobile, remarketing, YouTube, and social based on business needs and analytical findings."),
        26: ("Agency Scale: ", "Built and managed 75+ Google Ads accounts totaling more than $276K per month and created a standardized quality system adopted across the agency."),
        27: ("Search & Conversion Analysis: ", "Used industry analytical platforms to assess accounts and translate findings into website, SEO, channel, and campaign priorities."),
        28: ("Testing & Attribution: ", "Created conversion and attribution models and multivariate tests to improve audience, message, channel, and landing-page decisions."),
        29: ("SEO, Content & CRO: ", "Performed keyword and competitive research, monitored SEO/SEM with multiple tools, optimized landing pages, developed ad copy and creative, and ran A/B and multivariate tests."),
        30: ("Ecommerce Leadership: ", "Led an eight-person ecommerce department responsible for paid search, outreach, lead generation, and customer acquisition across the United States, Canada, and the UK."),
        31: ("Revenue & Campaign Scale: ", "Managed 150+ campaigns and a $400K budget across Google, Bing, Gemini, and Amazon, producing more than $9M in revenue—35% of company revenue."),
        32: ("Web Collaboration: ", "Partnered with web development and used JavaScript-assisted automation to align campaigns, site work, promotions, and business goals."),
        33: ("Performance Planning: ", "Set annual goals, projections, and KPIs with senior management and reported business performance using Magento, Google Analytics, ACCTivate, and QuickBooks."),
    }
    for index, (lead, body) in bullets.items():
        set_labeled(paragraphs[index], lead, body)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


if __name__ == "__main__":
    build_resume(OUTPUT)
    export_and_validate(OUTPUT, PDF, SCRATCH)
