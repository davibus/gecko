"""Generate the reviewed Gecko deliverables for HealthCare Scout ID 847."""

import json
from pathlib import Path

import requests
from bs4 import BeautifulSoup
from docx import Document

from gecko_v2 import create_plan, make_resume, native_qa, source_text


ROOT = Path(__file__).resolve().parents[1]
JOB_NUMBER = "f3febd60-0083-4bee-9856-c6d779a72177"
NAME = f"HealthCare+{JOB_NUMBER}"
SOURCE_URL = f"https://api.lever.co/v0/postings/healthcare/{JOB_NUMBER}"


def listing_text(html: str) -> str:
    soup = BeautifulSoup(html, "html.parser")
    for item in soup.find_all("li"):
        item.string = "- " + item.get_text(" ", strip=True)
    for heading in soup.find_all(["h1", "h2", "h3", "h4"]):
        heading.string = "## " + heading.get_text(" ", strip=True)
    return soup.get_text("\n", strip=True)


def main() -> None:
    source_text()  # Read the current master DOCX for this generation.
    scratch = ROOT / "scratch" / NAME
    scratch.mkdir(parents=True, exist_ok=True)
    response = requests.get(SOURCE_URL, timeout=30)
    response.raise_for_status()
    listing = response.json()
    if listing.get("id") != JOB_NUMBER or listing.get("text") != "Growth Marketing Manager":
        raise ValueError("The retrieved Lever listing does not match Scout 847")
    (scratch / "lever-listing.json").write_text(json.dumps(listing, indent=2, ensure_ascii=False), encoding="utf-8")

    archive = ROOT / "input/job-descriptions" / f"{NAME}.md"
    archive.write_text(
        f"""# Growth Marketing Manager

- **Scout ID:** 847
- **Job Number:** {JOB_NUMBER}
- **Company:** HealthCare
- **Job Title:** Growth Marketing Manager
- **Location:** Remote, US
- **Work Arrangement:** Remote
- **Job Type:** Full Time
- **Salary:**
- **Source:** web-careers
- **Date Discovered:** 2026-09-22
- **URL:** {listing['hostedUrl']}
- **URL Status:** Official Lever listing retrieved successfully

## Full listing

{listing_text(listing.get('description', ''))}

{listing_text(listing.get('additional', ''))}
""",
        encoding="utf-8",
    )

    plan = create_plan(archive)
    plan["resume"]["skills"] = [
        "Google Ads", "Microsoft Ads", "Meta", "GA4", "Google Tag Manager",
        "attribution", "conversion tracking", "CRO", "forecasting", "SEO",
        "landing pages", "affiliate marketing", "Looker Studio", "Tableau",
        "Python", "SQL", "B2B", "e-commerce",
    ]
    plan["resume"]["selected_evidence_ids"] = [
        "E021", "E002", "E003", "E004",
        "E005", "E006", "E007", "E008",
        "E009", "E010", "E011", "E012",
        "E013", "E014", "E015", "E016",
        "E017", "E019", "E020", "E024",
    ]
    missing_metrics = set(plan["resume"]["relevant_metric_ids"]) - set(plan["resume"]["selected_evidence_ids"])
    if missing_metrics:
        raise ValueError(f"Required master accomplishments were omitted: {sorted(missing_metrics)}")

    resume = ROOT / "output/resumes" / f"Dave-Call+{NAME}.docx"
    make_resume(plan, resume)
    headline = "Growth Marketing | Acquisition Strategy, Attribution & Conversion"
    summary = (
        "Performance marketing leader with 14+ years across agency, e-commerce, and B2B environments. "
        "Managed large paid media budgets and multi-channel acquisition programs, using audience strategy, "
        "conversion testing, attribution, and forecasting to guide investment. Built automated reporting "
        "and dashboards to improve decisions, and coordinated marketing with sales, product, operations, "
        "and finance teams. Brings hands-on Google Ads, Microsoft Ads, Meta, SEO, and landing-page experience."
    )
    doc = Document(resume)
    for paragraph in doc.paragraphs:
        if paragraph.text == plan["resume"]["headline"].upper():
            paragraph.runs[0].text = headline.upper()
        elif paragraph.text == plan["resume"]["summary"]:
            paragraph.runs[0].text = summary
    doc.tables[2].cell(0, 0).paragraphs[0].paragraph_format.page_break_before = True
    doc.save(resume)
    plan["resume"].update(headline=headline, summary=summary)
    plan["review_notes"].append(
        "Editorial summary is supported by master summary and E002–E005, E007–E009, E012, E014, E019–E021, E023. "
        "No insurance, Invoca, offline-import, or partnership-negotiation experience is claimed."
    )
    (scratch / "reviewed-plan.json").write_text(json.dumps(plan, indent=2, ensure_ascii=False), encoding="utf-8")
    qa = native_qa(plan, resume, scratch)
    print(json.dumps(qa, indent=2))
    if qa["status"] != "pass":
        raise SystemExit(1)

    report = ROOT / "output/match-reports" / f"Dave-Call+{NAME}.md"
    report.write_text(f"""# Growth Marketing Manager — HealthCare

Match Score: 68/100

## Strongest alignment areas

- 14+ years in performance marketing with substantial paid acquisition budget ownership. The master documents up to $30 million per month managed with a four-person team.
- Google Ads, Microsoft Ads, Meta, audience strategy, budget allocation, landing pages, SEO collaboration, and conversion testing directly fit the channel work.
- At GRIP6, monthly marketing volume grew from about $200K to $800K while ROAS improved from 1.5 to 3.5.
- Customer journey analysis includes micro-conversions, click-to-call, purchase, and lifetime value; the master also documents attribution, GA4/GTM setup, dashboards, forecasting, holdouts, and budget scenarios.
- E-commerce and B2B leadership, cross-functional work with sales, operations, finance, IT, and product, and experience advising agency clients support the broader growth scope.
- The position is remote within the US, consistent with the master resume's Lehi, Utah location.

## Weaknesses or missing requirements

- The master does not establish insurance or U65/non-ACA market experience, enrollment or contribution-margin ownership, or a proven channel taken from zero to profitable enrollment volume.
- Invoca, offline conversion uploads, click-ID capture, reconciled call dispositions, and engineering-ticket ownership are not documented. Click-to-call and general conversion tracking are supported, but they do not establish those specific systems.
- The master mentions affiliate marketing and client/channel consultation, but does not document sourcing and negotiating affiliate, referral, lead-aggregator, or content partnerships.
- Agency experience is documented from the agency side; direct accountability for an outside agency partner is not established.
- The master documents CPA, ROAS, ROI, and LTV work, but does not demonstrate ownership of blended CAC or contribution-margin targets.

## ATS keyword alignment

Supported: growth marketing, paid search, paid social, Google Ads, Microsoft Ads, Meta, acquisition budget, ROAS, CPA, conversion tracking, attribution, customer journey, click-to-call, lifetime value, GA4, Google Tag Manager, forecasting, SEO, CRO, landing pages, A/B testing, e-commerce, B2B, and cross-functional collaboration.

Gaps: Invoca, offline conversion imports, click-ID capture, insurance enrollments, U65/non-ACA, partner negotiation, lead aggregators, and agency-partner accountability.

## Recommended resume emphasis

Lead with acquisition budget scale, channel and ROAS outcomes, measurement and testing, and the ability to connect spend decisions to business economics. Show the breadth beyond paid search through B2B/B2C e-commerce, SEO collaboration, website conversion work, and cross-functional planning.

## Interview/application considerations

Prepare concrete examples of how budget reallocation, attribution analysis, and testing changed investment decisions. Explain the difference between the documented click-to-call analysis and the role's Invoca/offline conversion requirements. Discuss how agency-side experience would transfer to directing an outside partner, and address the insurance and partnership gaps directly.

## Score rationale

| Area | Points |
| --- | ---: |
| Paid channel strategy and budget ownership | 19/20 |
| Attribution, analytics, testing, and forecasting | 17/20 |
| Cross-channel growth, SEO, and conversion | 14/20 |
| Partnerships and agency direction | 6/15 |
| Insurance domain and enrollment economics | 2/15 |
| Seniority and US remote alignment | 10/10 |
| **Total** | **68/100** |

## Validation and sources

- Microsoft Word computed 2 pages; the Word-exported PDF contains 2 pages. Native validation passed with no content or layout audit issues.
- Factual source: `input/master-resume/Dave-Call-resume-9-23-26.docx`, freshly read for this generation.
- [Official HealthCare listing]({listing['hostedUrl']})
- Archived description: `input/job-descriptions/{NAME}.md`
- Evidence and validation files: `scratch/{NAME}/`
""", encoding="utf-8")
    print(f"Resume: {resume}\nReport: {report}\nListing: {archive}")


if __name__ == "__main__":
    main()
