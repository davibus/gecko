"""Regression tests for the evidence-first resume workflow."""

import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import gecko_v2 as v2


class GeckoV2Tests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.listing = v2.ROOT / "input/job-descriptions/Vitality-Medical+b1e13fda0b6ba9cb.md"
        cls.plan = v2.create_plan(cls.listing)

    def test_plan_precedes_generation_and_classifies_requirements(self):
        plan = self.plan
        self.assertEqual(plan["version"], 3)
        self.assertTrue(plan["requirements"]["required_skills"])
        self.assertTrue(plan["requirements"]["preferred_skills"])
        self.assertTrue(plan["requirements"]["responsibilities"])
        self.assertTrue(plan["requirements"]["tools"])
        self.assertTrue(plan["requirements"]["seniority_signals"])
        self.assertTrue(plan["requirements"]["industry_terminology"])
        self.assertTrue(plan["requirements"]["ats_keywords"])
        self.assertTrue(any(r["status"] == "gap" for r in plan["requirements"]["responsibilities"]))
        v2.verify_plan(plan)

    def test_resume_filename_includes_safe_title_and_keeps_number_last(self):
        self.assertEqual(
            v2.resume_filename("Acme / West", "Director, SEO / Growth: US?", "abc123"),
            "Dave-Call+Acme-West+Director,-SEO-Growth-US+abc123.docx",
        )
        self.assertTrue(v2.resume_filename("Acme", "Very Long Title " * 20, "abc123").endswith("+abc123.docx"))

    def test_manual_generator_checks_live_x_before_writing(self):
        import sys
        sys.path.insert(0, str(v2.ROOT / "job-scout"))
        with tempfile.TemporaryDirectory() as temp:
            listing = Path(temp) / "Acme+abc123.md"
            listing.write_text("- **Scout ID:** 123\n- **URL:** https://example.test/job\n", encoding="utf-8")
            tab = SimpleNamespace(
                headers={"Resume Created": 7},
                rows=[(2, {"Scout ID": "123", "Resume Created": " x ",
                           "Job URL": "https://example.test/job"})],
            )
            with patch("google_tracker.GoogleTracker") as tracker:
                tracker.return_value.scout.return_value = tab
                with self.assertRaisesRegex(RuntimeError, "SKIP RESUME"):
                    v2._assert_no_existing_tracker_resume({"listing": str(listing)})
                tab.rows[0][1]["Resume Created"] = ""
                v2._assert_no_existing_tracker_resume({"listing": str(listing)})

    def test_changed_evidence_and_unsupported_skill_are_rejected(self):
        tampered = json.loads(json.dumps(self.plan))
        tampered["evidence"][0]["quote"] += " Invented result."
        with self.assertRaisesRegex(ValueError, "Unsupported or altered evidence"):
            v2.verify_plan(tampered)
        tampered = json.loads(json.dumps(self.plan))
        tampered["resume"]["tools_platforms"].append("Kubernetes")
        with self.assertRaisesRegex(ValueError, "Tools & Platforms"):
            v2.verify_plan(tampered)

    def test_generated_docx_keeps_source_claims_and_layout(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "resume.docx"
            v2.make_resume(self.plan, path)
            self.assertEqual(v2.inspect_docx(self.plan, path), [])
            from docx import Document
            doc = Document(path)
            bullet = next(p for p in doc.paragraphs if p.style.name == "List Bullet")
            bullet.text += " Unsupported metric $999M."
            doc.save(path)
            self.assertIn("Resume bullets differ from the source-backed tailoring plan.", v2.inspect_docx(self.plan, path))

    def test_official_tool_names_do_not_trigger_keyword_stuffing(self):
        listing = v2.ROOT / "input/job-descriptions/Pattern+3245845963235237630.md"
        plan = v2.create_plan(listing)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "resume.docx"
            v2.make_resume(plan, path)
            issues = v2.inspect_docx(plan, path)
        self.assertFalse(any("keyword stuffing: ads" in issue.casefold() for issue in issues))

    def test_native_word_failure_blocks_pass(self):
        with tempfile.TemporaryDirectory() as tmp:
            docx = Path(tmp) / "resume.docx"
            v2.make_resume(self.plan, docx)
            with patch.object(v2.subprocess, "run") as run:
                run.return_value.returncode = 1
                run.return_value.stderr = "Word bridge unavailable"
                run.return_value.stdout = ""
                report = v2.native_qa(self.plan, docx, Path(tmp))
            self.assertEqual(run.call_args.kwargs["creationflags"], v2.windows_creationflags())
            self.assertEqual(report["status"], "fail")
            self.assertTrue(any("Native Word pagination" in issue for issue in report["issues"]))

    def test_relevant_metric_cannot_disappear_silently(self):
        metric = self.plan["resume"]["relevant_metric_ids"][0]
        altered = json.loads(json.dumps(self.plan))
        altered["resume"]["selected_evidence_ids"].remove(metric)
        if metric in altered["resume"]["selected_results_ids"]:
            altered["resume"]["selected_results_ids"].remove(metric)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "resume.docx"
            v2.make_resume(altered, path)
            self.assertTrue(any("Relevant quantified accomplishment omitted" in issue
                                for issue in v2.inspect_docx(altered, path)))

    def test_reads_updated_master_archive_instead_of_cached_content(self):
        with tempfile.TemporaryDirectory() as tmp:
            master = Path(tmp) / "Dave-Call-Resume.txt"
            master.write_text("First master version\n", encoding="utf-8")
            with patch.object(v2, "ROOT", Path(tmp)), patch.object(v2, "MASTER", master):
                first = v2.source_text()
                first_hash = v2.source_hashes()
                master.write_text("First master version\nNew approved accomplishment\n", encoding="utf-8")
                self.assertIn("New approved accomplishment", v2.source_text())
                self.assertNotEqual(first, v2.source_text())
                self.assertNotEqual(first_hash, v2.source_hashes())

    def test_old_files_and_format_reference_are_not_evidence(self):
        self.assertEqual(v2.MASTER.name, "Dave-Call-Resume.txt")
        self.assertEqual(v2.FORMAT_MODEL.name, "MODEL-GECKO-PRODUCT_Dave_Call_Resume.pdf")
        self.assertEqual({Path(path).as_posix() for path in self.plan["source_sha256"]},
                         {"input/master-resume/Dave-Call-Resume.txt"})
        self.assertEqual(self.plan["format_model"],
                         "input/master-resume/MODEL-GECKO-PRODUCT_Dave_Call_Resume.pdf")
        self.assertEqual(self.plan["format_model_sha256"], v2.formatting_model_hash())
        self.assertTrue(all(e["source"].endswith("Dave-Call-Resume.txt")
                            for e in self.plan["evidence"]))
        self.assertNotIn("Spanish", self.plan["resume"]["contact"])
        self.assertNotIn("Spanish", self.plan["resume"]["certifications"])

    def test_missing_required_authorities_fail_without_fallback(self):
        with tempfile.TemporaryDirectory() as tmp:
            missing_master = Path(tmp) / "Dave-Call-Resume.txt"
            missing_model = Path(tmp) / "MODEL-GECKO-PRODUCT_Dave_Call_Resume.pdf"
            with patch.object(v2, "MASTER", missing_master), patch.object(v2, "FORMAT_MODEL", v2.FORMAT_MODEL):
                with self.assertRaisesRegex(FileNotFoundError, "Required Gecko master resume could not be found"):
                    v2.create_plan(self.listing)
            with patch.object(v2, "MASTER", v2.MASTER), patch.object(v2, "FORMAT_MODEL", missing_model):
                with self.assertRaisesRegex(FileNotFoundError, "Required Gecko formatting model could not be found"):
                    v2.make_resume(self.plan, Path(tmp) / "resume.docx")

    def test_generated_docx_uses_required_model_typography_and_hierarchy(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "resume.docx"
            v2.make_resume(self.plan, path)
            from docx import Document
            doc = Document(path)
            self.assertEqual(doc.styles["Normal"].font.name, "EB Garamond")
            self.assertAlmostEqual(doc.styles["Normal"].font.size.pt, 11)
            headings = [p.text for p in doc.paragraphs if p.text in [s.upper() for s in v2.SECTIONS]]
            self.assertEqual(headings, [s.upper() for s in v2.SECTIONS])

    def test_mandatory_strengths_tools_and_single_line_job_headers(self):
        self.assertGreaterEqual(len(self.plan["resume"]["core_strengths"]), 8)
        self.assertGreaterEqual(len(self.plan["resume"]["tools_platforms"]), 8)
        self.assertEqual(
            v2.formatting_content_issues(
                self.plan["resume"]["core_strengths"], self.plan["resume"]["tools_platforms"]),
            [],
        )
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "resume.docx"
            v2.make_resume(self.plan, path)
            from docx import Document
            doc = Document(path)
            self.assertGreaterEqual(len([c for r in doc.tables[0].rows for c in r.cells if c.text.strip()]), 8)
            evidence = {item["id"]: item for item in self.plan["evidence"]}
            selected_jobs = {evidence[eid]["job"] for eid in self.plan["resume"]["selected_evidence_ids"]}
            jobs = [job for index, job in enumerate(self.plan["resume"]["jobs"]) if index in selected_jobs]
            self.assertEqual(len(doc.tables) - 1, len(jobs))
            for table, job in zip(doc.tables[1:], jobs):
                self.assertEqual(len(table.rows), 1)
                self.assertEqual(table.cell(0, 0).text, v2.employment_heading(job))
                self.assertEqual(len(table.cell(0, 0).paragraphs), 1)
                self.assertEqual(len(table.cell(0, 0).paragraphs[0].runs), 1)
                self.assertTrue(table.cell(0, 0).paragraphs[0].runs[0].bold)
                self.assertFalse(table.cell(0, 0).paragraphs[0]._p.xpath(".//w:br | .//w:tab"))
                self.assertEqual(table.cell(0, 1).text, v2.employment_dates(job))
                self.assertEqual(table.cell(0, 1).paragraphs[0].alignment, 2)
                self.assertAlmostEqual(table.cell(0, 0).width.inches, 6.2, places=1)
                self.assertAlmostEqual(table.cell(0, 1).width.inches, 1.3, places=1)

    def test_heading_formatter_preserves_punctuation_and_rejects_missing_fields(self):
        cases = (
            ("Manager", "Acme", "Provo, UT"),
            ("Director of E-Commerce / Marketing", "Intercon Inc.", "Salt Lake City, UT"),
            ("VP, Growth & Analytics", "Smith-Jones, LLC", "West Valley City, UT"),
            ("Director, Marketing Strategy & Development", "Bowen Law Professional Group", "Salt Lake City, UT"),
        )
        for title, company, location in cases:
            self.assertEqual(v2.employment_heading({"title": title, "company": company, "location": location}),
                             f"{title} - {company} - {location}")
        for missing in ("title", "company", "location"):
            job = {"title": "Manager", "company": "Acme", "location": "Provo, UT"}
            job[missing] = ""
            with self.assertRaises(ValueError):
                v2.employment_heading(job)
        with self.assertRaises(ValueError):
            v2.employment_heading({"title": "Manager\nGrowth", "company": "Acme", "location": "Provo, UT"})
        with self.assertRaises(ValueError):
            v2.employment_heading({"title": "Manager", "company": "Acme", "location": "Utah"})

    def test_employment_dates_preserve_verified_months_and_present(self):
        self.assertEqual(v2.employment_dates({"dates": "November 2022 - June 2023 (Contract)"}),
                         "Nov 2022–Jun 2023")
        self.assertEqual(v2.employment_dates({"dates": "March 2021 - Present"}),
                         "Mar 2021–Present")
        for source in ("", "2022-2025", "November 2022 - Tomorrow"):
            with self.assertRaises(ValueError):
                v2.employment_dates({"dates": source})

    def test_generation_auto_corrects_section_counts_and_capitalization(self):
        altered = json.loads(json.dumps(self.plan))
        altered["resume"]["core_strengths"] = ["marketing analytics"]
        altered["resume"]["tools_platforms"] = ["google ads"]
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "resume.docx"
            v2.make_resume(altered, path)
            self.assertGreaterEqual(len(altered["resume"]["core_strengths"]), 8)
            self.assertGreaterEqual(len(altered["resume"]["tools_platforms"]), 8)
            self.assertEqual(v2.inspect_docx(altered, path), [])

    def test_qa_rejects_missing_dates_misalignment_and_split_heading(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "resume.docx"
            v2.make_resume(self.plan, path)
            from docx import Document
            doc = Document(path)
            table = doc.tables[1]
            table.cell(0, 1).text = ""
            doc.save(path)
            issues = v2.inspect_docx(self.plan, path)
            self.assertTrue(any("Employment dates differ" in issue for issue in issues))
            table.cell(0, 1).text = v2.employment_dates(self.plan["resume"]["jobs"][0])
            table.cell(0, 1).paragraphs[0].alignment = 0
            doc.save(path)
            self.assertTrue(any("Employment dates are not right-aligned" in issue for issue in v2.inspect_docx(self.plan, path)))
            table.cell(0, 0).paragraphs[0].add_run(" extra")
            doc.save(path)
            self.assertTrue(any("split across paragraphs or runs" in issue for issue in v2.inspect_docx(self.plan, path)))

    def test_current_master_listed_marketo_is_supported(self):
        listing = "## Required Qualifications\n- Must have hands-on Marketo experience in marketing operations."
        result = v2.requirements(listing, v2.source_text(), self.plan["evidence"])
        self.assertEqual(result["tools"][0]["text"], "Marketo")
        self.assertEqual(result["tools"][0]["status"], "supported")

    def test_unlisted_requirement_remains_a_gap(self):
        listing = "## Required Qualifications\n- Must have hands-on Core Web Vitals experience."
        result = v2.requirements(listing, v2.source_text(), self.plan["evidence"])
        self.assertEqual(result["required_skills"][0]["status"], "gap")

    def test_flattened_aggregator_prose_still_yields_qualitative_requirements(self):
        listing = (
            "## Job description\n"
            "Company background and product information. "
            "Professional Experience/Background to be successful in this role: "
            "Ability to manage paid search campaigns and marketing analytics "
            "Develop reporting and forecasting for cross-functional teams "
            "Minimum 5 years of digital marketing experience "
            "Expected Outcomes in 3, 6, or 12 months: "
            "Work with stakeholders to improve customer acquisition and retention "
            "Represent marketing performance to executive leadership"
        )
        result = v2.requirements(listing, v2.source_text(), self.plan["evidence"])
        core = result["required_skills"] + result["responsibilities"]
        self.assertGreaterEqual(len(core), 4)
        self.assertTrue(any(item["status"] in {"supported", "partial", "gap"} for item in core))

if __name__ == "__main__":
    unittest.main()
