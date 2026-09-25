"""Generate the reviewed Apply? selections with direct master-DOCX evidence."""

import argparse
import json
from pathlib import Path

import docx
from docx.shared import Inches, RGBColor

from generate_wpromote_resume import build_resume as build_layout
from generate_wpromote_resume import set_labeled, set_plain
from generate_vitality_medical_resume import NAVY
from gecko_v2 import (source_text, source_hashes, extract_evidence, master_identity,
                      master_jobs, listing_metadata, resume_filename)

ROOT = Path(__file__).resolve().parents[1]
CONFIG = {
    45: {
        "company": "Capital-One", "key": "5887543800",
        "headline": "AGENCY CLIENT STRATEGY | PAID MEDIA & PERFORMANCE ANALYTICS",
        "summary": "Digital marketing leader with 14+ years in paid media, agency account strategy, and client-facing performance analysis. Managed 75+ Google Ads accounts and more than $276K in combined monthly spend while advising clients on business goals, channel allocation, campaign structure, and growth priorities. Presented data-backed recommendations, managed teams and large budgets, and coordinated with business stakeholders to improve acquisition outcomes.",
        "skills": [
            ["Agency Client Work: ", "Client consultation, business-goal discovery, account management, performance presentations, and cross-functional coordination."],
            ["Paid Media & Platforms: ", "Google Ads, Microsoft Ads, Meta, Amazon, Search, Shopping, Display, YouTube, remarketing, and campaign optimization."],
            ["Analytics & Economics: ", "Tableau, Looker Studio, GA4, attribution, conversion analysis, ROI, ROAS, CPA, LTV, and forecasting."],
            ["Planning & Growth: ", "Budget allocation, audience and competitive research, campaign testing, investment scenarios, and client recommendations."],
            ["Leadership & Reporting: ", "Four-person paid media team management, eight-person ecommerce team leadership, automated reporting, and KPI presentations."],
        ],
        "order": [[2,3,1,0],[1,2,3,0],[0,3,2,1],[1,3,2,0],[0,3,2,1]],
    },
    60: {
        "company": "Bowman", "key": "5879577416",
        "headline": "MARKETING LEADERSHIP | CROSS-FUNCTIONAL STRATEGY & DELIVERY",
        "summary": "Digital marketing leader with 14+ years directing ecommerce, client strategy, campaign performance, and cross-functional initiatives. Led an eight-person ecommerce team, coordinated priorities with sales, operations, finance, IT, and product partners, and used project-level cost-benefit analysis to guide growth investments. Experienced presenting results to clients and senior leaders and building repeatable reporting and account-management processes.",
        "skills": [
            ["Leadership & Collaboration: ", "Eight-person team leadership, agency client consultation, sales and operations coordination, and senior leadership presentations."],
            ["Marketing Strategy: ", "Business-goal discovery, audience and competitive research, messaging recommendations, campaign planning, and growth prioritization."],
            ["Processes & Quality: ", "Account quality assurance, repeatable agency processes, project prioritization, cross-functional reporting, and cost-benefit analysis."],
            ["Measurement: ", "GA4, Google Tag Manager, Tableau, Looker Studio, Funnel.io, attribution, conversion analysis, and KPI reporting."],
            ["Digital Execution: ", "Paid search, social, SEO collaboration, ecommerce, lead generation, A/B testing, and automated reporting."],
        ],
        "order": [[2,3,1,0],[3,2,1,0],[0,3,2,1],[1,3,2,0],[0,3,2,1]],
    },
    51: {
        "company": "Acxiom", "key": "5839615490",
        "headline": "AGENCY CLIENT STRATEGY | DIGITAL MARKETING & ANALYTICS",
        "summary": "Client-facing digital marketing leader with 14+ years across agency account strategy, paid media, ecommerce, and analytics. Managed 75+ Google Ads accounts, consulted with clients on business goals and channel investment, and translated performance data into recommendations. Led teams, built reporting processes, and worked with sales, finance, operations, IT, and product stakeholders to prioritize growth opportunities.",
        "skills": [
            ["Client Strategy: ", "Agency consultation, business-goal discovery, account management, stakeholder presentations, and cross-channel recommendations."],
            ["Growth & Economics: ", "Budget allocation, customer acquisition, ROI, ROAS, CPA, LTV, forecasting, and cost-benefit analysis."],
            ["Analytics & Reporting: ", "Tableau, Looker Studio, Funnel.io, GA4, SQL, Python, Databricks, attribution, and KPI reporting."],
            ["Cross-Functional Delivery: ", "Sales, finance, operations, IT, product, website, and creative coordination; repeatable agency processes."],
            ["Paid Media Scope: ", "Google Ads, Microsoft Ads, Meta, Amazon, Search, Shopping, Display, YouTube, lead generation, and campaign optimization."],
        ],
        "order": [[2,3,1,0],[1,2,3,0],[0,3,2,1],[1,3,2,0],[0,3,2,1]],
    },
    703: {
        "company": "Lumin-Digital", "key": "d96affe1-9ce7-4c3b-b40e-223e6c82a537",
        "headline": "B2B DIGITAL MARKETING | CUSTOMER INSIGHT & COMPETITIVE ANALYSIS",
        "summary": "Digital marketing leader with 14+ years across agency client strategy, B2B ecommerce, acquisition, and performance analysis. Conducted keyword and competitive research, developed audience and messaging recommendations, and advised clients on website, SEO, email, social, and paid-media opportunities. Directed marketing for a B2B distributor and coordinated with sales, product, finance, IT, and operations teams to evaluate growth investments.",
        "skills": [
            ["Customer & Market Insight: ", "Client consultation, audience analysis, customer journeys, keyword and competitive research, and messaging recommendations."],
            ["B2B & Cross-Functional Work: ", "B2B ecommerce, sales and product coordination, project prioritization, and senior leadership presentations."],
            ["Campaign & Conversion Strategy: ", "Paid search, social, SEO collaboration, email channel experience, lead generation, landing pages, and A/B testing."],
            ["Business Analysis: ", "Cost-benefit analysis, ROI, ROAS, CPA, annual projections, attribution, forecasting, and conversion reporting."],
            ["Analytics & Tools: ", "GA4, Google Tag Manager, Tableau, Looker Studio, Funnel.io, Excel, SQL, Python, and Databricks."],
        ],
        "order": [[1,3,2,0],[3,2,1,0],[0,3,2,1],[1,2,3,0],[3,0,2,1]],
    },
    14: {
        "company": "Acxiom", "key": "5879612472",
        "headline": "DIGITAL MARKETING | CAMPAIGN ANALYTICS & OPERATIONS",
        "summary": "Digital marketing professional with 14+ years coordinating campaigns, measuring performance, and improving reporting across agency and ecommerce settings. Configured GA4 and Google Tag Manager measurement, built Tableau and Looker Studio dashboards, and automated weekly reporting with Python, SQL, Databricks, Excel, and Funnel.io. Worked with sales, product, operations, finance, and IT teams to turn data into practical recommendations.",
        "skills": [
            ["Campaign Support: ", "Paid search, social, ecommerce, lead generation, audience analysis, website recommendations, and channel coordination."],
            ["Reporting & Dashboards: ", "GA4, Google Tag Manager, Tableau, Looker Studio, Funnel.io, KPI reporting, attribution, and conversion analysis."],
            ["Data & Automation: ", "Excel, Python, SQL, Databricks, marketing databases, automated reporting, and monitoring scripts."],
            ["Testing & Improvement: ", "A/B and multivariate testing, landing pages, cost-benefit analysis, forecasting, and budget scenarios."],
            ["Team Collaboration: ", "Marketing, sales, product, finance, operations, and IT coordination; client communication and process improvement."],
        ],
        "order": [[2,1,3,0],[2,3,1,0],[0,2,3,1],[3,2,1,0],[3,0,2,1]],
    },
    23: {
        "company": "Trellix", "key": "5849329600",
        "headline": "B2B DIGITAL MARKETING | CHANNEL PERFORMANCE & ANALYTICS",
        "summary": "Digital marketing leader with 14+ years across paid media, agency client strategy, B2B ecommerce, and cross-channel reporting. Coordinated marketing, sales, operations, finance, IT, and product stakeholders for a distributor serving more than 650 retail partners. Managed large media budgets, built automated reporting, and translated attribution, conversion, and ROI analysis into clear investment recommendations.",
        "skills": [
            ["B2B & Channel Collaboration: ", "B2B retail relationships, sales coordination, agency client strategy, cross-functional planning, and lead generation."],
            ["Campaign & Budget Management: ", "Paid search, social, Google Ads, Microsoft Ads, Meta, Amazon, budget allocation, pacing, and campaign optimization."],
            ["Measurement & Reporting: ", "GA4, Google Tag Manager, Tableau, Looker Studio, Funnel.io, attribution, conversion analysis, CPA, ROI, and KPI reporting."],
            ["Planning & Testing: ", "Audience and keyword research, competitive analysis, A/B testing, regression models, forecasting, and cost-benefit analysis."],
            ["Leadership & Delivery: ", "Four-person paid media team management, eight-person ecommerce team leadership, client presentations, and senior leadership recommendations."],
        ],
        "order": [[2,3,1,0],[3,2,1,0],[0,2,3,1],[1,3,2,0],[0,3,2,1]],
    },
    721: {
        "company": "Decile-Group", "key": "b8d5a786-5a7e-4a89-8108-5ec50bde48e8",
        "headline": "GROWTH MARKETING | ACQUISITION, TESTING & ANALYTICS",
        "summary": "Digital marketing professional with 14+ years growing acquisition across agency, ecommerce, and multi-channel programs. Advised clients on paid media, SEO, email, social, and website opportunities; developed audience and messaging recommendations; and measured conversion and revenue outcomes. Led an eight-person ecommerce team and used dashboards, forecasts, and structured testing to guide growth decisions.",
        "skills": [
            ["Growth & Channels: ", "Paid search, social, SEO collaboration, email channel experience, lead generation, ecommerce, and audience strategy."],
            ["Testing & Conversion: ", "A/B and multivariate testing, landing pages, attribution, customer journeys, conversion analysis, and campaign optimization."],
            ["Analytics & Reporting: ", "GA4, Google Tag Manager, Tableau, Looker Studio, Funnel.io, KPI dashboards, ROAS, ROI, and forecasting."],
            ["Client & Leadership Work: ", "Agency client consultation, business-goal discovery, senior leadership presentations, cross-functional coordination, and team leadership."],
            ["Technical Tools: ", "Python, SQL, Databricks, Excel, automated reporting, monitoring scripts, and marketing databases."],
        ],
        "order": [[1,2,3,0],[1,2,3,0],[0,2,3,1],[1,2,3,0],[0,3,2,1]],
    },
    714: {
        "company": "RainFocus", "key": "3aaf4676-cec2-4e54-a77e-02075e283ba7",
        "headline": "DIGITAL MARKETING STRATEGY | BUSINESS ANALYSIS & CROSS-FUNCTIONAL DELIVERY",
        "summary": "Digital marketing and ecommerce leader with 14+ years connecting customer, campaign, and financial data to business decisions. Directed digital marketing for a B2B distributor, coordinating with sales, finance, IT, operations, and product teams and using cost-benefit analysis to prioritize investments. Experienced in client research, messaging recommendations, KPI presentations, and leading an eight-person ecommerce team.",
        "skills": [
            ["Market & Customer Insight: ", "Client consultation, customer journeys, audience analysis, keyword and competitive research, and messaging recommendations."],
            ["Business Cases & Economics: ", "Cost-benefit analysis, annual projections, ROI, ROAS, CPA, forecasting, regression models, and budget scenarios."],
            ["Cross-Functional Leadership: ", "Sales, finance, IT, operations, product, website, and creative coordination; team management and leadership presentations."],
            ["Marketing Analytics: ", "GA4, Google Tag Manager, Tableau, Looker Studio, Funnel.io, SQL, Python, Databricks, and KPI reporting."],
            ["Digital Execution: ", "Paid search, social, ecommerce, lead generation, conversion testing, landing pages, and repeatable agency processes."],
        ],
        "order": [[1,3,2,0],[3,2,1,0],[3,0,2,1],[1,2,3,0],[3,0,2,1]],
    },
    471: {
        "company": "Lemon.io", "key": "2091129",
        "headline": "MARKETING ANALYTICS | PYTHON, DATABRICKS & MODELING",
        "summary": "Digital marketing analyst and leader with 14+ years turning campaign and customer data into practical decisions. Built Tableau dashboards and automated weekly reporting across more than 10 channels using Python, Databricks, SQL, Excel, and Funnel.io. Applied regression models, holdouts, geo tests, attribution, and forecasting to evaluate investment options, and communicated findings to clients and senior leaders.",
        "skills": [
            ["Data & Programming: ", "Python, Databricks, SQL, Excel, marketing databases, automated reporting, and monitoring scripts."],
            ["Analytics & Visualization: ", "Tableau, Looker Studio, GA4, Funnel.io, dashboards, KPI reporting, and conversion analysis."],
            ["Modeling & Experiments: ", "Regression models, holdouts, geo tests, forecasting, A/B and multivariate testing, and budget scenarios."],
            ["Business Communication: ", "Client consultation, senior leadership presentations, cross-functional reporting, and investment recommendations."],
            ["Marketing Measurement: ", "Attribution, customer journeys, LTV, CPA, ROI, ROAS, campaign monitoring, and channel performance."],
        ],
        "order": [[2,3,1,0],[2,1,3,0],[2,3,0,1],[3,2,1,0],[3,2,0,1]],
    },
    36: {
        "company": "Marrina-Decisions", "key": "5847768748",
        "headline": "MARKETING ANALYTICS | REPORTING AUTOMATION & CLIENT STRATEGY",
        "summary": "Digital marketing and analytics professional with 14+ years in acquisition, client strategy, and campaign measurement. Built Tableau and Looker Studio dashboards, automated weekly reporting across more than 10 channels with Python, SQL, Databricks, Excel, and Funnel.io, and maintained marketing databases. Experienced coordinating with sales, product, finance, operations, and IT and turning performance data into practical recommendations.",
        "skills": [
            ["Reporting & Dashboards: ", "Tableau, Looker Studio, Funnel.io, GA4, Google Tag Manager, KPI reporting, and automated weekly reporting."],
            ["Marketing Data: ", "Python, SQL, Databricks, Excel, marketing databases, conversion tracking, attribution, and data analysis."],
            ["Campaign Optimization: ", "Paid search, social, lead generation, audience analysis, A/B and multivariate testing, and budget allocation."],
            ["Client & Team Delivery: ", "Agency consultation, performance presentations, cross-functional coordination, and repeatable account-management processes."],
            ["Business Planning: ", "ROI, ROAS, CPA, forecasting, regression models, cost-benefit analysis, and senior leadership recommendations."],
        ],
        "order": [[2,3,1,0],[2,3,1,0],[2,0,3,1],[3,1,2,0],[3,0,2,1]],
    },
    476: {
        "company": "iMerit-Technology", "key": "2091126",
        "headline": "DATA ANALYSIS | QUALITY CONTROL & CLEAR COMMUNICATION",
        "summary": "Digital marketing and analytics professional with 14+ years evaluating campaign data, checking performance, and communicating recommendations. Built automated weekly reporting across more than 10 channels, developed a repeatable agency account-management system, and used testing and attribution to assess results. Experienced presenting findings to clients and senior leaders and documenting learnings across teams.",
        "skills": [
            ["Analytical Review: ", "Campaign performance analysis, attribution, conversion analysis, customer journeys, regression models, and forecasting."],
            ["Quality & Testing: ", "Account quality assurance, repeatable processes, A/B and multivariate testing, holdouts, and monitoring scripts."],
            ["Research & Evidence: ", "Keyword and competitive research, audience analysis, cost-benefit analysis, KPI reporting, and data-backed recommendations."],
            ["Communication: ", "Client presentations, senior leadership reporting, cross-functional collaboration, and documented learnings."],
            ["Tools: ", "Tableau, Looker Studio, Python, SQL, Databricks, Excel, Funnel.io, GA4, and Google Tag Manager."],
        ],
        "order": [[2,3,1,0],[2,3,1,0],[2,3,0,1],[3,2,1,0],[3,0,2,1]],
    },
    696: {
        "company": "Instrumentl", "key": "a2c6c454-00a1-4ab0-9c02-cddfffec17b5",
        "headline": "MARKETING ANALYTICS | REPORTING AUTOMATION & CROSS-CHANNEL SYSTEMS",
        "summary": "Digital marketing and analytics professional with 14+ years across acquisition, ecommerce, agency strategy, and performance reporting. Built Tableau and Looker Studio dashboards, automated weekly reporting across more than 10 channels with Python, SQL, Databricks, Excel, and Funnel.io, and maintained marketing databases and cross-channel systems. Experienced aligning marketing, sales, finance, operations, IT, and product stakeholders around metrics and investment decisions.",
        "skills": [
            ["Marketing Data & Reporting: ", "Tableau, Looker Studio, Funnel.io, SQL, Python, Databricks, Excel, marketing databases, dashboards, and automated reporting."],
            ["Campaign Measurement: ", "GA4, Google Tag Manager, attribution, conversion tracking, KPI reporting, ROI, ROAS, customer journeys, and channel performance."],
            ["Experimentation & Planning: ", "A/B and multivariate testing, holdouts, geo tests, regression models, forecasting, budget scenarios, and cost-benefit analysis."],
            ["Cross-Channel Execution: ", "Paid search, social, SEO collaboration, lead generation, ecommerce, audience strategy, and client consultation."],
            ["Stakeholder Delivery: ", "Senior leadership presentations, cross-functional coordination, project prioritization, and repeatable account-management processes."],
        ],
        "order": [[2,3,1,0],[2,3,1,0],[2,0,3,1],[3,1,2,0],[3,0,2,1]],
    },
    698: {
        "company": "Smart-Working-Solutions", "key": "0c605546-9908-4d2e-9f44-197273172f82",
        "headline": "SEARCH MARKETING | SEO COLLABORATION & ANALYTICS",
        "summary": "Search and digital marketing professional with 14+ years in paid search, ecommerce, agency consulting, and performance analysis. Conducted keyword and competitive research, advised clients on SEO and website opportunities, and coordinated channel strategy with web and creative teams. Configured GA4 and Google Tag Manager measurement, built Looker Studio and Tableau dashboards, and used conversion testing and attribution to guide growth decisions.",
        "skills": [
            ["Search & Website Strategy: ", "Keyword and competitive research, SEO collaboration, landing pages, website recommendations, and audience strategy."],
            ["Analytics & Tracking: ", "GA4, Google Tag Manager, Looker Studio, Tableau, Funnel.io, conversion tracking, attribution, and KPI reporting."],
            ["Testing & Optimization: ", "A/B and multivariate testing, conversion paths, ad copy, bidding experiments, ROAS, CPA, and customer journey analysis."],
            ["Campaign Channels: ", "Google Ads, Microsoft Ads, Search, Shopping, Display, YouTube, remarketing, social, email channel experience, and ecommerce."],
            ["Reporting & Collaboration: ", "Python, SQL, Databricks, Excel, automated reporting, client presentations, and cross-functional coordination."],
        ],
        "order": [[1,2,3,0],[2,3,1,0],[0,2,3,1],[2,1,3,0],[2,3,0,1]],
    },
    704: {
        "company": "TeamSnap", "key": "14124f03-3002-4ac5-96e2-dac3a39610a2",
        "headline": "MARKETING STRATEGY | CLIENT PARTNERSHIP & CAMPAIGN MEASUREMENT",
        "summary": "Digital marketing leader with 14+ years across agency client strategy, multi-channel acquisition, ecommerce, and performance measurement. Consulted with clients on business goals, audiences, budget allocation, SEO, website, email, social, and paid media opportunities. Led an eight-person ecommerce team, coordinated with sales and other business functions, and built reporting that translates campaign data into clear recommendations.",
        "skills": [
            ["Integrated Campaigns: ", "Paid search, social, email channel experience, SEO collaboration, lead generation, audience strategy, and website planning."],
            ["Client & Sales Partnership: ", "Agency consultation, business-goal discovery, client presentations, cross-functional coordination, and team leadership."],
            ["Measurement & Insights: ", "GA4, Google Tag Manager, Tableau, Looker Studio, Funnel.io, attribution, conversion analysis, KPI reporting, and ROI."],
            ["Testing & Messaging: ", "Keyword and competitive research, targeting and messaging recommendations, ad copy, landing pages, and A/B and multivariate testing."],
            ["Program Delivery: ", "Campaign monitoring, budget allocation, reporting automation, project prioritization, and repeatable account-management processes."],
        ],
        "order": [[1,2,3,0],[3,2,1,0],[0,2,3,1],[1,3,2,0],[0,3,2,1]],
    },
    724: {
        "company": "Prosper", "key": "be2de9d8-63db-445b-962d-79ee01b64548",
        "headline": "PERFORMANCE MARKETING LEADERSHIP | GROWTH & ANALYTICS",
        "summary": "Performance marketing leader with 14+ years managing acquisition, ecommerce, agency programs, and multi-channel investment. Managed up to $30 million per month in paid media with a four-person team and led an eight-person ecommerce group. Uses attribution, customer journey analysis, regression models, and forecasting to guide budget decisions, while presenting practical recommendations to clients and senior leaders.",
        "skills": [
            ["Growth & Acquisition: ", "Google Ads, Microsoft Ads, Meta, Amazon, paid search, social, lead generation, ecommerce, and customer acquisition."],
            ["Investment & Economics: ", "Budget allocation, ROI, ROAS, CPA, LTV, forecasting, regression models, holdouts, and cost-benefit analysis."],
            ["Customer & Market Insight: ", "Customer journeys, audience analysis, keyword and competitive research, attribution, conversion paths, and campaign testing."],
            ["Leadership & Communication: ", "Four-person paid media team management, eight-person ecommerce team leadership, client consultation, and senior leadership presentations."],
            ["Reporting & Optimization: ", "Tableau, Looker Studio, GA4, Google Tag Manager, SQL, Python, Databricks, automated reporting, and KPI planning."],
        ],
        "order": [[3,1,2,0],[1,2,3,0],[0,3,2,1],[1,2,3,0],[0,3,2,1]],
    },
    469: {
        "company": "Lemon.io", "key": "2091131",
        "headline": "MARKETING ANALYTICS | PYTHON, SQL & AUTOMATION",
        "summary": "Digital marketing and analytics professional with 14+ years using data and automation to improve campaign decisions. Built Tableau dashboards and automated reporting across more than 10 channels using Python, SQL, Databricks, Excel, and Funnel.io. Developed monitoring scripts, applied regression models and experimental analysis, and translated complex findings for clients and business leaders. Experience centers on marketing analytics and campaign systems.",
        "skills": [
            ["Programming & Data: ", "Python, SQL, Databricks, Excel, automated reporting, marketing databases, and monitoring scripts."],
            ["Analytics & Modeling: ", "Tableau, Looker Studio, regression models, forecasting, holdouts, geo tests, attribution, and conversion analysis."],
            ["Automation & Testing: ", "Weekly reporting automation, campaign monitoring, A/B and multivariate testing, and repeatable account-management processes."],
            ["Business Communication: ", "Client consultation, senior leadership presentations, cross-functional coordination, and practical investment recommendations."],
            ["Digital Systems: ", "GA4, Google Tag Manager, Funnel.io, ecommerce platforms, campaign measurement, KPI reporting, and channel data analysis."],
        ],
        "order": [[2,3,1,0],[2,3,1,0],[2,0,3,1],[3,2,1,0],[3,2,0,1]],
    },
    715: {
        "company": "Seer-Interactive", "key": "4d7a9ca8-e79b-4654-8a49-1b9976227bb7",
        "headline": "DIGITAL MARKETING | PROGRAM MEASUREMENT & CROSS-CHANNEL GROWTH",
        "summary": "Digital marketing professional with 14+ years across agency consulting, customer acquisition, ecommerce, and performance measurement. Managed multi-channel programs, advised clients on paid, organic, email, and website opportunities, and built repeatable reporting processes. Led an eight-person ecommerce team and partnered with sales, operations, finance, IT, and product stakeholders. Brings hands-on analytics, automation, and practical communication to program planning and performance decisions.",
        "skills": [
            ["Campaigns & Distribution: ", "Paid search, social, email channel experience, SEO collaboration, audience strategy, lead generation, landing pages, and cross-channel planning."],
            ["Measurement & Reporting: ", "GA4, Google Tag Manager, Tableau, Looker Studio, Funnel.io, attribution, conversion analysis, KPI dashboards, forecasting, and ROI."],
            ["Client & Team Partnership: ", "Agency client consultation, performance presentations, stakeholder coordination, project prioritization, and eight-person team leadership."],
            ["Testing & Optimization: ", "Keyword and competitive research, messaging recommendations, A/B and multivariate testing, conversion paths, budget allocation, and campaign monitoring."],
            ["Technical Execution: ", "Python, SQL, Databricks, Excel, reporting automation, marketing databases, and repeatable account-management processes."],
        ],
        "order": [[2,1,3,0],[2,3,1,0],[0,2,3,1],[1,3,2,0],[0,3,2,1]],
    },
    710: {
        "company": "Saviynt", "key": "3fb45888-3ba9-4b70-9988-61255c93409b",
        "headline": "MARKETING ANALYTICS | DATA VISUALIZATION & BUSINESS INSIGHTS",
        "summary": "Marketing analytics and digital marketing professional with 14+ years turning campaign, customer, and financial data into practical business decisions. Built Tableau and Looker Studio dashboards, automated reporting across more than 10 channels with Python, SQL, Databricks, Excel, and Funnel.io, and used attribution, regression models, and forecasts to guide investment. Experienced presenting KPIs and recommendations to clients and senior leaders and coordinating reporting across marketing, sales, finance, operations, and IT.",
        "skills": [
            ["Analytics & Visualization: ", "Tableau, Looker Studio, GA4, Google Tag Manager, Funnel.io, dashboards, KPI reporting, and senior leadership presentations."],
            ["Data & Automation: ", "SQL, Python, Databricks, Excel, marketing databases, automated reporting, monitoring scripts, and cross-channel data analysis."],
            ["Measurement & Economics: ", "Attribution, conversion tracking, customer journeys, ROI, ROAS, CPA, LTV, campaign measurement, financial reporting, and forecasting."],
            ["Modeling & Testing: ", "Regression models, holdouts, geo tests, A/B and multivariate testing, scenario analysis, budget allocation, and cost-benefit analysis."],
            ["Business Partnership: ", "Client presentations, senior leadership reporting, B2B ecommerce, cross-functional collaboration, and practical investment recommendations."],
        ],
        "order": [[2,3,1,0],[2,1,3,0],[2,0,3,1],[3,1,2,0],[3,2,0,1]],
    },
    708: {
        "company": "Jobscan", "key": "103299c7-8c4d-4ccd-b281-da31b5b0beb6",
        "headline": "GROWTH MARKETING | CONVERSION, ANALYTICS & ACQUISITION",
        "summary": "Digital marketing professional with 14+ years across customer acquisition, ecommerce, and multi-channel growth. Combines keyword research, ad copy, SEO collaboration, landing-page optimization, and conversion testing with GA4 and performance analysis. Led an eight-person ecommerce team, advised agency clients on website and channel strategy, and helped expand a B2B distributor into multiple B2C channels. Brings a hands-on approach to translating customer journey data into tests, priorities, and measurable business results.",
        "skills": [
            ["Conversion & Customer Journeys: ", "Landing-page optimization, conversion rate optimization, A/B and multivariate testing, attribution, micro-conversions, LTV, customer acquisition, and audience analysis."],
            ["Search & Messaging: ", "SEO, keyword and competitive research, ad copy, targeting and messaging recommendations, paid search, audience segmentation, and website strategy."],
            ["Analytics & Reporting: ", "GA4, Google Tag Manager, Looker Studio, Tableau, Funnel.io, Excel, KPI reporting, regression models, forecasting, and financial analysis."],
            ["Commerce & Collaboration: ", "B2B/B2C ecommerce, Shopify, Magento, email channel experience, client presentations, team leadership, cross-functional coordination, and documentation of learnings."],
            ["Execution & Automation: ", "Python, SQL, Databricks, automated reporting, monitoring scripts, campaign quality assurance, Google Ads, Microsoft Ads, Meta, and Amazon."],
        ],
        "order": [[1,2,3,0],[1,2,3,0],[1,0,2,3],[2,1,3,0],[2,0,3,1]],
    },
    37: {
        "company": "Bamboo-Insurance", "key": "5846794031",
        "headline": "DIRECTOR OF PERFORMANCE MARKETING | STRATEGY & ANALYTICS",
        "summary": "Performance marketing leader with 14+ years across paid search, ecommerce, and multi-channel acquisition. Managed up to $30 million per month in paid media with a four-person team. Combines hands-on Google, Microsoft Ads, and Meta execution with attribution, customer journey analysis, forecasting, and budget allocation. Experienced in ad copy, audience and bidding tests, reporting automation, and presenting investment recommendations, with a practical focus on acquisition efficiency and revenue growth.",
        "skills": [
            ["Paid Media & Acquisition: ", "Google Ads, Microsoft Ads, Meta, Amazon, Instagram, TikTok, Search, Shopping, Performance Max, Display, YouTube, remarketing, and SEO."],
            ["Testing & Creative: ", "Ad copy, keyword and competitive research, audience segmentation, targeting and messaging recommendations, bid experiments, landing pages, A/B testing, and multivariate testing."],
            ["Measurement & Economics: ", "GA4, Google Tag Manager, attribution, CPC, CPA, ROAS, LTV, conversion tracking, incremental testing, holdouts, regression models, forecasting, and budget scenarios."],
            ["Reporting & Automation: ", "Excel, Tableau, Looker Studio, Funnel.io, Python, SQL, Databricks, automated reporting, campaign monitoring, KPI development, and financial analysis."],
            ["Leadership & Delivery: ", "Budget allocation, senior leadership reporting, client consultation, cross-functional collaboration, team management, campaign builds, quality assurance, pacing, and daily optimization."],
        ],
        "order": [[3,0,1,2],[0,1,2,3],[3,2,0,1],[1,2,3,0],[2,3,0,1]],
    },
}

LABELS = [
    ["Bidding & Audience Analysis: ", "Customer Journey Analysis: ", "Reporting Automation: ", "Investment Planning: "],
    ["Channel Ownership: ", "Growth & Efficiency: ", "Measurement & Visibility: ", "Cross-Functional Execution: "],
    ["Digital Leadership: ", "B2C Expansion: ", "Data & Reporting: ", "Commercial Priorities: "],
    ["Agency Account Scale: ", "Client Strategy: ", "Research & Testing: ", "Reporting & Repeatable Processes: "],
    ["Team Leadership: ", "Campaign Execution: ", "Conversion Optimization: ", "Executive Planning: "],
]


def build_resume(scout_id: int):
    cfg = CONFIG[scout_id]
    scratch = ROOT / "scratch" / f"{cfg['company']}+{cfg['key']}"
    listing_path = ROOT / "input/job-descriptions" / f"{cfg['company']}+{cfg['key']}.md"
    title = cfg.get("title") or listing_metadata(listing_path.read_text(encoding="utf-8-sig"), listing_path)["title"]
    output = ROOT / "output/resumes" / resume_filename(cfg["company"], title, cfg["key"])
    # Re-read the authoritative DOCX for each job, including all tables.
    source = source_text()
    evidence = extract_evidence(source)
    identity, jobs = master_identity(), master_jobs()
    build_layout(output, scratch_dir=scratch)
    doc = docx.Document(output)
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(.45)
    p = doc.paragraphs
    set_plain(p[0], identity["name"], size=19, bold=True, color=NAVY)
    set_plain(p[1], cfg["headline"], size=11.5, bold=True, color=RGBColor(40,70,110))
    set_plain(p[2], identity["contact"], size=9.5)
    for index, title in [(3,"PROFESSIONAL SUMMARY"),(5,"CORE COMPETENCIES & TECHNICAL SKILLS"),(11,"PROFESSIONAL EXPERIENCE"),(34,"EDUCATION & CERTIFICATIONS")]:
        set_plain(p[index], title, bold=True, color=NAVY)
    set_plain(p[4], cfg["summary"])
    for index,(lead,body) in enumerate(cfg["skills"],6):
        set_labeled(p[index],lead,body,lead_color=NAVY)
    used=[]
    empty_bullet_paragraphs=[]
    for job_index,start in enumerate([12,16,21,25,29]):
        pool=[e for e in evidence if e["job"]==job_index]
        # The current master includes additional legacy text between the GRIP6
        # bullets and the next table-backed role. Keep only clean, attributable
        # bullets rather than forcing that text under the GRIP6 job header.
        pool=[e for e in pool if "Digital Marketing Manager" not in e["quote"]]
        if len(pool) < 3:
            raise ValueError(f"Master resume has too few clean bullets for job {job_index}")
        clean_limit = 3 if job_index == 1 else 4
        selected=[i for i in cfg["order"][job_index] if i < min(clean_limit, len(pool))]
        for offset,source_index in enumerate(selected):
            e=pool[source_index]
            set_labeled(p[start+offset],LABELS[job_index][source_index],e["quote"])
            used.append(e)
        for offset in range(len(selected),4):
            empty_bullet_paragraphs.append(p[start+offset])
    for paragraph in empty_bullet_paragraphs:
        paragraph._element.getparent().remove(paragraph._element)
    # Selected-results evidence explicitly supports this metric at LifeSpan.
    result=next(e for e in evidence if "$9 million in revenue" in e["quote"])
    set_labeled(p[33],"Revenue Impact: ",result["quote"])
    used.append(result)
    for table,job in zip(doc.tables,jobs,strict=True):
        set_plain(table.cell(0,0).paragraphs[0],f"{job['title']} | {job['company']} — {job['location']}",bold=True,color=NAVY)
        table.cell(0,1).text=""
    set_plain(p[35],identity["education"],bold=True)
    p[36]._element.getparent().remove(p[36]._element)
    set_plain(p[37],identity["certifications"])
    doc.save(output)
    (scratch/"generation-evidence.json").write_text(json.dumps({"scout_id":scout_id,"source_hashes":source_hashes(),"config":cfg,"selected_evidence":used},indent=2),encoding="utf-8")
    print(output)


if __name__=="__main__":
    parser=argparse.ArgumentParser()
    parser.add_argument("scout_id",type=int,choices=list(CONFIG))
    build_resume(parser.parse_args().scout_id)
