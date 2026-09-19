"""Build and validate the Gecko resume for NRTC job b629ab7dd39d6b83."""

from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.shared import RGBColor

from generate_marksmen_growth_resume import build_resume as build_base_resume
from generate_vitality_medical_resume import MUTED, NAVY, export_and_validate, set_run


ROOT = Path(__file__).resolve().parents[1]
JOB_KEY = "b629ab7dd39d6b83"
COMPANY = "NRTC"
OUTPUT = ROOT / "output" / "resumes" / f"Dave-Call+{COMPANY}+{JOB_KEY}.docx"
SCRATCH = ROOT / "scratch" / f"{COMPANY}+{JOB_KEY}"
TEMPLATE = SCRATCH / "base-template.docx"
PDF = SCRATCH / "resume-preview.pdf"


def clear_content(paragraph):
    """Remove runs while preserving paragraph properties and styling."""
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_plain(paragraph, text, *, size=11, bold=False, color=None):
    clear_content(paragraph)
    set_run(paragraph.add_run(text), size=size, bold=bold, color=color)


def set_labeled(paragraph, lead, body, *, lead_color=None):
    clear_content(paragraph)
    set_run(paragraph.add_run(lead), bold=True, color=lead_color)
    set_run(paragraph.add_run(body))


def build_resume(path: Path):
    SCRATCH.mkdir(parents=True, exist_ok=True)
    build_base_resume(TEMPLATE)
    doc = docx.Document(TEMPLATE)
    paragraphs = doc.paragraphs

    set_plain(
        paragraphs[1],
        "DIGITAL MARKETING SPECIALIST | PAID SEARCH & SOCIAL",
        size=11.5,
        bold=True,
        color=RGBColor(40, 70, 110),
    )
    set_plain(
        paragraphs[4],
        "Hands-on digital marketer with 14+ years of experience developing, executing, and optimizing paid search, "
        "paid social, organic content, and customer acquisition programs across B2B, B2C, and agency environments. "
        "Combines campaign strategy and creative development with conversion tracking, analytics, SEO, testing, and "
        "clear client communication. Experienced managing multiple accounts and substantial media budgets across "
        "Google, Meta, LinkedIn, TikTok, Pinterest, Snapchat, YouTube, and other platforms. Known for turning complex "
        "performance data into practical recommendations, persuasive content, and measurable business results.",
    )

    skills = {
        6: ("Paid Search & Social: ", "Google Ads, Bing Ads, Meta/Facebook/Instagram, LinkedIn, TikTok, Pinterest, Snapchat, YouTube, Twitter/X, Amazon, campaign research, setup, optimization, audience development, budgeting, and multi-account management."),
        7: ("Measurement & Optimization: ", "GA4, Google Tag Manager, conversion tracking, Adobe Analytics, Tableau, Looker Studio, Funnel.io, attribution, CRO, A/B and multivariate testing, LTV, SQL, Python, and Excel."),
        8: ("Content, Social & SEO: ", "Organic social strategy, content development and publication, copywriting, community management, Hootsuite, Buffer, SEO audits, on-page and mobile SEO, keyword research, Search Console, SEMrush, and Moz."),
        9: ("Client & Campaign Strategy: ", "B2B and B2C marketing, client consulting, relationship management, market and audience analysis, persuasive recommendations, cross-functional collaboration, project management, and executive reporting."),
        10: ("AI & Digital Tools: ", "ChatGPT, Claude, Perplexity, Codex, Cursor, AntiGravity, generative-AI research and workflow support, WordPress, Shopify, Magento, HTML5, CSS3, JavaScript, APIs, and marketing automation."),
    }
    for index, (lead, body) in skills.items():
        set_labeled(paragraphs[index], lead, body, lead_color=NAVY)

    bullets = {
        12: ("Portfolio & Team Leadership: ", "Managed $30 million per month with a team of four, using customer, channel, and financial analysis to guide investment and acquisition decisions."),
        13: ("Paid Media Optimization: ", "Analyzed paid-search strategy, bidding, audiences, brand and non-brand performance, promotions, and Performance Max within a large customer-acquisition portfolio."),
        14: ("Conversion Measurement: ", "Evaluated data-driven attribution, lifetime value, click-to-call activity, holdouts, and full-funnel reach-to-conversion tests to improve decisions."),
        15: ("Automated Reporting: ", "Used Python, SQL, Databricks, Excel, Funnel.io, and Tableau to automate accurate weekly reporting across more than 10 marketing platforms."),
        16: ("Cross-Platform Campaigns: ", "Managed Google, Bing, Amazon, Facebook, Instagram, and TikTok advertising across a $400K monthly budget."),
        17: ("Measurable Growth: ", "Increased monthly volume from $200K to $800K while improving ROAS from 1.5 to 3.5 through hands-on analysis and optimization."),
        18: ("Tracking & Analytics: ", "Implemented custom GA4 and Google Tag Manager reporting and built Looker Studio dashboards to connect marketing performance with operating priorities."),
        19: ("Integrated Digital Strategy: ", "Guided SEO, website optimization, email effectiveness, content, and cross-platform budget decisions in collaboration with creative, technical, and leadership teams."),
        21: ("B2B and B2C Growth: ", "Led ecommerce and marketing for a large B2B distributor, built a functional digital environment, and expanded into B2C platforms while supporting 650+ retail partners."),
        22: ("Audience & Market Development: ", "Used daily financial and market analysis to identify niches, guide product development, and assess new channels and customer opportunities."),
        23: ("Digital Expansion: ", "Expanded Amazon operations into CastleGate, Target.com, Costco.com, and other platforms while improving product data and digital merchandising."),
        24: ("Cross-Functional Communication: ", "Worked with IT, operations, finance, product development, HR, and sales leaders to implement technology, process, and operating improvements."),
        25: ("Multi-Client Strategy: ", "Advised clients on websites, SEO, email, paid search, display, mobile, remarketing, YouTube, and social based on business objectives, audiences, and budgets."),
        26: ("Hands-On Account Scale: ", "Built and managed 75+ Google Ads accounts totaling more than $276K per month while maintaining an agency-wide service-quality system."),
        27: ("Audits & Recommendations: ", "Analyzed campaign and site performance, presented strategic recommendations, and developed conversion and attribution models to improve client results."),
        28: ("Testing & Automation: ", "Designed multivariate tests and custom scripts for campaign monitoring, budget oversight, and faster identification of performance issues."),
        29: ("Team & Market Leadership: ", "Led an eight-person ecommerce team responsible for paid search, outreach, lead generation, and customer acquisition in the United States, Canada, and the UK."),
        30: ("Revenue Impact: ", "Managed 150+ campaigns and a $400K budget across Google, Bing, Gemini, and Amazon, producing more than $9M in revenue—35% of company revenue."),
        31: ("Campaign & Content Execution: ", "Performed keyword, audience, and competitive research; created search and display copy and creative; and managed email, shopping, video, social, local, and remarketing programs."),
        32: ("Conversion Optimization: ", "Analyzed funnels and conversion paths, optimized landing pages, and ran A/B and multivariate tests to improve acquisition performance."),
        33: ("Measurement & Reporting: ", "Established goals, projections, and KPIs with senior management and reported performance using Google Analytics, Magento, ACCTivate, and QuickBooks."),
    }
    for index, (lead, body) in bullets.items():
        set_labeled(paragraphs[index], lead, body)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


if __name__ == "__main__":
    build_resume(OUTPUT)
    export_and_validate(OUTPUT, PDF, SCRATCH)
