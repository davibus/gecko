"""Build and validate the Gecko resume for KoboToolbox job 2091141."""

from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.shared import RGBColor

from generate_marksmen_growth_resume import build_resume as build_base_resume
from generate_vitality_medical_resume import NAVY, export_and_validate, set_run


ROOT = Path(__file__).resolve().parents[1]
JOB_KEY = "2091141"
COMPANY = "KoboToolbox"
OUTPUT = ROOT / "output" / "resumes" / f"Dave-Call+{COMPANY}+{JOB_KEY}.docx"
SCRATCH = ROOT / "scratch" / f"{COMPANY}+{JOB_KEY}"
PDF = SCRATCH / "resume-preview.pdf"


def clear_content(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_plain(paragraph, text, *, size=11, bold=False, color=None):
    clear_content(paragraph)
    set_run(paragraph.add_run(text), size=size, bold=bold, color=color)


def set_labeled(paragraph, lead, body, *, lead_color=None):
    clear_content(paragraph)
    set_run(paragraph.add_run(lead), bold=True, color=lead_color)
    set_run(paragraph.add_run(body))


def build_resume(path: Path, scratch_dir: Path | None = None):
    scratch = Path(scratch_dir) if scratch_dir is not None else SCRATCH
    template = scratch / "base-template.docx"
    scratch.mkdir(parents=True, exist_ok=True)
    build_base_resume(template)
    doc = docx.Document(template)
    paragraphs = doc.paragraphs

    set_plain(
        paragraphs[1],
        "WEB TECHNOLOGY, DATA & AUTOMATION PROFESSIONAL",
        size=11.5,
        bold=True,
        color=RGBColor(40, 70, 110),
    )
    set_plain(
        paragraphs[4],
        "Technology and digital operations leader with 14+ years applying web technologies, scripting, data systems, "
        "automation, and analytics to real business problems. Proficient in HTML5, CSS3, JavaScript, Python, SQL, PHP, "
        "VBA, Git, XML, and API-based workflows, with experience supporting Magento, Shopify, WordPress, databases, "
        "reporting systems, and cross-functional web initiatives. Known for learning new tools quickly, investigating "
        "complex systems, documenting repeatable processes, and collaborating with technical and business teams to "
        "deliver reliable, measurable improvements.",
    )

    skills = {
        6: ("Web & Programming: ", "HTML5, CSS3, JavaScript, Python, SQL, PHP, VBA, XML, object-oriented programming, Git, GitHub, VS Code, APIs, relational databases, and web development support."),
        7: ("Platforms & Data Systems: ", "Magento, Shopify, WordPress, Salesforce, HubSpot, Google Ads API, ERP systems, marketing databases, GA4, Google Tag Manager, Databricks, Tableau, Looker Studio, and Adobe Analytics."),
        8: ("Automation & Troubleshooting: ", "Custom scripts, workflow automation, campaign monitoring, data integration, reporting pipelines, database consistency, process improvement, technical research, and issue investigation."),
        9: ("AI-Assisted Technical Work: ", "Codex, ChatGPT, Claude, Perplexity, Cursor, and AntiGravity for code-assisted workflows, research, analysis, scripting, documentation, troubleshooting, and quality review."),
        10: ("Collaboration & Delivery: ", "Cross-functional project management, requirements translation, stakeholder communication, distributed-team collaboration, prioritization, documentation, customer support, and continuous learning."),
    }
    for index, (lead, body) in skills.items():
        set_labeled(paragraphs[index], lead, body, lead_color=NAVY)

    bullets = {
        12: ("Automated Data Systems: ", "Used Python, SQL, Databricks, Excel, Funnel.io, and Tableau to build an automated weekly reporting workflow spanning more than 10 platforms."),
        13: ("Complex Analysis: ", "Investigated attribution, customer value, bidding, audience, promotion, and channel data and translated findings into testable recommendations and operating decisions."),
        14: ("Technical Evaluation: ", "Assessed Click-to-Call and iSpot systems and built dashboards and statistical models to improve visibility into customer journeys and program performance."),
        15: ("Team & Portfolio Leadership: ", "Managed $30 million per month with a team of four, coordinating analysis, testing, reporting, and investment decisions across a complex portfolio."),
        16: ("Web Measurement Implementation: ", "Configured custom GA4 reporting and Google Tag Manager implementations while partnering on website optimization and technical SEO."),
        17: ("Cross-Functional Dashboards: ", "Built Looker Studio reporting that connected marketing data with inventory, engineering, management, web development, email, and operating priorities."),
        18: ("Systems-Based Planning: ", "Created a robust spreadsheet methodology used across departments to connect inventory data, product availability, operating needs, and business decisions."),
        19: ("Digital Operations: ", "Managed interconnected advertising, ecommerce, analytics, and reporting platforms while improving monthly volume from $200K to $800K and ROAS from 1.5 to 3.5."),
        21: ("Digital Transformation: ", "Helped move a large traditional distributor into a functional digital environment by improving information architecture, ERP workflows, databases, collaboration, and ecommerce systems."),
        22: ("Database & Reporting Integrity: ", "Developed and maintained multiple company databases to improve consistency across departments and support accurate financial and marketing reporting."),
        23: ("Web & Commerce Platforms: ", "Added multiple B2C ecommerce platforms, supported Amazon Vendor and Seller Central operations, and expanded into Target.com, Costco.com, CastleGate, and other channels."),
        24: ("Technical Project Leadership: ", "Worked with IT, operations, finance, product development, HR, and sales leaders while managing changing priorities, systems implementations, reporting, and process improvement."),
        25: ("JavaScript Automation: ", "Developed custom scripts for continuous campaign monitoring, budget oversight, and faster identification of performance issues across client accounts."),
        26: ("Reusable Operating Systems: ", "Created a standardized quality system for 75+ accounts that was adopted across the agency to improve consistency and service delivery."),
        27: ("Data Integration & Visualization: ", "Consolidated data sources with Funnel.io and presented actionable findings through Looker Studio while using analytical platforms to investigate account behavior."),
        28: ("Testing & Models: ", "Created conversion and attribution models and multivariate tests to evaluate audiences, channels, messages, and landing-page behavior."),
        29: ("Web Team Collaboration: ", "Partnered with web developers on site performance, SEO/SEM, landing pages, conversion paths, and changes needed to support promotions and business goals."),
        30: ("Human-Assisted Automation: ", "Used JavaScript-based automation to adjust campaign behavior around promotions and business goals while retaining human oversight."),
        31: ("Technical Analytics: ", "Used Magento, Google Analytics, ACCTivate, QuickBooks, platform editors, and Amazon analytics to report KPIs, projections, conversion paths, and business performance."),
        32: ("Research & Optimization: ", "Conducted keyword, competitive, funnel, and acquisition analysis; developed experiments; and improved landing pages and conversion rates through structured testing."),
        33: ("Team Leadership: ", "Led an eight-person ecommerce department and coordinated paid search, outreach, customer acquisition, analytics, web collaboration, and operational planning across three countries."),
    }
    for index, (lead, body) in bullets.items():
        set_labeled(paragraphs[index], lead, body)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


if __name__ == "__main__":
    build_resume(OUTPUT)
    export_and_validate(OUTPUT, PDF, SCRATCH)
