"""Build and validate the Gecko resume for AMBA job 5873087057."""

from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.shared import RGBColor

from generate_extra_space_storage_resume import build_resume as build_base_resume
from generate_vitality_medical_resume import NAVY, export_and_validate, set_run


ROOT = Path(__file__).resolve().parents[1]
JOB_KEY = "5873087057"
COMPANY = "Association-Member-Benefits-Advisors"
OUTPUT = ROOT / "output" / "resumes" / f"Dave-Call+{COMPANY}+{JOB_KEY}.docx"
SCRATCH = ROOT / "scratch" / f"{COMPANY}+{JOB_KEY}"
TEMPLATE = SCRATCH / "base-template.docx"
PDF = SCRATCH / "resume-preview.pdf"


def clear_content(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_plain(paragraph, text, *, size=11, bold=False, color=None):
    clear_content(paragraph)
    set_run(paragraph.add_run(text), size=size, bold=bold, color=color)


def set_labeled(paragraph, lead, body, *, lead_color=None):
    clear_content(paragraph)
    set_run(paragraph.add_run(lead), bold=True, color=lead_color)
    set_run(paragraph.add_run(body))


def build_resume(path: Path):
    SCRATCH.mkdir(parents=True, exist_ok=True)
    build_base_resume(TEMPLATE)
    doc = docx.Document(TEMPLATE)
    paragraphs = doc.paragraphs

    set_plain(
        paragraphs[1],
        "SENIOR PAID SEARCH STRATEGIST | PPC & ANALYTICS",
        size=11.5,
        bold=True,
        color=RGBColor(40, 70, 110),
    )
    set_plain(
        paragraphs[4],
        "Hands-on paid-search strategist with 14+ years of experience building, managing, testing, and improving "
        "high-volume campaigns. Deep background in Google Ads, Microsoft Advertising, account structure, keyword "
        "and search-query analysis, negative keywords, smart bidding, ad copy, budget allocation, Quality Score, "
        "and performance reporting. Combines careful execution with GA4, Google Tag Manager, advanced analytics, "
        "clear communication, cross-functional collaboration, and a consistent focus on ROI and conversion growth.",
    )

    skills = {
        6: ("Paid Search Strategy & Execution: ", "Google Ads, Microsoft/Bing Ads, Google Ads Editor, Bing Editor, campaign and ad-group structure, keyword research, match types, search-query analysis, negative keywords, ad copy, extensions, bidding, and account optimization."),
        7: ("Performance & Budget Management: ", "CTR, CPC, CPA, ROAS, Quality Score, Ad Rank, budget pacing and allocation, smart bidding, bid experiments, account-health monitoring, anomaly detection, forecasting, and ROI analysis."),
        8: ("Measurement & Conversion: ", "GA4, Google Tag Manager, conversion tracking, attribution, funnels, conversion paths, landing-page optimization, CRO, A/B and multivariate testing, LTV, Tableau, Looker Studio, Funnel.io, and Adobe Analytics."),
        9: ("Collaboration, Technology & AI: ", "Creative and ad-copy development, SEO/SEM alignment, vendor management, project management, Excel, Python, SQL, JavaScript, Google Ads API, ChatGPT, Claude, Perplexity, Codex, Cursor, and AntiGravity."),
    }
    for index, (lead, body) in skills.items():
        set_labeled(paragraphs[index], lead, body, lead_color=NAVY)

    bullets = {
        11: ("Paid Search Strategy: ", "Evaluated max CPC versus smart bidding, brand and non-brand CPA, Performance Max and Shopping mix, audiences, attribution, and incremental CPA across scale changes."),
        12: ("Testing & Forecasting: ", "Analyzed bidding, holdouts, geo promotions, reach-to-conversion funnels, lifetime value, and regression scenarios to guide investment and optimization decisions."),
        13: ("Automated Reporting: ", "Used Python, SQL, Databricks, Excel, Funnel.io, and Tableau to automate accurate weekly reporting across more than 10 marketing platforms."),
        14: ("Cross-Platform Management: ", "Managed Google Ads, Bing Ads, Amazon, Facebook, Instagram, and TikTok across a $400K monthly budget, aligning channel investment with performance goals."),
        15: ("Performance Improvement: ", "Increased monthly volume from $200K to $800K while improving ROAS from 1.5 to 3.5 through hands-on analysis, testing, and optimization."),
        16: ("Conversion Measurement: ", "Implemented custom GA4 and Google Tag Manager reporting and built Looker Studio KPIs connecting campaign results with website, email, inventory, and budget decisions."),
        18: ("Account Architecture & Scale: ", "Built and managed 75+ Google Ads accounts totaling more than $276K per month across search, display, mobile, remarketing, YouTube, and social."),
        19: ("Monitoring & Quality: ", "Created scripts for continuous campaign monitoring and a standardized account-quality system adopted agency-wide across clients with different budgets and needs."),
        20: ("Optimization & Testing: ", "Used analytical tools to make data-driven recommendations and created conversion models, attribution models, and multivariate tests across the customer journey."),
        21: ("Client & Cross-Functional Reporting: ", "Advised clients on channel investment, consolidated performance data with Funnel.io, and presented actionable recommendations through Looker Studio reporting."),
        22: ("Campaign Operations: ", "Managed 150+ campaigns and a $400K budget across Google Ads, Bing Ads, Gemini, and Amazon, producing more than $9M in revenue—35% of company revenue."),
        23: ("Search Optimization: ", "Performed keyword, search-query, competitor, and channel analysis; optimized landing pages; created ad copy; and developed bid strategies and experiments."),
        24: ("Platform Analysis: ", "Used Excel, Google Ads Editor, Bing Editor, Google Analytics, Magento, ACCTivate, and QuickBooks for campaign, performance, and business reporting."),
        25: ("Campaign Automation: ", "Implemented human-assisted JavaScript automation based on promotions and business goals while maintaining direct oversight of campaign decisions."),
        26: ("Process & Data Leadership: ", "Led ecommerce and marketing while implementing digital systems, databases, and reporting that increased productivity, transparency, and collaboration."),
        27: ("Cross-Functional Delivery: ", "Worked with IT, operations, finance, product development, HR, and sales while managing website, technology, marketing, and project-level priorities."),
    }
    for index, (lead, body) in bullets.items():
        set_labeled(paragraphs[index], lead, body)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


if __name__ == "__main__":
    build_resume(OUTPUT)
    export_and_validate(OUTPUT, PDF, SCRATCH)
