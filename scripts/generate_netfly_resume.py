"""Build and validate the Gecko resume for NETFLY job aefcefac5c3401d3."""

from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.shared import RGBColor

from generate_marksmen_growth_resume import build_resume as build_base_resume
from generate_vitality_medical_resume import NAVY, export_and_validate, set_run


ROOT = Path(__file__).resolve().parents[1]
JOB_KEY = "aefcefac5c3401d3"
COMPANY = "NETFLY"
OUTPUT = ROOT / "output" / "resumes" / f"Dave-Call+{COMPANY}+{JOB_KEY}.docx"
SCRATCH = ROOT / "scratch" / f"{COMPANY}+{JOB_KEY}"
TEMPLATE = SCRATCH / "base-template.docx"
PDF = SCRATCH / "resume-preview.pdf"


def clear_content(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_plain(paragraph, text, *, size=11, bold=False, color=None):
    clear_content(paragraph)
    set_run(paragraph.add_run(text), size=size, bold=bold, color=color)


def set_labeled(paragraph, lead, body, *, lead_color=None):
    clear_content(paragraph)
    set_run(paragraph.add_run(lead), bold=True, color=lead_color)
    set_run(paragraph.add_run(body))


def build_resume(path: Path):
    SCRATCH.mkdir(parents=True, exist_ok=True)
    build_base_resume(TEMPLATE)
    doc = docx.Document(TEMPLATE)
    paragraphs = doc.paragraphs

    set_plain(
        paragraphs[1],
        "SENIOR PERFORMANCE MARKETING | META & YOUTUBE",
        size=11.5,
        bold=True,
        color=RGBColor(40, 70, 110),
    )
    set_plain(
        paragraphs[4],
        "Performance marketing leader with 14+ years of hands-on experience building and scaling customer-acquisition "
        "programs across paid social, video, search, display, and ecommerce. Combines channel strategy with direct "
        "execution across Meta, YouTube, TikTok, Google, audience development, creative testing, landing-page "
        "optimization, attribution, budget allocation, and performance reporting. Proven managing multi-million-dollar "
        "annual media portfolios and creating measurable growth through rigorous experimentation, analytics, automation, "
        "and cross-functional problem solving. Comfortable challenging assumptions, making difficult investment decisions, "
        "and building repeatable systems rather than simply maintaining campaigns.",
    )

    skills = {
        6: ("Paid Acquisition & Video: ", "Meta/Facebook/Instagram, YouTube, TikTok, Google Ads, Bing Ads, Amazon, LinkedIn, display, video, paid social, remarketing, campaign architecture, audience development, and budget allocation."),
        7: ("Measurement & Attribution: ", "GA4, Google Tag Manager, conversion tracking, CRM integration, data-driven attribution, LTV, CAC, ROAS, Tableau, Looker Studio, Funnel.io, Adobe Analytics, and performance dashboards."),
        8: ("Creative Testing & Conversion: ", "Creative strategy, ad copy, audience and message testing, landing-page optimization, conversion-path analysis, CRO, A/B and multivariate testing, offers, and full-funnel experimentation."),
        9: ("Systems & Analytics: ", "Python, SQL, Databricks, Excel, JavaScript, APIs, automated reporting, campaign-monitoring scripts, regression analysis, forecasting, financial analysis, and scalable operating processes."),
        10: ("Leadership & Technology: ", "Customer-acquisition strategy, team leadership, executive communication, cross-functional collaboration, ChatGPT, Claude, Perplexity, Codex, Cursor, AntiGravity, CRM, web development, and marketing automation."),
    }
    for index, (lead, body) in skills.items():
        set_labeled(paragraphs[index], lead, body, lead_color=NAVY)

    bullets = {
        12: ("Portfolio & Team Leadership: ", "Managed $30 million per month with a team of four, using customer, channel, and financial analysis to guide investment and acquisition decisions."),
        13: ("Acquisition Strategy: ", "Evaluated paid-media investment across channels, audiences, bidding approaches, brand and non-brand activity, promotions, and Performance Max."),
        14: ("Attribution & Experimentation: ", "Analyzed data-driven attribution, lifetime value, click-to-call activity, holdouts, and full-funnel reach-to-conversion tests to improve decisions."),
        15: ("Decision Systems: ", "Used Python, SQL, Databricks, Excel, Funnel.io, and Tableau to automate accurate weekly reporting across more than 10 marketing platforms."),
        16: ("Cross-Platform Ownership: ", "Managed Google, Bing, Amazon, Facebook, Instagram, and TikTok advertising across a $400K monthly marketing budget."),
        17: ("Measurable Growth: ", "Increased monthly volume from $200K to $800K while improving ROAS from 1.5 to 3.5 through hands-on analysis, testing, and optimization."),
        18: ("Tracking & Reporting: ", "Implemented custom GA4 and Google Tag Manager reporting and built Looker Studio dashboards connecting acquisition performance with operating priorities."),
        19: ("Conversion Leadership: ", "Guided website optimization, SEO, email effectiveness, content, and cross-platform budget decisions with creative, technical, inventory, and leadership teams."),
        21: ("Growth-System Development: ", "Led ecommerce and marketing for a large B2B distributor, built a functional digital environment, and expanded into B2C channels while supporting 650+ retail partners."),
        22: ("Opportunity Analysis: ", "Used daily financial and market analysis to identify niches, guide product development, evaluate channel economics, and prioritize new revenue opportunities."),
        23: ("Platform Expansion: ", "Expanded Amazon operations into CastleGate, Target.com, Costco.com, and other channels while improving product data, digital merchandising, and operating processes."),
        24: ("Builder Mindset: ", "Worked across IT, operations, finance, product development, HR, and sales to implement technology, process, cultural, and operating change."),
        25: ("Multi-Channel Acquisition: ", "Advised clients on paid search, display, mobile, remarketing, YouTube, social, websites, SEO, and email based on customer needs, business goals, and performance."),
        26: ("Hands-On Campaign Scale: ", "Built and managed 75+ Google Ads accounts totaling more than $276K per month while creating a standardized service-quality system used agency-wide."),
        27: ("Attribution & Testing: ", "Developed conversion and attribution models and multivariate tests to improve audience, message, channel, and landing-page decisions."),
        28: ("Campaign Automation: ", "Created custom scripts for continuous campaign monitoring, budget oversight, and faster identification of performance problems."),
        29: ("Acquisition Leadership: ", "Led an eight-person ecommerce team responsible for paid search, outreach, lead generation, and customer acquisition in the United States, Canada, and the UK."),
        30: ("Revenue Impact: ", "Managed 150+ campaigns and a $400K budget across Google, Bing, Gemini, and Amazon, producing more than $9M in revenue—35% of company revenue."),
        31: ("Creative & Channel Execution: ", "Performed keyword, audience, and competitive research; created ad copy and creative; and managed display, video, social, email, shopping, local, and remarketing programs."),
        32: ("Landing-Page Optimization: ", "Built attribution models, analyzed funnels and conversion paths, optimized landing pages, and ran A/B and multivariate tests to improve acquisition performance."),
        33: ("Performance Management: ", "Established annual goals, projections, and KPIs with senior management and reported business performance using Google Analytics, Magento, ACCTivate, and QuickBooks."),
    }
    for index, (lead, body) in bullets.items():
        set_labeled(paragraphs[index], lead, body)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


if __name__ == "__main__":
    build_resume(OUTPUT)
    export_and_validate(OUTPUT, PDF, SCRATCH)
