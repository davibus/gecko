"""Conservatively validate local Job Scout links without changing workbook structure.

The canonical Gecko tracker is Google Sheets. This utility exists for an explicitly
supplied local workbook and writes only Job Scout column I (Notes) and column J
(Website status). It uses openpyxl as requested and saves through an atomic replace.
"""

from __future__ import annotations

import argparse
import errno
import html
import ipaddress
import json
import os
import re
import socket
import sqlite3
import ssl
import sys
import tempfile
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Iterable
from urllib.error import HTTPError, URLError
from urllib.parse import parse_qs, urljoin, urlsplit
from urllib.request import HTTPRedirectHandler, Request, build_opener

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell


ROOT = Path(__file__).resolve().parents[1]
JOB_SCOUT_DIR = ROOT / "job-scout"
if str(JOB_SCOUT_DIR) not in sys.path:
    sys.path.insert(0, str(JOB_SCOUT_DIR))

from link_validation import official_api_url  # noqa: E402
from url_resolution import classify_url  # noqa: E402


DEFAULT_WORKBOOK = ROOT / "output" / "job-tracker.xlsx"
DEFAULT_DATABASE = JOB_SCOUT_DIR / "data" / "jobs.sqlite3"
SHEET_NAME = "Job Scout"
NOTES_COLUMN = 9
WEBSITE_STATUS_COLUMN = 10
JOB_URL_COLUMN = 21
REMOVED_NOTE = "Doesn't exist"

USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/131.0.0.0 Safari/537.36 GeckoJobLinkChecker/1.0"
)

EXPIRED_RE = re.compile(
    r"\b(?:this|the|that|requested)?\s*(?:job|position|posting|vacancy|opportunity)\s+"
    r"(?:is|has been|was)?\s*(?:no longer available|no longer open|no longer exists|"
    r"unavailable|closed|expired|removed|filled|not found)\b"
    r"|\bno longer accepting applications\b"
    r"|\bapplications? (?:are|is) (?:now )?closed\b"
    r"|\bwe (?:are|'re) no longer accepting applications\b",
    re.I,
)
BOT_RE = re.compile(
    r"cloudflare|checking your browser|verify you are human|captcha|access denied|"
    r"bot protection|automated access",
    re.I,
)
PARKED_RE = re.compile(
    r"\b(?:this )?domain (?:name )?(?:has expired|is expired|is parked|is for sale)\b"
    r"|\bbuy this domain\b|\bparked (?:free )?courtesy\b",
    re.I,
)
SCRIPT_STYLE_RE = re.compile(r"<(?:script|style)\b[^>]*>.*?</(?:script|style)>", re.I | re.S)
TAG_RE = re.compile(r"<[^>]+>")
JSON_LD_RE = re.compile(
    r"<script[^>]+type=[\"']application/ld\+json[\"'][^>]*>(.*?)</script>",
    re.I | re.S,
)
HYPERLINK_FORMULA_RE = re.compile(r'^=HYPERLINK\(\s*[\"\']([^\"\']+)', re.I)

WEBSITE_SOURCE_HEADERS = {
    "company website",
    "company website url",
    "company url",
    "employer website",
    "employer website url",
    "employer url",
    "website url",
    "corporate website",
    "homepage",
}
NON_COMPANY_HOSTS = {
    "indeed.com", "linkedin.com", "glassdoor.com", "ziprecruiter.com",
    "monster.com", "careerbuilder.com", "simplyhired.com", "jooble.org",
    "google.com", "bing.com", "facebook.com", "instagram.com", "x.com",
    "twitter.com", "youtube.com",
}


@dataclass(frozen=True)
class FetchResult:
    requested_url: str
    final_url: str = ""
    http_status: int | None = None
    content: str = ""
    error_kind: str = ""
    reason: str = ""


@dataclass(frozen=True)
class CheckResult:
    status: str  # exists, removed, unknown
    url: str
    final_url: str = ""
    reason: str = ""
    http_status: int | None = None
    content: str = ""
    retryable: bool = False


@dataclass
class Summary:
    populated_rows: int = 0
    checked: int = 0
    jobs_existing: int = 0
    jobs_removed: int = 0
    job_status_unknown: int = 0
    job_urls_missing: int = 0
    websites_existing: int = 0
    no_website: int = 0
    website_status_unknown: int = 0
    website_sources_missing: int = 0
    notes_changed: int = 0
    website_cells_changed: int = 0
    warnings: list[str] = field(default_factory=list)


class PublicRedirectHandler(HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        validate_public_url(newurl)
        return super().redirect_request(req, fp, code, msg, headers, newurl)


class ThrottledFetcher:
    def __init__(self, *, timeout: float = 12.0, delay: float = 0.75):
        self.timeout = timeout
        self.delay = max(0.0, delay)
        self._last_request = 0.0
        self._opener = build_opener(PublicRedirectHandler())

    def _throttle(self) -> None:
        remaining = self.delay - (time.monotonic() - self._last_request)
        if remaining > 0:
            time.sleep(remaining)

    def fetch(self, url: str, *, limit: int = 750_000) -> FetchResult:
        try:
            validate_public_url(url)
            self._throttle()
            request = Request(
                url,
                headers={
                    "User-Agent": USER_AGENT,
                    "Accept": "text/html,application/xhtml+xml,application/json;q=0.9,*/*;q=0.8",
                    "Accept-Language": "en-US,en;q=0.8",
                },
            )
            with self._opener.open(request, timeout=self.timeout) as response:
                self._last_request = time.monotonic()
                final_url = response.geturl()
                validate_public_url(final_url)
                body = response.read(limit).decode("utf-8", errors="replace")
                return FetchResult(url, final_url, response.status, body)
        except HTTPError as error:
            self._last_request = time.monotonic()
            body = ""
            try:
                body = error.read(limit).decode("utf-8", errors="replace")
            except (OSError, ValueError):
                pass
            return FetchResult(
                url,
                error.geturl() or url,
                error.code,
                body,
                "http",
                f"HTTP {error.code}",
            )
        except (URLError, TimeoutError, OSError, ValueError) as error:
            self._last_request = time.monotonic()
            cause = getattr(error, "reason", error)
            if isinstance(cause, (ssl.SSLError, ssl.CertificateError)):
                kind = "ssl"
            elif isinstance(cause, (TimeoutError, socket.timeout)):
                kind = "timeout"
            elif isinstance(cause, socket.gaierror):
                kind = "dns"
            elif isinstance(cause, ValueError):
                kind = "invalid"
            else:
                kind = "network"
            return FetchResult(url, error_kind=kind, reason=f"{type(cause).__name__}: {cause}")


def normalize_header(value) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def visible_text(content: str) -> str:
    value = SCRIPT_STYLE_RE.sub(" ", content)
    value = TAG_RE.sub(" ", value)
    return re.sub(r"\s+", " ", html.unescape(value)).strip()


def _resolved_addresses(host: str) -> list[ipaddress.IPv4Address | ipaddress.IPv6Address]:
    values = []
    for item in socket.getaddrinfo(host, None, type=socket.SOCK_STREAM):
        address = ipaddress.ip_address(item[4][0].split("%", 1)[0])
        if address not in values:
            values.append(address)
    return values


def validate_public_url(url: str) -> None:
    parts = urlsplit(url)
    if parts.scheme not in {"http", "https"} or not parts.hostname or parts.username or parts.password:
        raise ValueError("Refusing an invalid or non-public HTTP URL")
    host = parts.hostname.casefold().rstrip(".")
    if host == "localhost" or host.endswith(".localhost"):
        raise ValueError("Refusing a local HTTP destination")
    try:
        literal = ipaddress.ip_address(host)
    except ValueError:
        literal = None
    if literal is not None and not literal.is_global:
        raise ValueError("Refusing a private or reserved HTTP destination")
    if literal is None:
        addresses = _resolved_addresses(host)
        if not addresses or any(not address.is_global for address in addresses):
            raise ValueError("Refusing a hostname that resolves to a private or reserved address")


def host_is_public(host: str) -> tuple[str, str]:
    """Return (status, reason): public, nonexistent, or unknown."""
    try:
        addresses = _resolved_addresses(host)
    except socket.gaierror as error:
        permanent_codes = {getattr(socket, "EAI_NONAME", -2), getattr(socket, "EAI_NODATA", -5)}
        if error.errno in permanent_codes:
            return "nonexistent", f"DNS reports that {host} does not exist"
        return "unknown", f"Temporary DNS failure: {error}"
    except OSError as error:
        return "unknown", f"Temporary DNS failure: {error}"
    if not addresses:
        return "unknown", "DNS returned no addresses"
    if any(not address.is_global for address in addresses):
        return "unknown", "Domain resolves to a private or reserved address"
    return "public", ", ".join(str(address) for address in addresses[:3])


def _host_matches(host: str, patterns: Iterable[str]) -> bool:
    host = host.casefold().removeprefix("www.")
    return any(host == pattern or host.endswith("." + pattern) for pattern in patterns)


def _is_job_specific(url: str) -> bool:
    parts = urlsplit(url)
    query = parse_qs(parts.query)
    if any(key.casefold() in {"jk", "jobid", "job_id", "gh_jid", "lever-source"} for key in query):
        return True
    return bool(re.search(r"/(?:job|jobs|viewjob|positions?|careers?)/[^/?#]+", parts.path, re.I))


def _generic_removed_redirect(original_url: str, final_url: str) -> bool:
    if not final_url or not _is_job_specific(original_url):
        return False
    original = urlsplit(original_url)
    final = urlsplit(final_url)
    original_host = (original.hostname or "").casefold().removeprefix("www.")
    final_host = (final.hostname or "").casefold().removeprefix("www.")
    if original_host != final_host and not (
        original_host.endswith("." + final_host) or final_host.endswith("." + original_host)
    ):
        return False
    if "indeed." in original_host:
        original_jk = parse_qs(original.query).get("jk", [""])[0]
        final_jk = parse_qs(final.query).get("jk", [""])[0]
        return bool(original_jk and original_jk != final_jk)
    generic_paths = {"", "/", "/jobs", "/jobs/", "/search", "/careers", "/career"}
    return final.path.casefold() in generic_paths and not _is_job_specific(final_url)


def _classify_job_fetch(url: str, result: FetchResult) -> CheckResult:
    if result.error_kind:
        return CheckResult(
            "unknown", url, result.final_url, result.reason, result.http_status,
            result.content, result.error_kind in {"timeout", "network", "dns"},
        )
    status = result.http_status or 0
    text = visible_text(result.content)
    if status in {404, 410}:
        return CheckResult("removed", url, result.final_url, f"HTTP {status}", status, result.content)
    if status in {403, 429}:
        return CheckResult("unknown", url, result.final_url, f"HTTP {status}", status, result.content)
    if status >= 500:
        return CheckResult(
            "unknown", url, result.final_url, f"HTTP {status}", status,
            result.content, retryable=True,
        )
    if status < 200 or status >= 400:
        return CheckResult("unknown", url, result.final_url, f"HTTP {status}", status, result.content)
    if BOT_RE.search(text[:5000]):
        return CheckResult(
            "unknown", url, result.final_url, "Bot protection or access challenge",
            status, result.content,
        )
    if EXPIRED_RE.search(text):
        return CheckResult(
            "removed", url, result.final_url,
            "Page explicitly says the job is closed or unavailable", status, result.content,
        )
    if _generic_removed_redirect(url, result.final_url):
        return CheckResult(
            "removed", url, result.final_url,
            "Job board redirected the job-specific URL to a generic page", status, result.content,
        )
    return CheckResult("exists", url, result.final_url, "Job listing loaded", status, result.content)


def check_job_url(
    url: str,
    fetcher: ThrottledFetcher,
    *,
    retries: int = 1,
) -> CheckResult:
    api_url = official_api_url(url)
    if api_url:
        api_result = _classify_job_fetch(api_url, fetcher.fetch(api_url))
        if api_result.status == "removed":
            page_result = _classify_job_fetch(url, fetcher.fetch(url))
            if page_result.status == "exists":
                return page_result
            return CheckResult(
                "removed", url, api_result.final_url,
                "Official ATS API confirms the posting is absent",
                api_result.http_status, api_result.content,
            )

    result = _classify_job_fetch(url, fetcher.fetch(url))
    for _ in range(max(0, retries)):
        if result.status != "unknown" or not result.retryable:
            break
        result = _classify_job_fetch(url, fetcher.fetch(url))
    return result


def normalize_website_url(value: str) -> str:
    value = str(value or "").strip()
    if not value:
        return ""
    if value.startswith("www.") or "://" not in value:
        value = "https://" + value.lstrip("/")
    parts = urlsplit(value)
    if parts.scheme not in {"http", "https"} or not parts.hostname:
        return ""
    try:
        ipaddress.ip_address(parts.hostname)
    except ValueError:
        if "." not in parts.hostname:
            return ""
    return value


def company_site_from_url(value: str) -> str:
    url = normalize_website_url(value)
    if not url:
        return ""
    host = (urlsplit(url).hostname or "").casefold().removeprefix("www.")
    if _host_matches(host, NON_COMPANY_HOSTS):
        return ""
    if classify_url(url) in {"aggregator_intermediary", "official_ats", "dead_unavailable"}:
        return ""
    parts = urlsplit(url)
    return f"{parts.scheme}://{parts.netloc}/"


def _iter_json_nodes(value):
    if isinstance(value, dict):
        yield value
        for child in value.values():
            yield from _iter_json_nodes(child)
    elif isinstance(value, list):
        for child in value:
            yield from _iter_json_nodes(child)


def website_from_job_content(content: str, page_url: str) -> str:
    for block in JSON_LD_RE.findall(content or ""):
        try:
            payload = json.loads(html.unescape(block).strip())
        except (json.JSONDecodeError, TypeError):
            continue
        for node in _iter_json_nodes(payload):
            types = node.get("@type", "")
            if isinstance(types, list):
                is_job = any(str(value).casefold() == "jobposting" for value in types)
            else:
                is_job = str(types).casefold() == "jobposting"
            if not is_job:
                continue
            organization = node.get("hiringOrganization") or {}
            if not isinstance(organization, dict):
                continue
            for key in ("url", "sameAs"):
                values = organization.get(key) or []
                if not isinstance(values, list):
                    values = [values]
                for value in values:
                    site = company_site_from_url(urljoin(page_url, str(value)))
                    if site:
                        return site
    return ""


def check_company_website(
    url: str,
    fetcher: ThrottledFetcher,
    *,
    retries: int = 1,
) -> CheckResult:
    url = normalize_website_url(url)
    if not url:
        return CheckResult("unknown", url, reason="No usable company website URL")
    host = urlsplit(url).hostname or ""
    dns_status, dns_reason = host_is_public(host)
    dns_attempts = 0
    while dns_status == "unknown" and dns_attempts < max(0, retries):
        time.sleep(getattr(fetcher, "delay", 0.0))
        dns_status, dns_reason = host_is_public(host)
        dns_attempts += 1
    if dns_status == "nonexistent":
        return CheckResult("removed", url, reason=dns_reason)
    if dns_status != "public":
        return CheckResult("unknown", url, reason=dns_reason)

    result = fetcher.fetch(url)
    attempts = 0
    while result.error_kind in {"timeout", "network", "dns"} and attempts < max(0, retries):
        result = fetcher.fetch(url)
        attempts += 1
    if result.error_kind == "ssl":
        return CheckResult("exists", url, reason="Domain resolves; site has an SSL problem")
    if result.error_kind:
        return CheckResult("unknown", url, reason=result.reason, retryable=False)
    text = visible_text(result.content)
    if PARKED_RE.search(text[:20_000]):
        return CheckResult(
            "removed", url, result.final_url,
            "The domain is parked, for sale, or expired", result.http_status, result.content,
        )
    return CheckResult(
        "exists", url, result.final_url,
        f"Domain resolves and returned HTTP {result.http_status}",
        result.http_status, result.content,
    )


def cell_url(cell: Cell) -> str:
    if cell.hyperlink and cell.hyperlink.target:
        return normalize_website_url(cell.hyperlink.target)
    value = cell.value
    if isinstance(value, str):
        formula = HYPERLINK_FORMULA_RE.match(value)
        if formula:
            return normalize_website_url(formula.group(1))
        return normalize_website_url(value)
    return ""


class ProjectWebsiteLookup:
    """Read-only lookup for direct employer domains already stored by Job Scout."""

    def __init__(self, path: Path = DEFAULT_DATABASE):
        self.path = path

    def find(self, scout_id, company: str, title: str) -> str:
        if not self.path.is_file():
            return ""
        connection = sqlite3.connect(f"file:{self.path.as_posix()}?mode=ro", uri=True)
        connection.row_factory = sqlite3.Row
        try:
            row = None
            if str(scout_id or "").strip().isdigit():
                row = connection.execute(
                    """
                    SELECT j.url, j.canonical_url, e.enriched_source_url,
                           u.authoritative_url, u.destination_type, u.redirect_url
                    FROM jobs j
                    LEFT JOIN enrichments e ON e.job_id = j.id
                    LEFT JOIN url_resolutions u ON u.job_id = j.id
                    WHERE j.id = ?
                    """,
                    (int(scout_id),),
                ).fetchone()
            if row is None and company and title:
                row = connection.execute(
                    """
                    SELECT j.url, j.canonical_url, e.enriched_source_url,
                           u.authoritative_url, u.destination_type, u.redirect_url
                    FROM jobs j
                    LEFT JOIN enrichments e ON e.job_id = j.id
                    LEFT JOIN url_resolutions u ON u.job_id = j.id
                    WHERE lower(trim(j.company)) = lower(trim(?))
                      AND lower(trim(j.title)) = lower(trim(?))
                    ORDER BY j.id DESC LIMIT 1
                    """,
                    (company, title),
                ).fetchone()
            if row is None:
                return ""
            if row["destination_type"] == "official_employer_domain":
                site = company_site_from_url(row["authoritative_url"] or row["redirect_url"] or "")
                if site:
                    return site
            for field_name in ("enriched_source_url", "authoritative_url", "redirect_url", "canonical_url", "url"):
                site = company_site_from_url(row[field_name] or "")
                if site:
                    return site
            return ""
        except sqlite3.DatabaseError:
            return ""
        finally:
            connection.close()


def _worksheet(workbook):
    exact = [sheet for sheet in workbook.worksheets if sheet.title == SHEET_NAME]
    if exact:
        return exact[0]
    matches = [sheet for sheet in workbook.worksheets if sheet.title.strip().casefold() == SHEET_NAME.casefold()]
    if len(matches) == 1:
        return matches[0]
    raise ValueError(f"Workbook does not contain the required {SHEET_NAME!r} worksheet")


def _headers(worksheet) -> tuple[dict[str, int], list[str]]:
    values = [worksheet.cell(1, column).value for column in range(1, worksheet.max_column + 1)]
    mapping = {
        normalize_header(value): column
        for column, value in enumerate(values, start=1)
        if normalize_header(value)
    }
    return mapping, [str(value or "") for value in values]


def validate_layout(worksheet) -> tuple[dict[str, int], list[int]]:
    headers, raw_headers = _headers(worksheet)
    if normalize_header(worksheet.cell(1, NOTES_COLUMN).value) != "notes":
        raise ValueError(
            f"Expected column I to be 'Notes'; found {worksheet.cell(1, NOTES_COLUMN).value!r}"
        )
    if normalize_header(worksheet.cell(1, WEBSITE_STATUS_COLUMN).value) not in {
        "website", "website status"
    }:
        raise ValueError(
            f"Expected column J to be 'Website'; found {worksheet.cell(1, WEBSITE_STATUS_COLUMN).value!r}"
        )
    job_header = normalize_header(worksheet.cell(1, JOB_URL_COLUMN).value)
    if job_header not in {"job url", "job link", "url", "posting url", "application url"}:
        raise ValueError(
            f"Expected column U to contain job links; found {worksheet.cell(1, JOB_URL_COLUMN).value!r}"
        )
    if "company" not in headers:
        raise ValueError("Job Scout is missing a Company column")
    website_columns = [
        column for column, header in enumerate(raw_headers, start=1)
        if column != WEBSITE_STATUS_COLUMN and normalize_header(header) in WEBSITE_SOURCE_HEADERS
    ]
    return headers, website_columns


def _row_is_populated(worksheet, row: int) -> bool:
    return any(
        worksheet.cell(row, column).value not in (None, "")
        for column in range(1, worksheet.max_column + 1)
    )


def _append_removed_note(cell: Cell) -> bool:
    if cell.data_type == "f":
        return False
    current = str(cell.value or "").strip()
    if REMOVED_NOTE.casefold() in current.casefold():
        return False
    cell.value = f"{current}\n{REMOVED_NOTE}" if current else REMOVED_NOTE
    return True


def _set_website_status(cell: Cell, value: str) -> bool:
    if cell.data_type == "f" or cell.value == value:
        return False
    cell.value = value
    return True


def _explicit_website(worksheet, row: int, website_columns: list[int]) -> str:
    for column in website_columns:
        value = cell_url(worksheet.cell(row, column))
        if value:
            return value
    return ""


def _company_hyperlink(worksheet, row: int, company_column: int) -> str:
    cell = worksheet.cell(row, company_column)
    if cell.hyperlink and cell.hyperlink.target:
        return company_site_from_url(cell.hyperlink.target)
    return ""


def process_workbook(
    workbook,
    *,
    job_checker: Callable[[str], CheckResult],
    website_checker: Callable[[str], CheckResult],
    project_lookup: ProjectWebsiteLookup | None = None,
    limit: int | None = None,
    start_row: int = 2,
    progress: Callable[[str], None] = print,
) -> Summary:
    worksheet = _worksheet(workbook)
    headers, website_columns = validate_layout(worksheet)
    company_column = headers["company"]
    title_column = headers.get("job title") or headers.get("title")
    scout_id_column = headers.get("scout id")
    summary = Summary()

    rows = [
        row for row in range(max(2, start_row), worksheet.max_row + 1)
        if _row_is_populated(worksheet, row)
    ]
    if limit is not None:
        rows = rows[: max(0, limit)]

    for row in rows:
        summary.populated_rows += 1
        company = str(worksheet.cell(row, company_column).value or "").strip()
        title = str(worksheet.cell(row, title_column).value or "").strip() if title_column else ""
        scout_id = worksheet.cell(row, scout_id_column).value if scout_id_column else None
        job_url = cell_url(worksheet.cell(row, JOB_URL_COLUMN))
        progress(f"Checking row {row}: {company or '(company missing)'}")

        job_result = CheckResult("unknown", job_url, reason="Job URL is empty")
        if job_url:
            summary.checked += 1
            job_result = job_checker(job_url)
            if job_result.status == "exists":
                summary.jobs_existing += 1
                progress("  Job: EXISTS")
            elif job_result.status == "removed":
                summary.jobs_removed += 1
                notes = worksheet.cell(row, NOTES_COLUMN)
                if _append_removed_note(notes):
                    summary.notes_changed += 1
                elif notes.data_type == "f":
                    summary.warnings.append(f"I{row} is a formula; confirmed removal was not written")
                progress("  Job: DOESN'T EXIST")
            else:
                summary.job_status_unknown += 1
                progress(f"  Job: UNKNOWN ({job_result.reason})")
        else:
            summary.job_urls_missing += 1
            progress("  Job: SKIPPED (no URL in column U)")

        website_url = _explicit_website(worksheet, row, website_columns)
        if not website_url:
            website_url = _company_hyperlink(worksheet, row, company_column)
        if not website_url and job_result.content:
            website_url = website_from_job_content(
                job_result.content, job_result.final_url or job_url
            )
        if not website_url:
            website_url = company_site_from_url(job_result.final_url or job_url)
        if not website_url and project_lookup is not None:
            website_url = project_lookup.find(scout_id, company, title)

        if website_url:
            website_result = website_checker(website_url)
            if website_result.status == "exists":
                summary.websites_existing += 1
                website_cell = worksheet.cell(row, WEBSITE_STATUS_COLUMN)
                if _set_website_status(website_cell, "Yes"):
                    summary.website_cells_changed += 1
                elif website_cell.data_type == "f":
                    summary.warnings.append(f"J{row} is a formula; website status was not written")
                progress("  Website: YES")
            elif website_result.status == "removed":
                summary.no_website += 1
                website_cell = worksheet.cell(row, WEBSITE_STATUS_COLUMN)
                if _set_website_status(website_cell, "No Website"):
                    summary.website_cells_changed += 1
                elif website_cell.data_type == "f":
                    summary.warnings.append(f"J{row} is a formula; website status was not written")
                progress("  Website: NO WEBSITE")
            else:
                summary.website_status_unknown += 1
                progress(f"  Website: UNKNOWN ({website_result.reason})")
        else:
            summary.website_sources_missing += 1
            progress("  Website: UNKNOWN (no company website URL found)")
        progress("")
    return summary


def safe_save(workbook, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{destination.stem}-",
        suffix=destination.suffix,
        dir=destination.parent,
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        workbook.save(temporary)
        os.replace(temporary, destination)
    finally:
        temporary.unlink(missing_ok=True)


def _project_path(value: str | Path) -> Path:
    path = Path(value)
    return path if path.is_absolute() else ROOT / path


def _print_summary(summary: Summary, *, saved_to: Path | None, dry_run: bool) -> None:
    print("Summary")
    print(f"Populated rows: {summary.populated_rows}")
    print(f"Checked: {summary.checked}")
    print(f"Jobs existing: {summary.jobs_existing}")
    print(f"Jobs removed: {summary.jobs_removed}")
    print(f"Job status unknown: {summary.job_status_unknown}")
    print(f"Job URLs missing: {summary.job_urls_missing}")
    print(f"Websites existing: {summary.websites_existing}")
    print(f"No website: {summary.no_website}")
    print(f"Website status unknown: {summary.website_status_unknown}")
    print(f"Website source missing: {summary.website_sources_missing}")
    print(f"Notes cells changed: {summary.notes_changed}")
    print(f"Website cells changed: {summary.website_cells_changed}")
    if summary.warnings:
        print("Warnings:")
        for warning in summary.warnings:
            print(f"  - {warning}")
    if dry_run:
        print("Dry run: workbook was not saved")
    elif saved_to is not None:
        print(f"Saved: {saved_to}")


def run(args: argparse.Namespace) -> int:
    source = _project_path(args.workbook)
    if not source.is_file():
        raise FileNotFoundError(
            errno.ENOENT,
            "Workbook not found. Restore/export the requested workbook before running the check",
            str(source),
        )
    destination = _project_path(args.output) if args.output else source
    keep_vba = source.suffix.casefold() == ".xlsm"
    workbook = load_workbook(source, keep_links=True, keep_vba=keep_vba, data_only=False)
    fetcher = ThrottledFetcher(timeout=args.timeout, delay=args.delay)
    project_lookup = ProjectWebsiteLookup(_project_path(args.database)) if args.database else None
    summary = process_workbook(
        workbook,
        job_checker=lambda url: check_job_url(url, fetcher, retries=args.retries),
        website_checker=lambda url: check_company_website(url, fetcher, retries=args.retries),
        project_lookup=project_lookup,
        limit=args.limit,
        start_row=args.start_row,
    )
    if not args.dry_run:
        safe_save(workbook, destination)
    _print_summary(summary, saved_to=destination, dry_run=args.dry_run)
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--workbook", default=str(DEFAULT_WORKBOOK.relative_to(ROOT)),
        help="Workbook to inspect and update (default: output/job-tracker.xlsx)",
    )
    parser.add_argument("--output", help="Save to a separate workbook instead of updating in place")
    parser.add_argument(
        "--database", default=str(DEFAULT_DATABASE.relative_to(ROOT)),
        help="Optional read-only Job Scout SQLite data used to find employer domains",
    )
    parser.add_argument("--limit", type=int, help="Process only the first N populated rows")
    parser.add_argument("--start-row", type=int, default=2, help="First worksheet row to process")
    parser.add_argument("--timeout", type=float, default=12.0, help="Per-request timeout in seconds")
    parser.add_argument("--delay", type=float, default=0.75, help="Minimum delay between HTTP requests")
    parser.add_argument("--retries", type=int, default=1, help="Retries for temporary network failures")
    parser.add_argument("--dry-run", action="store_true", help="Check links without saving workbook changes")
    return parser


def main() -> int:
    try:
        return run(build_parser().parse_args())
    except (FileNotFoundError, ValueError, OSError) as error:
        print(f"Job-link check failed: {error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
