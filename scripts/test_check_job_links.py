from __future__ import annotations

import copy
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from check_job_links import (  # noqa: E402
    CheckResult,
    FetchResult,
    _classify_job_fetch,
    _generic_removed_redirect,
    check_company_website,
    check_job_url,
    process_workbook,
    safe_save,
)


HEADERS = [
    "Scout ID", "Source", "Company", "Job Title", "Gecko Status", "Apply?",
    "Resume Created", "Applied", "Notes", "Website", "Match Score",
    "Evidence Confidence", "Match Status", "Location", "Work Arrangement",
    "Employment Type", "Salary", "Date Posted", "Date Found", "Last Seen",
    "Job URL", "Company Website URL",
]


def workbook_fixture() -> Workbook:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Job Scout"
    worksheet.append(HEADERS)
    worksheet.append([
        1, "test", "Acme", "Paid Search Manager", "New", "yes", "", "",
        "Recruiter contacted", "", 90, 95, "Strong", "Remote", "remote",
        "full-time", "", "", "", "", "https://jobs.test/removed",
        "https://acme.test",
    ])
    worksheet.append([
        2, "test", "Example", "Growth Manager", "New", "yes", "", "",
        "", "Old value", 80, 90, "Near", "Utah", "hybrid", "full-time",
        "", "", "", "", "https://jobs.test/exists", "https://missing.test",
    ])
    worksheet.append([
        3, "test", "Blocked", "SEO Manager", "New", "yes", "", "",
        "Keep this", "Keep website", 70, 80, "Near", "Utah", "remote",
        "full-time", "", "", "", "", "https://jobs.test/unknown", "",
    ])
    worksheet.freeze_panes = "K2"
    worksheet.auto_filter.ref = "A1:V4"
    worksheet.row_dimensions[2].height = 31
    worksheet.column_dimensions["I"].width = 30
    worksheet.column_dimensions["H"].hidden = True
    worksheet.row_dimensions[4].hidden = True
    worksheet["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    worksheet["A1"].font = Font(color="FFFFFF", bold=True)
    worksheet["K2"] = "=1+1"
    worksheet["U2"].hyperlink = "https://jobs.test/removed"
    worksheet.conditional_formatting.add(
        "A2:V4",
        FormulaRule(formula=['$E2="New"'], fill=PatternFill("solid", fgColor="FFF2CC")),
    )
    table = Table(displayName="JobScoutTable", ref="A1:V4")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    worksheet.add_table(table)
    return workbook


def structure_snapshot(workbook: Workbook) -> dict:
    result = {"sheetnames": workbook.sheetnames, "sheets": []}
    for worksheet in workbook.worksheets:
        cells = {}
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.value is None and not cell.has_style and not cell.hyperlink:
                    continue
                cells[cell.coordinate] = {
                    "value": cell.value,
                    "style": copy.copy(cell._style) if cell.has_style else None,
                    "number_format": cell.number_format,
                    "hyperlink": cell.hyperlink.target if cell.hyperlink else None,
                }
        result["sheets"].append({
            "title": worksheet.title,
            "freeze": str(worksheet.freeze_panes),
            "filter": worksheet.auto_filter.ref,
            "dimensions": worksheet.calculate_dimension(),
            "row_heights": {key: value.height for key, value in worksheet.row_dimensions.items()},
            "row_hidden": {key: value.hidden for key, value in worksheet.row_dimensions.items()},
            "column_widths": {key: value.width for key, value in worksheet.column_dimensions.items()},
            "column_hidden": {key: value.hidden for key, value in worksheet.column_dimensions.items()},
            "tables": [(table.name, table.ref, table.tableStyleInfo.name) for table in worksheet.tables.values()],
            "conditional_formats": [str(item) for item in worksheet.conditional_formatting],
            "cells": cells,
        })
    return result


class ClassificationTests(unittest.TestCase):
    def test_dead_and_unknown_rules_are_conservative(self):
        self.assertEqual(
            _classify_job_fetch("https://jobs.test/1", FetchResult("x", "x", 404)).status,
            "removed",
        )
        self.assertEqual(
            _classify_job_fetch("https://jobs.test/1", FetchResult("x", "x", 403)).status,
            "unknown",
        )
        self.assertEqual(
            _classify_job_fetch(
                "https://jobs.test/1", FetchResult("x", "x", 200, "This job is no longer available")
            ).status,
            "removed",
        )
        self.assertEqual(
            _classify_job_fetch(
                "https://jobs.test/1", FetchResult("x", "x", 200, "Checking your browser - Cloudflare")
            ).status,
            "unknown",
        )

    def test_indeed_generic_redirect_is_removed(self):
        self.assertTrue(_generic_removed_redirect(
            "https://www.indeed.com/viewjob?jk=abc123", "https://www.indeed.com/jobs"
        ))
        self.assertFalse(_generic_removed_redirect(
            "https://www.indeed.com/viewjob?jk=abc123",
            "https://www.indeed.com/viewjob?jk=abc123",
        ))
        self.assertTrue(_generic_removed_redirect(
            "https://careers.example.com/jobs/paid-search-manager",
            "https://careers.example.com/careers",
        ))

    def test_resolved_domain_with_ssl_problem_counts_as_existing(self):
        fetcher = type("Fetcher", (), {
            "fetch": lambda self, url: FetchResult(url, error_kind="ssl", reason="certificate")
        })()
        with patch("check_job_links.host_is_public", return_value=("public", "203.0.113.1")):
            result = check_company_website("https://example.test", fetcher)
        self.assertEqual(result.status, "exists")

    def test_nxdomain_is_no_website(self):
        fetcher = type("Fetcher", (), {"fetch": lambda self, url: None})()
        with patch(
            "check_job_links.host_is_public",
            return_value=("nonexistent", "DNS reports that example.test does not exist"),
        ):
            result = check_company_website("https://example.test", fetcher)
        self.assertEqual(result.status, "removed")

    def test_temporary_job_failure_is_retried_once(self):
        class Fetcher:
            def __init__(self):
                self.results = [
                    FetchResult("x", error_kind="timeout", reason="timed out"),
                    FetchResult("x", "https://jobs.example.test/1", 200, "Paid Search Manager"),
                ]

            def fetch(self, _url):
                return self.results.pop(0)

        fetcher = Fetcher()
        with patch("check_job_links.validate_public_url", return_value=None):
            result = check_job_url("https://jobs.example.test/1", fetcher, retries=1)
        self.assertEqual(result.status, "exists")
        self.assertEqual(fetcher.results, [])


class WorkbookPreservationTests(unittest.TestCase):
    def test_only_notes_and_website_values_change(self):
        workbook = workbook_fixture()
        before = structure_snapshot(workbook)

        job_results = {
            "https://jobs.test/removed": CheckResult("removed", "", reason="HTTP 404"),
            "https://jobs.test/exists": CheckResult("exists", ""),
            "https://jobs.test/unknown": CheckResult("unknown", "", reason="HTTP 403"),
        }
        website_results = {
            "https://acme.test": CheckResult("exists", ""),
            "https://missing.test": CheckResult("removed", ""),
            "https://jobs.test": CheckResult("unknown", "", reason="No confirmed company site"),
        }
        summary = process_workbook(
            workbook,
            job_checker=lambda url: job_results[url],
            website_checker=lambda url: website_results[url.rstrip("/")],
            project_lookup=None,
            progress=lambda _message: None,
        )

        self.assertEqual(workbook["Job Scout"]["I2"].value, "Recruiter contacted\nDoesn't exist")
        self.assertEqual(workbook["Job Scout"]["J2"].value, "Yes")
        self.assertEqual(workbook["Job Scout"]["I3"].value, "")
        self.assertEqual(workbook["Job Scout"]["J3"].value, "No Website")
        self.assertEqual(workbook["Job Scout"]["I4"].value, "Keep this")
        self.assertEqual(workbook["Job Scout"]["J4"].value, "Keep website")
        self.assertEqual((summary.jobs_removed, summary.jobs_existing, summary.job_status_unknown), (1, 1, 1))

        after = structure_snapshot(workbook)
        allowed = {"I2", "J2", "J3"}
        before_cells = before["sheets"][0].pop("cells")
        after_cells = after["sheets"][0].pop("cells")
        self.assertEqual(before, after)
        for coordinate in set(before_cells) | set(after_cells):
            left = before_cells.get(coordinate)
            right = after_cells.get(coordinate)
            if coordinate in allowed:
                left = {**left, "value": right["value"]}
            self.assertEqual(left, right, coordinate)

    def test_round_trip_preserves_structure(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsx"
            output = Path(directory) / "output.xlsx"
            workbook_fixture().save(source)
            workbook = load_workbook(source, keep_links=True)
            before = structure_snapshot(workbook)
            process_workbook(
                workbook,
                job_checker=lambda url: CheckResult("exists", url),
                website_checker=lambda url: CheckResult("exists", url),
                project_lookup=None,
                progress=lambda _message: None,
            )
            safe_save(workbook, output)
            reopened = load_workbook(output, keep_links=True)
            after = structure_snapshot(reopened)

            before_cells = before["sheets"][0].pop("cells")
            after_cells = after["sheets"][0].pop("cells")
            self.assertEqual(before, after)
            for coordinate, left in before_cells.items():
                right = after_cells[coordinate]
                if coordinate in {"J2", "J3", "J4"}:
                    left = {**left, "value": "Yes"}
                self.assertEqual(left, right, coordinate)


if __name__ == "__main__":
    unittest.main()
