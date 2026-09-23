#!/usr/bin/env python
"""Gecko Job Scout command line interface."""

from __future__ import annotations

import argparse
import json
import os
import sys
import tempfile
from dataclasses import replace
from pathlib import Path
from openpyxl import load_workbook

from enrichment import enrich_and_rescore
from deduplicate import find_duplicate
from google_sheets_sync import sync_workbook_to_google, tracker_backend
from handoff import archive_listing
from link_validation import DailyLinkValidator
from models import RawListing
from normalize import normalize
from normalize import canonicalize_url
from preferences import load_preferences
from resume_evidence import extract_resume_text
from review import build_review_queue, format_review_queue, queue_to_json
from scoring import score_job
from service import discover, discover_remotive_full_feed
from sources import (
    AdzunaProvider, IndeedProvider, JoobleProvider, ProviderError,
    RemotiveProvider, WebCareerProvider,
)
from sources.base import SearchRequest
from storage import DEFAULT_DB, JobStore
from tracker_sync import sync_job_scout
from url_resolution import resolve_and_store, url_status_label


PROJECT_ROOT = Path(__file__).resolve().parent.parent
MASTER_RESUME = PROJECT_ROOT / "input" / "master-resume" / "Dave-Call-resume-9-23-26.docx"
DEFAULT_SOURCE = "adzuna"
DEFAULT_TRACKER = PROJECT_ROOT / "output" / "job-tracker.xlsx"
CORE_PROVIDERS = ("adzuna", "jooble", "remotive", "web-careers")


def load_local_environment(path: Path = PROJECT_ROOT / ".env.local") -> None:
    """Load simple KEY=VALUE entries without replacing shell-provided credentials."""
    if not path.is_file():
        return
    for number, raw_line in enumerate(path.read_text(encoding="utf-8-sig").splitlines(), 1):
        line = raw_line.strip()
        if not line or line.startswith("#"):
            continue
        if line.startswith("export "):
            line = line[7:].strip()
        if "=" not in line:
            raise ValueError(f"Invalid environment entry at {path.name}:{number}")
        name, value = line.split("=", 1)
        name = name.strip()
        if not name or not name.replace("_", "").isalnum() or name[0].isdigit():
            raise ValueError(f"Invalid environment variable name at {path.name}:{number}")
        value = value.strip()
        if len(value) >= 2 and value[0] == value[-1] and value[0] in "\"'":
            value = value[1:-1]
        os.environ.setdefault(name, value)


def providers():
    return {
        provider.name: provider
        for provider in (
            AdzunaProvider(), JoobleProvider(), RemotiveProvider(),
            IndeedProvider(), WebCareerProvider(),
        )
    }


def get_job(store: JobStore, job_id: int):
    job = store.get(job_id)
    if not job:
        raise ValueError(f"Job {job_id} was not found")
    return job


def sync_tracker(
    store: JobStore,
    tracker_path: str | Path,
    *,
    force_google: bool = False,
    jobs=None,
    append_only: bool = False,
    remove_scout_ids: set[int] | None = None,
):
    summary = sync_job_scout(
        store.all() if jobs is None else jobs, tracker_path, append_only=append_only,
        remove_scout_ids=remove_scout_ids,
    )
    result = {
        "job_scout_worksheet": {
            "rows": summary.rows,
            "confirmed": summary.confirmed,
            "provisional": summary.provisional,
            "near_matches": summary.near_matches,
            "tracker_preserved": summary.tracker_preserved,
        }
    }
    if force_google or tracker_backend() == "google-sheets":
        google = sync_workbook_to_google(tracker_path, remove_scout_ids=remove_scout_ids)
        result["google_sheets"] = {
            "spreadsheet_id": google.spreadsheet_id,
            "job_tracker_rows": google.application_rows,
            "job_scout_rows": google.scout_rows,
        }
    print(json.dumps(result, indent=2))
    return summary


def sync_sheets(args, store, _preferences):
    """Synchronize both XLSX worksheets to Google Sheets on demand."""
    sync_tracker(store, args.tracker, force_google=True)
    return 0


def protected_job_ids(jobs, tracker_path: str | Path) -> set[int]:
    """Preserve application history even if a Scout lifecycle flag is stale."""
    protected = {job.id for job in jobs if job.id is not None and job.status not in {"new", "reviewing"}}
    path = Path(tracker_path)
    if not path.is_file():
        return protected
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Job Scout" in workbook:
            sheet = workbook["Job Scout"]
            records = list(sheet.values)
            if records:
                columns = {str(value): index for index, value in enumerate(records[0]) if value is not None}
                for row in records[1:]:
                    scout_id = row[columns["Scout ID"]] if "Scout ID" in columns else None
                    if scout_id in (None, ""):
                        continue
                    manual = any(row[columns[field]] not in (None, "") for field in ("Apply?", "Applied", "Contacted", "Resume Created") if field in columns)
                    lifecycle = str(row[columns["Gecko Status"]] or "").lower() if "Gecko Status" in columns else ""
                    if manual or lifecycle not in {"", "new", "reviewing"}:
                        protected.add(int(scout_id))
        if "Job Tracker" in workbook:
            sheet = workbook["Job Tracker"]
            records = list(sheet.values)
            if records:
                columns = {str(value): index for index, value in enumerate(records[0]) if value is not None}
                for row in records[1:]:
                    def value(field):
                        return row[columns[field]] if field in columns else None
                    link = canonicalize_url(str(value("Job Link") or ""))
                    company = str(value("Company") or "").casefold().strip()
                    title = str(value("Job Title") or "").casefold().strip()
                    number = str(value("Job Number") or "").casefold().strip()
                    for job in jobs:
                        if job.id is None:
                            continue
                        urls = {canonicalize_url(url) for url in (job.url, job.canonical_url, job.authoritative_url) if url}
                        if (link and link in urls) or (number and number == job.source_job_id.casefold()) or (company and title and company == job.company.casefold().strip() and title == job.title.casefold().strip()):
                            protected.add(job.id)
    finally:
        workbook.close()
    return protected


def scout_row_ids(tracker_path: str | Path) -> set[int]:
    path = Path(tracker_path)
    if not path.is_file():
        return set()
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Job Scout" not in workbook:
            return set()
        rows = workbook["Job Scout"].values
        headers = next(rows, ())
        if "Scout ID" not in headers:
            return set()
        column = headers.index("Scout ID")
        return {int(row[column]) for row in rows if len(row) > column and row[column] not in (None, "")}
    finally:
        workbook.close()


def sheet_only_active_jobs(tracker_path: str | Path, known_ids: set[int]):
    """Include active worksheet rows that have no matching local SQLite record."""
    path = Path(tracker_path)
    if not path.is_file():
        return []
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Job Scout" not in workbook:
            return []
        rows = workbook["Job Scout"].values
        headers = next(rows, ())
        columns = {str(value): index for index, value in enumerate(headers) if value is not None}
        result = []
        for row in rows:
            def value(header):
                column = columns.get(header)
                return row[column] if column is not None and column < len(row) else None
            try:
                scout_id = int(value("Scout ID"))
            except (TypeError, ValueError):
                continue
            if scout_id in known_ids:
                continue
            status = str(value("Gecko Status") or "new").strip().lower().replace(" ", "-")
            if status in {"ignored", "rejected"}:
                continue
            url = str(value("Job URL") or "")
            item = normalize(RawListing(source="worksheet", source_job_id=str(scout_id), url=url,
                                        title=str(value("Job Title") or ""),
                                        company=str(value("Company") or "")))
            item.id = scout_id
            item.status = status
            item.authoritative_url = str(value("Authoritative URL") or "")
            result.append(item)
        return result
    finally:
        workbook.close()


def search(args, store, preferences):
    daily_mode = bool(getattr(args, "daily_mode", False))
    link_validator = getattr(args, "link_validator", None) if daily_mode else None
    preexisting_ids = {job.id for job in store.all()}
    available = providers()
    if args.source == "all":
        selected = list(available)
    elif args.source == "core":
        selected = list(CORE_PROVIDERS)
    else:
        selected = [args.source]
    queries = args.query or preferences["target_roles"]
    locations = args.location or ["Utah", "Remote"]
    # Remotive is globally remote and uses category RSS, never Utah/title discovery.
    requests = (
        [SearchRequest(query, location, args.page, args.results)
         for query in queries for location in locations]
        if any(name != "remotive" for name in selected) else []
    )
    resume_text = extract_resume_text(MASTER_RESUME)
    aggregate = {
        "raw_retrieved": 0, "rss_retrieved": 0, "rss_unique": 0,
        "rss_duplicates_removed": 0, "successful_feeds": 0, "failed_feeds": 0,
        "api_retrieved": 0,
        "fetched": 0, "normalized": 0, "scored": 0,
        "added": 0, "updated": 0, "strong": 0, "provisional": 0, "weak": 0,
        "duplicates": 0, "cross_provider_duplicates": 0,
        "score_80_plus": 0, "score_70_79": 0,
        "score_below_70": 0,
        "filtered_before_scoring": 0,
        "newly_discovered": 0, "already_existed": 0,
        "new_job_ids": [], "existing_job_ids": [],
        "unique_remotive_jobs_available": 0, "unique_jobs_imported": 0,
        "eligibility_includes_usa": 0, "eligibility_worldwide": 0,
        "top_25_by_gecko_match_score": [],
        "source_counts": {name: 0 for name in selected}, "source_backends": {},
        "source_diagnostics": {},
        "google_cse": {},
        "skipped_sources": [], "source_errors": {},
    }
    successful_sources = 0
    for name in selected:
        provider = available[name]
        if name == "web-careers":
            web_diagnostics = provider.diagnostics()
            aggregate["source_diagnostics"][name] = web_diagnostics
            aggregate["google_cse"] = web_diagnostics["google_cse"]
            if provider.backend:
                aggregate["source_backends"][name] = provider.backend
        if not provider.configured():
            aggregate["skipped_sources"].append(name)
            continue
        try:
            if name == "remotive":
                summary = discover_remotive_full_feed(
                    provider, store, preferences, resume_text, getattr(args, "limit", 100),
                    preserve_existing=daily_mode, preexisting_ids=preexisting_ids,
                    link_validator=link_validator,
                )
            else:
                summary = discover(
                    provider, requests, store, preferences, resume_text, args.minimum_score,
                    preserve_existing=daily_mode, preexisting_ids=preexisting_ids,
                    link_validator=link_validator,
                )
        except ProviderError as error:
            aggregate["source_errors"][name] = str(error)
            print(f"Warning: {name} provider failed: {error}", file=sys.stderr)
            continue
        successful_sources += 1
        aggregate["source_counts"][name] = (
            (summary.unique_imported or summary.fetched)
            if name == "remotive" else summary.fetched
        )
        if summary.source_backend:
            aggregate["source_backends"][name] = summary.source_backend
        if name == "web-careers":
            web_diagnostics = provider.diagnostics(
                jobs_added=summary.added,
                backend_qualifying=summary.backend_qualifying,
                backend_added=summary.backend_added,
            )
            aggregate["source_diagnostics"][name] = web_diagnostics
            aggregate["google_cse"] = web_diagnostics["google_cse"]
        if name == "remotive":
            aggregate["source_diagnostics"][name] = {
                "api_endpoint": getattr(provider, "api_endpoint", ""),
                "rss_error": getattr(provider, "rss_error", ""),
                "discovery_location": "Remote",
                "feeds_attempted": summary.successful_feeds + summary.failed_feeds,
                "category_feeds": summary.feed_results,
            }
            aggregate["unique_remotive_jobs_available"] = summary.unique_available
            aggregate["unique_jobs_imported"] = summary.unique_imported
            aggregate["eligibility_includes_usa"] = summary.eligibility_includes_usa
            aggregate["eligibility_worldwide"] = summary.eligibility_worldwide
            aggregate["top_25_by_gecko_match_score"] = summary.top_jobs
        for key in (
            "raw_retrieved", "rss_retrieved", "rss_unique",
            "rss_duplicates_removed", "successful_feeds", "failed_feeds",
            "api_retrieved", "fetched",
            "normalized", "scored", "added", "updated",
            "strong", "provisional", "weak", "duplicates", "cross_provider_duplicates",
            "score_80_plus",
            "score_70_79", "score_below_70",
            "filtered_before_scoring",
        ):
            aggregate[key] += getattr(summary, key)
        for job_id in summary.new_job_ids:
            if job_id not in aggregate["new_job_ids"]:
                aggregate["new_job_ids"].append(job_id)
        for job_id in summary.existing_job_ids:
            if job_id not in aggregate["existing_job_ids"]:
                aggregate["existing_job_ids"].append(job_id)
    aggregate["newly_discovered"] = len(aggregate["new_job_ids"])
    aggregate["already_existed"] = len(aggregate["existing_job_ids"])
    args.run_result = aggregate
    print(json.dumps(aggregate, indent=2))
    if len(aggregate["skipped_sources"]) == len(selected):
        print("No selected source is configured. See job-scout/README.md.", file=sys.stderr)
        return 2
    if not successful_sources:
        print("All configured sources failed. Review source_errors above.", file=sys.stderr)
        return 1
    if daily_mode:
        new_jobs = [store.get(job_id) for job_id in aggregate["new_job_ids"]]
        sync_tracker(
            store, args.tracker, jobs=[job for job in new_jobs if job], append_only=True,
        )
    else:
        sync_tracker(store, args.tracker)
    return 0


def diagnose_remotive(args, store, preferences):
    """Exercise only Remotive and report each stage without changing saved jobs."""
    provider = RemotiveProvider()
    queries = args.query or [
        "digital marketing", "paid search", "PPC", "marketing analytics",
    ]
    unique: dict[str, RawListing] = {}
    try:
        for query in queries:
            request = SearchRequest(query=query, location="Remote", results_per_page=args.results)
            for raw in provider.search(request):
                key = raw.source_job_id or raw.url
                unique.setdefault(key, raw)
    except ProviderError as error:
        print(f"Error: remotive provider failed: {error}", file=sys.stderr)
        return 1

    resume_text = extract_resume_text(MASTER_RESUME)
    threshold = int(preferences["minimum_score"] if args.minimum_score is None else args.minimum_score)
    candidates = []
    for raw in unique.values():
        if not raw.source_job_id or not raw.url or not raw.title or not raw.description:
            continue
        job = normalize(raw)
        result = score_job(job, preferences, resume_text)
        job.match_score = result.total
        job.evidence_confidence = result.confidence
        job.provisional = result.provisional
        candidates.append(job)

    matched = [job for job in candidates if job.match_score >= threshold]
    other_source_jobs = [job for job in store.all() if job.source.casefold() != "remotive"]
    survivors = [job for job in matched if find_duplicate(job, other_source_jobs) is None]
    examples = sorted(candidates, key=lambda job: (-job.match_score, job.title.casefold()))[:args.examples]
    report = {
        "provider": "remotive",
        "endpoint": getattr(provider, "api_endpoint", provider.endpoint),
        "queries": queries,
        "minimum_score": threshold,
        "raw_retrieved": provider.raw_count,
        "query_matches": len(unique),
        "matched_gecko_filters": len(matched),
        "survived_deduplication": len(survivors),
        "source_counts": {"remotive": len(unique)},
        "examples": [
            {
                "title": job.title,
                "company": job.company,
                "url": job.url,
                "match_score": job.match_score,
            }
            for job in examples
        ],
    }
    print(json.dumps(report, indent=2))
    return 0


def diagnose_remotive_feeds(args, _store, preferences):
    """Dry-run the bounded RSS import and report acquisition/scoring diagnostics."""
    provider = RemotiveProvider()
    resume_text = extract_resume_text(MASTER_RESUME)
    with tempfile.TemporaryDirectory() as directory:
        with JobStore(Path(directory) / "remotive-diagnostic.sqlite3") as diagnostic_store:
            summary = discover_remotive_full_feed(
                provider, diagnostic_store, preferences, resume_text,
                getattr(args, "limit", 100),
            )
    report = {
        "provider": "remotive",
        "discovery_location": "Remote",
        "feeds_attempted": summary.successful_feeds + summary.failed_feeds,
        "category_feeds": provider.feed_results,
        "successful_category_feeds": provider.successful_feed_count,
        "failed_category_feeds": provider.failed_feed_count,
        "raw_rss_records": provider.rss_raw_count,
        "unique_rss_jobs": summary.unique_available,
        "rss_duplicates_removed": provider.rss_duplicate_count,
        "unique_jobs_imported": summary.unique_imported,
        "scores_80_plus": summary.score_80_plus,
        "scores_70_79": summary.score_70_79,
        "scores_below_70": summary.score_below_70,
    }
    print(json.dumps(report, indent=2))
    return 0 if summary.unique_available else 1


def diagnose_remotive_rss(args, store, preferences):
    """Backward-compatible alias for the category-feed diagnostic."""
    return diagnose_remotive_feeds(args, store, preferences)


def list_jobs(args, store, preferences):
    jobs = [job for job in store.all(retained_only=not args.all)
            if job.match_score >= args.minimum_score and (not args.status or job.status == args.status)]
    jobs.sort(key=lambda job: (-job.match_score, job.date_posted or "9999", job.company.lower()))
    if args.json:
        print(json.dumps([job.to_dict() for job in jobs], indent=2))
    elif not jobs:
        print("No matching saved jobs.")
    else:
        print(f"{'ID':>4}  {'Score':>5}  {'Confidence':>10}  {'Evidence':<12}  {'Status':<14}  Company - Title")
        for job in jobs:
            evidence = "provisional" if job.provisional else "confirmed"
            print(f"{job.id:>4}  {job.match_score:>3}/100  {job.evidence_confidence:>9}%  {evidence:<12}  {job.status:<14}  {job.company} - {job.title}")
    return 0


def show(args, store, _preferences):
    print(json.dumps(get_job(store, args.job_id).to_dict(), indent=2))
    return 0


def select(args, store, _preferences):
    job = get_job(store, args.job_id)
    path = archive_listing(job, PROJECT_ROOT)
    store.update_status(args.job_id, "selected")
    sync_tracker(store, args.tracker)
    print(f"Selected Job Scout #{args.job_id}. Archived listing: {path.relative_to(PROJECT_ROOT)}")
    print(f"Next: use Gecko for the job description at {path.relative_to(PROJECT_ROOT)}")
    print("No resume was generated and no application-tracker row was added.")
    return 0


def set_status(args, store, _preferences):
    store.update_status(args.job_id, args.status)
    sync_tracker(store, args.tracker)
    print(f"Job {args.job_id} status set to {args.status}")
    return 0


def rescore(args, store, preferences):
    """Recalculate saved listings without refetching or changing lifecycle status."""
    resume_text = extract_resume_text(MASTER_RESUME)
    threshold = int(preferences["minimum_score"] if args.minimum_score is None else args.minimum_score)
    changes = []
    for job in store.all():
        old_score = job.match_score
        scoring_job = replace(job, description=job.enriched_description) if job.enriched_description else job
        result = score_job(scoring_job, preferences, resume_text)
        job.match_score = result.total
        job.evidence_confidence = result.confidence
        job.provisional = result.provisional
        job.evidence_levels = result.evidence_levels
        job.match_strengths = result.strengths
        job.match_weaknesses = result.weaknesses
        store.update_scoring(job, retained=job.match_score >= threshold)
        changes.append({
            "id": job.id, "old_score": old_score, "new_score": job.match_score,
            "score_change": job.match_score - old_score,
            "evidence_confidence": job.evidence_confidence,
            "provisional": job.provisional,
        })
    changes.sort(key=lambda item: (-item["new_score"], item["id"]))
    print(json.dumps(changes, indent=2))
    sync_tracker(store, args.tracker)
    return 0


def review_jobs(args, store, _preferences):
    """Display a read-only, prioritized application review queue."""
    queue = build_review_queue(
        store.all(),
        minimum_score=args.minimum_score,
        limit=args.limit,
        confirmed_only=args.confirmed_only,
        include_closed=args.include_closed,
    )
    print(queue_to_json(queue) if args.json else format_review_queue(queue))
    return 0


def enrich_jobs(args, store, preferences):
    """Enrich one or all provisional 80+ Adzuna records and report before/after evidence."""
    if bool(args.job_id) == bool(args.provisional):
        raise ValueError("Specify either a job ID or --provisional")
    if args.provisional:
        targets = [job for job in store.all()
                   if job.source.lower() == "adzuna" and job.match_score >= 80 and job.provisional]
    else:
        job = get_job(store, args.job_id)
        if job.source.lower() != "adzuna" or job.match_score < 80 or not job.provisional:
            raise ValueError("Enrichment is limited to provisional Adzuna jobs scoring 80 or higher")
        targets = [job]
    resume_text = extract_resume_text(MASTER_RESUME)
    report = []
    for job in targets:
        job = enrich_and_rescore(job, store, preferences, resume_text)
        if job.match_score >= 80 and job.evidence_confidence >= 65:
            final_status = "confirmed strong match"
        elif job.match_score >= 80:
            final_status = "80+ provisional — full description recommended"
        else:
            final_status = "below threshold after enrichment"
        report.append({
            "id": job.id,
            "title": job.title,
            "company": job.company,
            "original_score": job.original_match_score,
            "enriched_score": job.enriched_match_score or None,
            "original_confidence": job.original_evidence_confidence,
            "enriched_confidence": job.enriched_evidence_confidence or None,
            "enrichment_status": job.enrichment_status,
            "enrichment_source": job.enriched_source_url,
            "enrichment_error": job.enrichment_error,
            "final_status": final_status,
        })
    print(json.dumps(report, indent=2))
    sync_tracker(store, args.tracker)
    return 0


def resolve_urls(args, store, _preferences):
    """Resolve one job or a bounded set of Jooble-sourced jobs."""
    if bool(args.job_id) == bool(args.jooble):
        raise ValueError("Specify either a job ID or --jooble")
    if args.jooble:
        targets = [
            job for job in store.all()
            if any(link.get("source", "").casefold() == "jooble" for link in job.source_links)
            and (args.force or job.url_verification_status == "not_attempted")
        ][:args.limit]
    else:
        targets = [get_job(store, args.job_id)]
    report = []
    for job in targets:
        original_score = job.match_score
        resolved = resolve_and_store(job, store)
        report.append({
            "id": resolved.id,
            "company": resolved.company,
            "title": resolved.title,
            "url_status": url_status_label(resolved),
            "authoritative_url": resolved.authoritative_url,
            "confidence": resolved.authoritative_url_confidence,
            "redirect_url": resolved.url_redirect_url,
            "match_score": resolved.match_score,
            "match_score_unchanged": resolved.match_score == original_score,
            "error": resolved.url_resolution_error,
        })
    print(json.dumps(report, indent=2))
    sync_tracker(store, args.tracker)
    return 0


def daily(args, store, preferences):
    """Run the complete discovery-to-review workflow without creating resumes."""
    validator = DailyLinkValidator()
    if isinstance(store, JobStore):
        if tracker_backend() == "google-sheets":
            if not Path(args.tracker).is_file():
                raise RuntimeError("XLSX backup is required to verify protected Google Sheets history before cleanup")
            # Read cloud manual fields into the backup before deciding what is safe to remove.
            sync_workbook_to_google(args.tracker)
        existing = [job for job in store.all() if job.status not in {"ignored", "rejected"}]
        existing.extend(sheet_only_active_jobs(args.tracker, {job.id for job in existing if job.id is not None}))
        protected = protected_job_ids(existing, args.tracker)
        checked = validator.check_existing(existing)
        dead = [(job, result) for job, result in checked if result.status == "dead"]
        removable = [(job, result) for job, result in dead if job.id not in protected]
        validator.protected_dead += len(dead) - len(removable)
        if removable:
            ids = {job.id for job, _ in removable}
            sync_tracker(store, args.tracker, jobs=[], append_only=True, remove_scout_ids=ids)
            uncleared = ids & scout_row_ids(args.tracker)
            if uncleared:
                raise RuntimeError(f"Confirmed-dead Scout rows were not cleared: {sorted(uncleared)}")
            for job, result in removable:
                store.delete_dead_unprotected(job.id)
                validator.record_removed(job, result)
    search_args = argparse.Namespace(
        source="core", query=None, location=None, page=1, results=args.results,
        minimum_score=None, tracker=args.tracker, limit=100, daily_mode=True,
        link_validator=validator,
    )
    result = search(search_args, store, preferences)
    print("Link validation:")
    for field in ("checked", "valid", "removed_dead", "temporary_failure", "protected_dead"):
        print(f"  {field}: {validator.report()[field]}")
    for removed in validator.removed:
        print(f"  removed: {removed['company']} | {removed['job_title']} | {removed['url']} | "
              f"{removed['reason_removed']} | HTTP/status {removed['http_status'] or 'n/a'}")
    if result:
        return result
    run_result = getattr(search_args, "run_result", {})
    new_jobs = [store.get(job_id) for job_id in run_result.get("new_job_ids", [])]
    queue = build_review_queue(
        [job for job in new_jobs if job], minimum_score=args.minimum_score,
        limit=args.limit, confirmed_only=False, include_closed=False,
    )
    qualifying = sum(len(jobs) for jobs in queue.values())
    if not qualifying:
        print("DAILY REVIEW: No new qualifying jobs were discovered in this run.")
    else:
        print(f"DAILY REVIEW: {qualifying} new qualifying job(s) discovered in this run.")
        print(format_review_queue(queue))
    return 0


def build_parser():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--db", default=str(DEFAULT_DB), help="SQLite database path")
    parser.add_argument("--tracker", default=str(DEFAULT_TRACKER), help="Tracker workbook path")
    parser.add_argument("--preferences", help="Alternate preferences JSON")
    sub = parser.add_subparsers(dest="command", required=True)
    find = sub.add_parser("search", help="Search configured providers and save scored results")
    find.add_argument(
        "--source", "--provider",
        choices=["all", "core", "adzuna", "jooble", "remotive", "indeed", "web-careers"],
        default=DEFAULT_SOURCE,
        help=(f"Provider to query (default: {DEFAULT_SOURCE}); core is Adzuna, Jooble, "
              "Remotive, and Web Careers (Brave when configured)"),
    )
    find.add_argument("--query", action="append", help="Repeat for multiple role queries; defaults to all target roles")
    find.add_argument(
        "--location", action="append",
        help="Repeat for non-Remotive providers; defaults to Utah and Remote. Remotive is always Remote.",
    )
    find.add_argument("--page", type=int, default=1)
    find.add_argument("--results", type=int, default=20)
    find.add_argument(
        "--limit", type=int, default=100,
        help="Unique Remotive jobs to import after within-provider deduplication (default: 100)",
    )
    find.add_argument("--minimum-score", type=int)
    find.set_defaults(function=search)
    diagnostic = sub.add_parser(
        "diagnose-remotive",
        help="Run a read-only Remotive-only API, filter, and deduplication diagnostic",
    )
    diagnostic.add_argument(
        "--query", action="append",
        help="Repeat to override the digital marketing / paid search / PPC / analytics queries",
    )
    diagnostic.add_argument(
        "--results", type=int, default=10_000,
        help="Maximum matching records considered per query (default: 10000)",
    )
    diagnostic.add_argument("--minimum-score", type=int)
    diagnostic.add_argument("--examples", type=int, default=5)
    diagnostic.set_defaults(function=diagnose_remotive)
    rss_diagnostic = sub.add_parser(
        "diagnose-remotive-rss",
        help="Backward-compatible alias for the Remotive category-feed diagnostic",
    )
    rss_diagnostic.add_argument("--limit", type=int, default=100)
    rss_diagnostic.set_defaults(function=diagnose_remotive_rss)
    feeds_diagnostic = sub.add_parser(
        "diagnose-remotive-feeds",
        help="Dry-run Remotive RSS acquisition, deduplication, and local scoring",
    )
    feeds_diagnostic.add_argument("--limit", type=int, default=100)
    feeds_diagnostic.set_defaults(function=diagnose_remotive_feeds)
    listing = sub.add_parser("list", help="Show strong saved matches")
    listing.add_argument("--minimum-score", type=int, default=80)
    listing.add_argument("--status", choices=["new", "reviewing", "selected", "resume-created", "applied", "contacted", "interview", "rejected", "offer", "ignored"])
    listing.add_argument("--all", action="store_true", help="Include weak results remembered as seen")
    listing.add_argument("--json", action="store_true")
    listing.set_defaults(function=list_jobs)
    detail = sub.add_parser("show", help="Show a complete saved record")
    detail.add_argument("job_id", type=int)
    detail.set_defaults(function=show)
    choose = sub.add_parser("select", help="Archive one job for the existing Gecko workflow")
    choose.add_argument("job_id", type=int)
    choose.set_defaults(function=select)
    status = sub.add_parser("status", help="Change a saved job's lifecycle status")
    status.add_argument("job_id", type=int)
    status.add_argument("status", choices=["new", "reviewing", "selected", "resume-created", "applied", "contacted", "interview", "rejected", "offer", "ignored"])
    status.set_defaults(function=set_status)
    rescore_parser = sub.add_parser("rescore", help="Recalculate all saved jobs with the current scoring model")
    rescore_parser.add_argument("--minimum-score", type=int)
    rescore_parser.set_defaults(function=rescore)
    enrich_parser = sub.add_parser("enrich", help="Enrich provisional 80+ Adzuna matches")
    enrich_parser.add_argument("job_id", nargs="?", type=int)
    enrich_parser.add_argument("--provisional", action="store_true", help="Enrich all provisional 80+ matches")
    enrich_parser.set_defaults(function=enrich_jobs)
    resolve_parser = sub.add_parser(
        "resolve-url", help="Verify provider links and find authoritative employer/ATS postings",
    )
    resolve_parser.add_argument("job_id", nargs="?", type=int)
    resolve_parser.add_argument("--jooble", action="store_true", help="Resolve unresolved Jooble-sourced jobs")
    resolve_parser.add_argument("--limit", type=int, default=20)
    resolve_parser.add_argument("--force", action="store_true", help="Retry jobs with an existing URL status")
    resolve_parser.set_defaults(function=resolve_urls)
    review_parser = sub.add_parser("review", help="Show the prioritized, read-only Job Scout review queue")
    review_parser.add_argument("--confirmed-only", action="store_true")
    review_parser.add_argument("--minimum-score", type=int, default=70)
    review_parser.add_argument("--limit", type=int, default=20)
    review_parser.add_argument("--json", action="store_true")
    review_parser.add_argument(
        "--include-closed", action="store_true",
        help="Include applied, rejected, and ignored jobs",
    )
    review_parser.set_defaults(function=review_jobs)
    daily_parser = sub.add_parser("daily", help="Search, deduplicate, score, enrich, sync, and review")
    daily_parser.add_argument("--results", type=int, default=20)
    daily_parser.add_argument("--minimum-score", type=int, default=70)
    daily_parser.add_argument("--limit", type=int, default=20)
    daily_parser.set_defaults(function=daily)
    sync_parser = sub.add_parser(
        "sync-sheets",
        help="Synchronize both tracker worksheets to Google Sheets and retain the XLSX backup",
    )
    sync_parser.set_defaults(function=sync_sheets)
    return parser


def main():
    args = build_parser().parse_args()
    try:
        load_local_environment()
        load_local_environment(PROJECT_ROOT / ".env.google-sheets.local")
        preferences = load_preferences(args.preferences)
        with JobStore(args.db) as store:
            return args.function(args, store, preferences)
    except (OSError, RuntimeError, ValueError, ProviderError) as error:
        print(f"Error: {error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
