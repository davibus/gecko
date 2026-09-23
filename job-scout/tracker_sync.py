"""Synchronize Job Scout history into an isolated worksheet in Gecko's tracker."""

from __future__ import annotations

from copy import copy
import os
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from tempfile import NamedTemporaryFile

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from models import JobListing
from normalize import canonicalize_url
from url_resolution import best_job_url, url_status_label
from excel_preservation import append_preserving_format, copy_row_format


SHEET_NAME = "Job Scout"
APPLICATION_SHEET = "Job Tracker"
LEGACY_HEADERS = [
    "Scout ID", "Company", "Job Title", "Match Score", "Evidence Confidence",
    "Match Status", "Location", "Work Arrangement", "Employment Type", "Salary",
    "Source", "Date Posted", "Date Found", "Last Seen", "Top Strengths",
    "Top Weaknesses", "Job URL", "Enrichment URL", "Gecko Status",
    "Resume Created", "Applied", "Contacted",
]
SOURCE_B_HEADERS = [
    "Scout ID", "Source", "Company", "Job Title", "Match Score", "Evidence Confidence",
    "Match Status", "Location", "Work Arrangement", "Employment Type", "Salary",
    "Date Posted", "Date Found", "Last Seen", "Top Strengths", "Top Weaknesses",
    "Job URL", "Enrichment URL", "Gecko Status", "Resume Created", "Applied", "Contacted",
]
RESUME_LINK_HEADERS = SOURCE_B_HEADERS[:20] + ["Resume Link"] + SOURCE_B_HEADERS[20:]
PREVIOUS_HEADERS = [header for header in RESUME_LINK_HEADERS if header not in {"Top Strengths", "Top Weaknesses"}]
URL_STATUSLESS_HEADERS = [
    header for header in PREVIOUS_HEADERS
    if header not in {"Applied", "Resume Link", "Contacted"}
] + ["Applied", "Resume Link", "Contacted"]
PRE_APPLY_HEADERS = (
    URL_STATUSLESS_HEADERS[:14]
    + ["URL Status", "Authoritative URL"]
    + URL_STATUSLESS_HEADERS[14:]
)
HEADERS = PRE_APPLY_HEADERS[:5] + ["Apply?"] + PRE_APPLY_HEADERS[5:]
STATUS_LABELS = {
    "new": "New", "reviewing": "Reviewing", "selected": "Selected",
    "resume-created": "Resume Created", "applied": "Applied", "contacted": "Contacted",
    "interview": "Interview", "rejected": "Rejected", "ignored": "Ignored", "offer": "Offer",
}
STATUS_RANK = {
    "New": 0, "Reviewing": 1, "Selected": 2, "Resume Created": 3,
    "Applied": 4, "Contacted": 5, "Interview": 6, "Rejected": 7,
    "Ignored": 7, "Offer": 8,
}
COLUMN_WIDTHS = [
    10, 12, 28, 44, 13, 12, 19, 24, 28, 18, 18, 20, 14, 14, 14,
    32, 48, 42, 42, 18, 16, 12, 48, 12,
]


@dataclass
class SyncSummary:
    rows: int
    confirmed: int
    provisional: int
    near_matches: int
    tracker_preserved: bool


def _parse_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not value:
        return None
    try:
        return datetime.fromisoformat(str(value).replace("Z", "+00:00")).date()
    except ValueError:
        return None


def _identity(value: str) -> str:
    return canonicalize_url(str(value or "")).lower()


def match_status(job: JobListing) -> str:
    if job.match_score >= 80 and job.evidence_confidence >= 65 and not job.provisional:
        return "Confirmed Strong Match"
    if job.match_score >= 80:
        return "Provisional 80+"
    if job.match_score >= 70:
        return "Near Match"
    return "Below Threshold"


def _merge_lifecycle(existing: dict, job: JobListing) -> str:
    database_status = STATUS_LABELS.get(job.status, job.status.replace("-", " ").title())
    candidates = [database_status, str(existing.get("Gecko Status") or "New")]
    if existing.get("Resume Created"):
        candidates.append("Resume Created")
    if existing.get("Applied"):
        candidates.append("Applied")
    if existing.get("Contacted"):
        candidates.append("Contacted")
    return max(candidates, key=lambda value: STATUS_RANK.get(value, 0))


def _join(values: list[str]) -> str:
    return "\n".join(f"• {value}" for value in values[:3])


def _job_row(job: JobListing, existing: dict | None = None) -> dict:
    existing = existing or {}
    lifecycle = _merge_lifecycle(existing, job)
    resume_created = existing.get("Resume Created") or ("X" if STATUS_RANK.get(lifecycle, 0) >= 3 else None)
    applied = existing.get("Applied")
    contacted = existing.get("Contacted")
    if not applied and job.status in {"applied", "contacted", "interview", "offer"}:
        applied = "X"
    if not contacted and job.status in {"contacted", "interview", "offer"}:
        contacted = "X"
    return {
        "Scout ID": job.id,
        "Company": job.company,
        "Job Title": job.title,
        "Match Score": job.match_score,
        "Apply?": existing.get("Apply?"),
        "Evidence Confidence": job.evidence_confidence / 100,
        "Match Status": match_status(job),
        "Location": job.location,
        "Work Arrangement": job.work_arrangement,
        "Employment Type": job.employment_type.lower().replace("_", "-") if job.employment_type else None,
        "Salary": job.salary or None,
        "Source": job.source,
        "Date Posted": _parse_date(job.date_posted),
        "Date Found": _parse_date(job.date_discovered),
        "Last Seen": _parse_date(job.last_seen or job.date_discovered),
        "URL Status": url_status_label(job),
        "Authoritative URL": job.authoritative_url or None,
        "Job URL": best_job_url(job),
        "Enrichment URL": job.enriched_source_url or None,
        "Gecko Status": lifecycle,
        "Resume Created": resume_created,
        "Resume Link": existing.get("Resume Link"),
        "Applied": applied,
        "Contacted": contacted,
    }


def _header_columns(ws) -> dict[str, int]:
    headers = [ws.cell(1, column).value for column in range(1, ws.max_column + 1)]
    while headers and headers[-1] in (None, ""):
        headers.pop()
    if len(headers) != len(set(headers)):
        raise ValueError("Existing Job Scout worksheet contains duplicate column names")
    columns = {str(header): index for index, header in enumerate(headers, 1) if header not in (None, "")}
    for header in HEADERS:
        if header in columns:
            continue
        column = ws.max_column + 1
        if ws.max_column >= 1:
            source = ws.cell(1, ws.max_column)
            target = ws.cell(1, column)
            if source.has_style:
                target._style = copy(source._style)
        ws.cell(1, column, header)
        columns[header] = column
    return columns


def _existing_rows(ws, columns: dict[str, int] | None = None) -> list[dict]:
    if ws.max_row < 2:
        return []
    columns = columns or _header_columns(ws)
    return [
        {
            **{header: ws.cell(row, columns[header]).value for header in HEADERS},
            "_worksheet_row": row,
        }
        for row in range(2, ws.max_row + 1)
        if any(ws.cell(row, columns[header]).value not in (None, "") for header in HEADERS)
    ]


def _migrate_schema(ws) -> None:
    """Add missing managed columns without moving or recreating existing cells."""
    _header_columns(ws)


def _sort_key(record: dict):
    confidence = record.get("Evidence Confidence") or 0
    if confidence > 1:
        confidence /= 100
    posted = _parse_date(record.get("Date Posted"))
    return (
        -(record.get("Match Score") or 0),
        -confidence,
        -(posted.toordinal() if posted else 0),
        str(record.get("Company") or "").lower(),
    )


def _style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=False)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    end_column = get_column_letter(len(HEADERS))
    ws.auto_filter.ref = f"A1:{end_column}{max(ws.max_row, 1)}"
    ws.sheet_properties.tabColor = "5B9BD5"
    for index, width in enumerate(COLUMN_WIDTHS, 1):
        ws.column_dimensions[ws.cell(1, index).column_letter].width = width
    for row in range(2, ws.max_row + 1):
        ws.cell(row, HEADERS.index("Match Score") + 1).number_format = "0"
        ws.cell(row, HEADERS.index("Evidence Confidence") + 1).number_format = "0%"
        for header in ("Date Posted", "Date Found", "Last Seen"):
            column = HEADERS.index(header) + 1
            ws.cell(row, column).number_format = "mm/dd/yyyy"
        for column in range(1, len(HEADERS) + 1):
            ws.cell(row, column).alignment = Alignment(vertical="top", wrap_text=False)
        for header in ("Authoritative URL", "Job URL", "Enrichment URL", "Resume Link"):
            column = HEADERS.index(header) + 1
            cell = ws.cell(row, column)
            if cell.value:
                cell.hyperlink = str(cell.value)
                cell.style = "Hyperlink"
    end = max(ws.max_row, 2)
    match_column = get_column_letter(HEADERS.index("Match Status") + 1)
    status_column = get_column_letter(HEADERS.index("Gecko Status") + 1)
    applied_column = get_column_letter(HEADERS.index("Applied") + 1)
    ws.conditional_formatting.add(
        f"A2:{end_column}{end}",
        FormulaRule(formula=[f'${match_column}2="Confirmed Strong Match"'], fill=PatternFill("solid", fgColor="E2F0D9")),
    )
    ws.conditional_formatting.add(
        f"A2:{end_column}{end}",
        FormulaRule(formula=[f'${match_column}2="Near Match"'], fill=PatternFill("solid", fgColor="FFF2CC")),
    )
    ws.conditional_formatting.add(
        f"A2:{end_column}{end}",
        FormulaRule(
            formula=[f'OR(${status_column}2="Ignored",${status_column}2="Rejected",${status_column}2="Applied",${applied_column}2<>"")'],
            fill=PatternFill("solid", fgColor="E7E6E6"), stopIfTrue=True,
        ),
    )


def _save_atomic(workbook, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with NamedTemporaryFile(prefix="job-scout-sync-", suffix=".xlsx", dir=path.parent, delete=False) as temp:
        temp_path = Path(temp.name)
    try:
        workbook.save(temp_path)
        os.replace(temp_path, path)
    finally:
        if temp_path.exists():
            temp_path.unlink()


def sync_job_scout(
    jobs: list[JobListing], tracker_path: str | Path, *, append_only: bool = False,
    remove_scout_ids: set[int] | None = None,
) -> SyncSummary:
    """Sync Scout jobs while preserving workbook history and manual fields.

    ``append_only`` is used by the daily workflow: jobs that already have a
    worksheet identity are left byte-for-byte equivalent at the cell-value
    level, while genuinely new identities are appended.
    """
    tracker_path = Path(tracker_path)
    workbook_exists = tracker_path.exists()
    if tracker_path.exists():
        workbook = load_workbook(tracker_path)
    else:
        workbook = Workbook()
        workbook.active.title = APPLICATION_SHEET
    tracker_preserved = APPLICATION_SHEET in workbook.sheetnames
    if not tracker_preserved:
        raise ValueError(f"Workbook must contain the existing {APPLICATION_SHEET!r} worksheet")
    sheet_exists = SHEET_NAME in workbook.sheetnames
    ws = workbook[SHEET_NAME] if sheet_exists else workbook.create_sheet(SHEET_NAME)
    if ws.max_row == 1 and ws.cell(1, 1).value is None:
        for column, header in enumerate(HEADERS, 1):
            ws.cell(1, column, header)

    headers_before = [ws.cell(1, column).value for column in range(1, ws.max_column + 1)]
    _migrate_schema(ws)

    columns = _header_columns(ws)
    headers_after = [ws.cell(1, column).value for column in range(1, ws.max_column + 1)]
    changed = not sheet_exists or headers_before != headers_after
    rows = _existing_rows(ws, columns)
    if remove_scout_ids:
        removable = {str(job_id) for job_id in remove_scout_ids}
        kept = []
        for row in rows:
            if str(row.get("Scout ID")) not in removable:
                kept.append(row)
                continue
            if any(row.get(field) not in (None, "") for field in ("Apply?", "Applied", "Contacted", "Resume Created")) or str(row.get("Gecko Status") or "").lower() not in {"", "new", "reviewing"}:
                kept.append(row)
                continue
            row_number = row["_worksheet_row"]
            for column in range(1, ws.max_column + 1):
                cell = ws.cell(row_number, column)
                if cell.value is not None or cell.hyperlink is not None:
                    cell.value = None
                    cell.hyperlink = None
                    changed = True
        rows = kept
    by_id = {str(row["Scout ID"]): row for row in rows if row.get("Scout ID") not in (None, "")}
    by_identity = {}
    for row in rows:
        for column in ("Job URL", "Enrichment URL"):
            identity = _identity(row.get(column))
            if identity:
                by_identity[identity] = row

    for job in jobs:
        existing = by_id.get(str(job.id))
        if existing is None:
            for url in (job.canonical_url, job.url, job.enriched_source_url):
                existing = by_identity.get(_identity(url))
                if existing is not None:
                    break
        updated = _job_row(job, existing)
        if existing is None:
            row_number = append_preserving_format(ws, set(HEADERS))
            changed = True
            updated["_worksheet_row"] = row_number
            for header, value in updated.items():
                if header in columns:
                    ws.cell(row_number, columns[header], value)
            rows.append(updated)
            existing = updated
        elif not append_only:
            row_number = existing["_worksheet_row"]
            for header, value in updated.items():
                if header in columns and ws.cell(row_number, columns[header]).value != value:
                    ws.cell(row_number, columns[header], value)
                    changed = True
            existing.update(updated)
        row_number = existing["_worksheet_row"]
        for header in ("Authoritative URL", "Job URL", "Enrichment URL", "Resume Link"):
            cell = ws.cell(row_number, columns[header])
            target = str(cell.value) if cell.value else None
            current = cell.hyperlink.target if cell.hyperlink else None
            if current != target:
                cell.hyperlink = target
                changed = True
    if not workbook_exists or not sheet_exists:
        _style_sheet(ws)
        changed = True
    if changed:
        _save_atomic(workbook, tracker_path)

    return SyncSummary(
        rows=len(rows),
        confirmed=sum(row.get("Match Status") == "Confirmed Strong Match" for row in rows),
        provisional=sum(row.get("Match Status") == "Provisional 80+" for row in rows),
        near_matches=sum(row.get("Match Status") == "Near Match" for row in rows),
        tracker_preserved=tracker_preserved,
    )
