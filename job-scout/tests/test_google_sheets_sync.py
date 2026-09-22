from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table


SCOUT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SCOUT_ROOT))

from google_sheets_sync import (
    APPLICATION_HEADERS,
    GoogleSheetsConfig,
    _records_from_values,
    merge_application_rows,
    merge_scout_rows,
    sync_workbook_to_google,
)
from tracker_sync import HEADERS as SCOUT_HEADERS


def application_record(job_number="job-1", applied="", contacted=""):
    return {
        "Resume #": 1, "Company": "Example", "Job Title": "Manager", "Pay": "$100,000",
        "Job Number": job_number, "Match Score": "90/100", "Job Link": "https://example.test/job/1",
        "Resume Link": "resumes/example.docx", "Date Created": "09/19/2026",
        "Applied": applied, "Contacted": contacted, "Source": "Test",
        "Date Found": "09/18/2026", "Status": "resume-created",
    }


def scout_record(scout_id=1, score=90, status="New", applied="", contacted="", url=None):
    record = {header: "" for header in SCOUT_HEADERS}
    record.update({
        "Scout ID": scout_id, "Company": "Example", "Job Title": "Manager",
        "Match Score": score, "Evidence Confidence": 0.9,
        "Match Status": "Confirmed Strong Match", "Location": "Remote",
        "Source": "test", "Date Posted": "09/18/2026", "Date Found": "09/18/2026",
        "Last Seen": "09/19/2026", "Top Strengths": "Strong", "Top Weaknesses": "Gap",
        "Job URL": url or "https://example.test/job/1", "Gecko Status": status,
        "Resume Created": "X" if status == "Resume Created" else "",
        "Applied": applied, "Contacted": contacted,
    })
    return record


def values(headers, records):
    return [headers] + [[record.get(header, "") for header in headers] for record in records]


class FakeSheetsApi:
    """Small mocked Google API surface used to verify requests and responses."""

    def __init__(self, application_rows=None, scout_rows=None):
        self.remote = {
            "Job Tracker": values(APPLICATION_HEADERS, application_rows or []),
            "Job Scout": values(SCOUT_HEADERS, scout_rows or []),
        }
        self.writes = {}
        self.clears = []
        self.batch_requests = []

    def service(self):
        outer = MagicMock()
        spreadsheets = outer.spreadsheets.return_value
        values_api = spreadsheets.values.return_value

        metadata = {
            "sheets": [
                {
                    "properties": {"sheetId": 1, "title": "Job Tracker", "index": 0, "gridProperties": {}},
                    "basicFilter": {
                        "range": {"sheetId": 1, "startRowIndex": 0, "endRowIndex": len(self.remote["Job Tracker"]), "startColumnIndex": 0, "endColumnIndex": len(self.remote["Job Tracker"][0])},
                        "criteria": {"9": {"hiddenValues": [""]}},
                    },
                    "conditionalFormats": [{"booleanRule": {}}],
                },
                {
                    "properties": {"sheetId": 2, "title": "Job Scout", "index": 1, "gridProperties": {}},
                    "basicFilter": {
                        "range": {"sheetId": 2, "startRowIndex": 0, "endRowIndex": len(self.remote["Job Scout"]), "startColumnIndex": 0, "endColumnIndex": len(self.remote["Job Scout"][0])},
                        "criteria": {"4": {"condition": {"type": "NUMBER_GREATER", "values": [{"userEnteredValue": "80"}]}}},
                    },
                    "conditionalFormats": [{"booleanRule": {}}],
                },
            ]
        }
        spreadsheets.get.side_effect = lambda **_kwargs: self._request(metadata)
        spreadsheets.batchUpdate.side_effect = self._batch
        values_api.get.side_effect = self._get
        values_api.clear.side_effect = self._clear
        values_api.update.side_effect = self._update
        return outer

    @staticmethod
    def _request(payload):
        request = MagicMock()
        request.execute.return_value = payload
        return request

    @staticmethod
    def _title(range_value):
        return range_value.split("!", 1)[0].strip("'").replace("''", "'")

    def _get(self, **kwargs):
        return self._request({"values": self.remote.get(self._title(kwargs["range"]), [])})

    def _clear(self, **kwargs):
        self.clears.append(self._title(kwargs["range"]))
        return self._request({})

    def _update(self, **kwargs):
        self.writes[self._title(kwargs["range"])] = kwargs["body"]["values"]
        return self._request({"updatedRows": len(kwargs["body"]["values"])})

    def _batch(self, **kwargs):
        self.batch_requests.extend(kwargs["body"]["requests"])
        return self._request({"replies": []})


class GoogleSheetsMergeTests(unittest.TestCase):
    def test_custom_and_reordered_scout_columns_are_preserved_by_header(self):
        row = scout_record()
        reordered = ["My Notes", *reversed(SCOUT_HEADERS)]
        source = values(reordered, [{**row, "My Notes": "keep me"}])
        records = _records_from_values(source, SCOUT_HEADERS, "Job Scout")
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]["Scout ID"], 1)
        self.assertEqual(records[0]["My Notes"], "keep me")

    def test_custom_columns_follow_rows_through_merge_and_sort(self):
        local = scout_record(1, score=96, url="https://example.test/job/1")
        remote = {**scout_record(1, score=80, url="https://example.test/job/1"), "My Notes": "keep me"}
        merged = merge_scout_rows([local], [remote])
        self.assertEqual(merged[0]["My Notes"], "keep me")

    def test_new_rows_append_and_history_is_preserved(self):
        merged = merge_scout_rows(
            [scout_record(2, score=95, url="https://example.test/job/2")],
            [scout_record(1, score=80, url="https://example.test/job/1")],
        )
        self.assertEqual([row["Scout ID"] for row in merged], [2, 1])

    def test_existing_job_updates_without_duplicate_by_canonical_url(self):
        local = scout_record(1, score=96, url="https://example.test/job/1")
        remote = scout_record(999, score=80, applied="X", url="https://example.test/job/1?utm_source=old")
        merged = merge_scout_rows([local], [remote])
        self.assertEqual(len(merged), 1)
        self.assertEqual(merged[0]["Match Score"], 96)
        self.assertEqual(merged[0]["Applied"], "X")

    def test_manual_application_fields_are_preserved(self):
        merged = merge_application_rows(
            [application_record()],
            [application_record(applied="X", contacted="manual note")],
        )
        self.assertEqual(len(merged), 1)
        self.assertEqual(merged[0]["Applied"], "X")
        self.assertEqual(merged[0]["Contacted"], "manual note")

    def test_select_and_resume_created_statuses_win_over_stale_remote_values(self):
        selected = merge_scout_rows([scout_record(status="Selected")], [scout_record(status="New")])
        self.assertEqual(selected[0]["Gecko Status"], "Selected")
        completed = merge_scout_rows([scout_record(status="Resume Created")], [scout_record(status="Selected")])
        self.assertEqual(completed[0]["Gecko Status"], "Resume Created")
        self.assertEqual(completed[0]["Resume Created"], "X")


class GoogleSheetsApiSyncTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.tracker = Path(self.temp.name) / "job-tracker.xlsx"
        workbook = Workbook()
        application = workbook.active
        application.title = "Job Tracker"
        application.append(APPLICATION_HEADERS)
        application.append([application_record().get(header) for header in APPLICATION_HEADERS])
        scout = workbook.create_sheet("Job Scout")
        scout.append(SCOUT_HEADERS)
        scout.append([scout_record(status="Resume Created").get(header) for header in SCOUT_HEADERS])
        workbook.save(self.tracker)
        self.config = GoogleSheetsConfig("spreadsheet-id", Path("unused.json"))

    def tearDown(self):
        self.temp.cleanup()

    def test_mocked_api_writes_both_tabs_and_keeps_xlsx_backup_valid(self):
        fake = FakeSheetsApi(
            application_rows=[application_record(applied="X")],
            scout_rows=[scout_record(status="Selected", applied="X", contacted="manual")],
        )
        summary = sync_workbook_to_google(
            self.tracker, config=self.config, service=fake.service()
        )
        self.assertEqual(summary.application_rows, 1)
        self.assertEqual(summary.scout_rows, 1)
        self.assertEqual(set(fake.writes), {"Job Tracker", "Job Scout"})
        self.assertEqual(fake.clears, [])
        self.assertTrue(fake.batch_requests)
        self.assertTrue(any("updateCells" in request for request in fake.batch_requests))
        self.assertFalse(any("repeatCell" in request for request in fake.batch_requests))
        self.assertFalse(any("deleteConditionalFormatRule" in request for request in fake.batch_requests))
        self.assertFalse(any("updateDimensionProperties" in request for request in fake.batch_requests))
        resume_links = [
            request["updateCells"]["rows"][0]["values"][0]["userEnteredFormat"]["textFormat"]["link"]["uri"]
            for request in fake.batch_requests if "updateCells" in request
        ]
        self.assertTrue(any(link.startswith("http://127.0.0.1:8765/open/") for link in resume_links))

        scout_write = fake.writes["Job Scout"]
        scout_row = dict(zip(scout_write[0], scout_write[1]))
        self.assertEqual(scout_row["Gecko Status"], "Resume Created")
        self.assertEqual(scout_row["Resume Created"], "X")
        self.assertEqual(scout_row["Applied"], "X")
        self.assertEqual(scout_row["Contacted"], "manual")

        backup = load_workbook(self.tracker)
        self.assertEqual(backup.sheetnames, ["Job Tracker", "Job Scout"])
        self.assertEqual(backup["Job Tracker"]["J2"].value, "X")
        scout_headers = {cell.value: cell.column for cell in backup["Job Scout"][1]}
        self.assertEqual(backup["Job Scout"].cell(2, scout_headers["Applied"]).value, "X")
        self.assertTrue(str(backup["Job Scout"].cell(2, scout_headers["Resume Link"]).value).startswith("file:///"))
        self.assertEqual(backup["Job Scout"].cell(2, scout_headers["Contacted"]).value, "manual")

    def test_api_sync_uses_remote_requests_without_a_google_sheet_file_lock(self):
        fake = FakeSheetsApi()
        sync_workbook_to_google(self.tracker, config=self.config, service=fake.service())
        self.assertIn("Job Tracker", fake.writes)
        self.assertIn("Job Scout", fake.writes)
        self.assertEqual(fake.clears, [])

    def test_api_sync_keeps_remote_column_order_custom_values_and_manual_fields(self):
        fake = FakeSheetsApi()
        application_headers = ["My App Notes", *reversed(APPLICATION_HEADERS)]
        scout_headers = ["My Scout Notes", *reversed(SCOUT_HEADERS)]
        fake.remote["Job Tracker"] = values(application_headers, [{
            **application_record(applied="X", contacted="manual"),
            "My App Notes": "retain application note",
        }])
        fake.remote["Job Scout"] = values(scout_headers, [{
            **scout_record(applied="X", contacted="manual"),
            "My Scout Notes": '=IF(A2<>"","retain scout note","")',
            "URL Status": "Verified - Official ATS",
            "Authoritative URL": "https://jobs.lever.co/example/123",
        }])

        sync_workbook_to_google(self.tracker, config=self.config, service=fake.service())

        self.assertEqual(fake.writes["Job Tracker"][0], application_headers)
        self.assertEqual(fake.writes["Job Scout"][0], scout_headers)
        application = dict(zip(application_headers, fake.writes["Job Tracker"][1]))
        scout = dict(zip(scout_headers, fake.writes["Job Scout"][1]))
        self.assertEqual(application["My App Notes"], "retain application note")
        self.assertEqual(application["Applied"], "X")
        self.assertEqual(application["Contacted"], "manual")
        self.assertEqual(scout["My Scout Notes"], '=IF(A2<>"","retain scout note","")')
        self.assertEqual(scout["Applied"], "X")
        self.assertEqual(scout["Contacted"], "manual")
        self.assertEqual(scout["URL Status"], "Verified - Official ATS")
        self.assertEqual(scout["Authoritative URL"], "https://jobs.lever.co/example/123")

    def test_google_sync_updates_xlsx_in_place_and_preserves_excel_formatting(self):
        workbook = load_workbook(self.tracker)
        application = workbook["Job Tracker"]
        scout = workbook["Job Scout"]
        application["C2"].fill = PatternFill("solid", fgColor="FFD966")
        application["C2"].font = Font(size=16, bold=True, color="5B9BD5")
        application.row_dimensions[2].height = 34
        application.column_dimensions["C"].width = 61
        application.freeze_panes = "D5"
        application.auto_filter.ref = "A1:N2"
        application.add_table(Table(displayName="Applications", ref="A1:N2"))
        app_validation = DataValidation(type="list", formula1='"Yes,No"')
        app_validation.add("J2")
        application.add_data_validation(app_validation)
        application["O1"] = "User Formula"
        application["O2"] = "=A2*10"

        scout["D2"].fill = PatternFill("solid", fgColor="C6E0B4")
        scout.row_dimensions[2].height = 29
        scout.column_dimensions["D"].width = 53
        scout.freeze_panes = "E6"
        scout.auto_filter.ref = f"A1:{scout.cell(1, scout.max_column).column_letter}2"
        scout.add_table(Table(displayName="ScoutRows", ref=f"A1:{scout.cell(1, scout.max_column).column_letter}2"))
        scout["X1"] = "User Formula"
        scout["X2"] = "=E2*2"
        workbook.save(self.tracker)

        fake = FakeSheetsApi(
            application_rows=[application_record("job-2")],
            scout_rows=[scout_record(2, url="https://example.test/job/2")],
        )
        sync_workbook_to_google(self.tracker, config=self.config, service=fake.service())

        workbook = load_workbook(self.tracker, data_only=False)
        application = workbook["Job Tracker"]
        scout = workbook["Job Scout"]
        self.assertEqual(application["C2"].fill.fgColor.rgb, "00FFD966")
        self.assertEqual(application["C2"].font.color.rgb, "005B9BD5")
        self.assertEqual(application["O2"].value, "=A2*10")
        self.assertEqual(application["O3"].value, "=A3*10")
        self.assertEqual(application.row_dimensions[2].height, 34)
        self.assertEqual(application.row_dimensions[3].height, 34)
        self.assertEqual(application.column_dimensions["C"].width, 61)
        self.assertEqual(application.freeze_panes, "D5")
        self.assertEqual(application.auto_filter.ref, "A1:N3")
        self.assertEqual(application.tables["Applications"].ref, "A1:N3")
        self.assertIn("J3", str(application.data_validations.dataValidation[0].sqref))
        self.assertEqual(scout["D2"].fill.fgColor.rgb, "00C6E0B4")
        self.assertEqual(scout["X2"].value, "=E2*2")
        self.assertEqual(scout["X3"].value, "=E3*2")
        self.assertEqual(scout.row_dimensions[3].height, 29)
        self.assertEqual(scout.column_dimensions["D"].width, 53)
        self.assertEqual(scout.freeze_panes, "E6")
        self.assertTrue(scout.tables["ScoutRows"].ref.endswith("3"))

        copy_requests = [request for request in fake.batch_requests if "copyPaste" in request]
        self.assertTrue(copy_requests)
        filter_requests = [request["setBasicFilter"]["filter"] for request in fake.batch_requests if "setBasicFilter" in request]
        self.assertEqual(len(filter_requests), 2)
        self.assertEqual(filter_requests[0]["criteria"], {"9": {"hiddenValues": [""]}})
        self.assertEqual(
            filter_requests[1]["criteria"],
            {"4": {"condition": {"type": "NUMBER_GREATER", "values": [{"userEnteredValue": "80"}]}}},
        )


if __name__ == "__main__":
    unittest.main()
