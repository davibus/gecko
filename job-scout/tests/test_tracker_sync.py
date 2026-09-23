from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.styles import Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table


SCOUT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SCOUT_ROOT))
SCRIPTS_ROOT = Path(__file__).resolve().parents[2] / "scripts"
sys.path.insert(0, str(SCRIPTS_ROOT))

from models import RawListing
from normalize import normalize
from tracker_sync import HEADERS, LEGACY_HEADERS, SHEET_NAME, sync_job_scout
from manage_job_tracker import HEADERS as TRACKER_HEADERS
from manage_job_tracker import (
    JobRecord, append_record, load_tracker, mark_scout_resume_created, upsert_record,
)


class TrackerSyncTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path = Path(self.temp.name) / "tracker.xlsx"
        workbook = Workbook()
        tracker = workbook.active
        tracker.title = "Job Tracker"
        tracker.append(["Existing", "Formula"])
        tracker.append(["Keep me", "=1+1"])
        tracker["A1"].font = Font(bold=True, color="FFFFFF")
        tracker["A1"].fill = PatternFill("solid", fgColor="1F4E78")
        notes = workbook.create_sheet("Notes")
        notes["A1"] = "Do not change"
        workbook.save(self.path)

    def tearDown(self):
        self.temp.cleanup()

    @staticmethod
    def cell(ws, header, row=2):
        return ws.cell(row, HEADERS.index(header) + 1)

    @staticmethod
    def job(job_id, score=82, confidence=80, *, posted="2026-09-01", status="new"):
        raw = RawListing(
            source="adzuna", source_job_id=str(job_id),
            url=f"https://www.adzuna.com/details/{job_id}?utm_source=test",
            title=f"Paid Search Manager {job_id}", company=f"Company {job_id}",
            location="Lehi, Utah", description="Paid search manager role.",
            employment_type="full_time", salary="$100,000", date_posted=posted,
        )
        job = normalize(raw, discovered="2026-09-10")
        job.id = job_id
        job.last_seen = "2026-09-19"
        job.match_score = score
        job.evidence_confidence = confidence
        job.match_strengths = ["Strong one", "Strong two", "Strong three", "Strong four"]
        job.match_weaknesses = ["Weak one", "Weak two", "Weak three", "Weak four"]
        job.status = status
        return job

    def test_creates_sheet_and_preserves_existing_tabs_data_formulas_and_style(self):
        before = load_workbook(self.path, data_only=False)
        original_style = before["Job Tracker"]["A1"].style_id
        summary = sync_job_scout([self.job(1)], self.path)
        after = load_workbook(self.path, data_only=False)
        self.assertEqual(after.sheetnames, ["Job Tracker", "Notes", SHEET_NAME])
        self.assertEqual(after["Job Tracker"]["A2"].value, "Keep me")
        self.assertEqual(after["Job Tracker"]["B2"].value, "=1+1")
        self.assertEqual(after["Job Tracker"]["A1"].style_id, original_style)
        self.assertEqual(after["Notes"]["A1"].value, "Do not change")
        self.assertTrue(summary.tracker_preserved)

    def test_new_jobs_append_with_required_headers(self):
        sync_job_scout([self.job(1), self.job(2)], self.path)
        ws = load_workbook(self.path)[SHEET_NAME]
        self.assertEqual([cell.value for cell in ws[1]], HEADERS)
        self.assertEqual(ws.max_row, 3)
        self.assertEqual({ws.cell(row, 1).value for row in (2, 3)}, {1, 2})

    def test_existing_and_enriched_jobs_update_without_duplicate(self):
        job = self.job(1, 81, 35)
        job.provisional = True
        sync_job_scout([job], self.path)
        job.match_score = 93
        job.evidence_confidence = 96
        job.provisional = False
        job.enriched_source_url = "https://careers.example.test/jobs/1"
        job.match_strengths = ["Updated strength"]
        sync_job_scout([job], self.path)
        ws = load_workbook(self.path)[SHEET_NAME]
        self.assertEqual(ws.max_row, 2)
        self.assertEqual(ws["E2"].value, 93)
        self.assertEqual(ws["H2"].value, "Confirmed Strong Match")
        self.assertEqual(ws["B2"].value, "adzuna")
        self.assertEqual(self.cell(ws, "Enrichment URL").value, "https://careers.example.test/jobs/1")
        self.assertNotIn("Top Strengths", [cell.value for cell in ws[1]])
        self.assertNotIn("Top Weaknesses", [cell.value for cell in ws[1]])

    def test_manual_applied_and_contacted_marks_are_preserved(self):
        sync_job_scout([self.job(1)], self.path)
        workbook = load_workbook(self.path)
        ws = workbook[SHEET_NAME]
        self.cell(ws, "Apply?").value = "Yes"
        self.cell(ws, "Applied").value = "X"
        self.cell(ws, "Contacted").value = "manual note"
        workbook.save(self.path)
        changed = self.job(1, 95, 99)
        sync_job_scout([changed], self.path)
        ws = load_workbook(self.path)[SHEET_NAME]
        self.assertEqual(self.cell(ws, "Apply?").value, "Yes")
        self.assertEqual(self.cell(ws, "Applied").value, "X")
        self.assertEqual(self.cell(ws, "Contacted").value, "manual note")

    def test_daily_append_only_leaves_existing_values_and_manual_fields_untouched(self):
        sync_job_scout([self.job(1, 82, 80)], self.path)
        workbook = load_workbook(self.path)
        ws = workbook[SHEET_NAME]
        self.cell(ws, "Applied").value = "X"
        self.cell(ws, "Contacted").value = "call Friday"
        workbook.save(self.path)
        before = {
            header: self.cell(load_workbook(self.path)[SHEET_NAME], header).value
            for header in HEADERS
        }

        changed_existing = self.job(1, 99, 99)
        changed_existing.company = "Must Not Replace Existing"
        sync_job_scout(
            [changed_existing, self.job(2, 88, 85)], self.path, append_only=True,
        )

        ws = load_workbook(self.path)[SHEET_NAME]
        existing_row = next(row for row in range(2, ws.max_row + 1) if ws.cell(row, 1).value == 1)
        after = {
            header: self.cell(ws, header, existing_row).value for header in HEADERS
        }
        self.assertEqual(after, before)
        self.assertEqual({ws.cell(row, 1).value for row in range(2, ws.max_row + 1)}, {1, 2})

    def test_legacy_layout_is_extended_without_moving_existing_columns_or_manual_fields(self):
        workbook = load_workbook(self.path)
        legacy = workbook.create_sheet(SHEET_NAME)
        legacy.append(LEGACY_HEADERS)
        record = {header: None for header in LEGACY_HEADERS}
        record.update({
            "Scout ID": 1, "Company": "Company 1", "Job Title": "Paid Search Manager 1",
            "Match Score": 82, "Evidence Confidence": 0.8, "Match Status": "Confirmed Strong Match",
            "Source": "adzuna", "Job URL": "https://www.adzuna.com/details/1",
            "Gecko Status": "Applied", "Applied": "X", "Contacted": "manual",
        })
        legacy.append([record[header] for header in LEGACY_HEADERS])
        workbook.save(self.path)
        sync_job_scout([self.job(1)], self.path)
        ws = load_workbook(self.path)[SHEET_NAME]
        headers = [cell.value for cell in ws[1]]
        self.assertEqual(headers[:len(LEGACY_HEADERS)], LEGACY_HEADERS)
        self.assertEqual(headers[len(LEGACY_HEADERS):], [
            header for header in HEADERS if header not in LEGACY_HEADERS
        ])
        columns = {header: index for index, header in enumerate(headers, 1)}
        self.assertEqual(ws.cell(2, columns["Source"]).value, "adzuna")
        self.assertEqual(ws.cell(2, columns["Applied"]).value, "X")
        self.assertEqual(ws.cell(2, columns["Contacted"]).value, "manual")

    def test_existing_row_order_and_hyperlinks_are_preserved(self):
        first = self.job(1, 90, 80, posted="2026-09-01")
        second = self.job(2, 91, 70, posted="2026-08-01")
        third = self.job(3, 90, 90, posted="2026-07-01")
        third.enriched_source_url = "https://careers.example.test/jobs/3"
        sync_job_scout([first, second, third], self.path)
        ws = load_workbook(self.path)[SHEET_NAME]
        self.assertEqual([ws.cell(row, 1).value for row in (2, 3, 4)], [1, 2, 3])
        for row in (2, 3, 4):
            self.assertIsNotNone(self.cell(ws, "Job URL", row).hyperlink)
        self.assertEqual(self.cell(ws, "Enrichment URL", 4).hyperlink.target, "https://careers.example.test/jobs/3")
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertEqual(ws.auto_filter.ref, "A1:X4")

    def test_normal_job_scout_update_preserves_user_excel_formatting(self):
        """A routine Scout append must retain user formatting and filter state."""
        sync_job_scout([self.job(1)], self.path)
        workbook = load_workbook(self.path)
        ws = workbook[SHEET_NAME]
        ws["E2"].fill = PatternFill("solid", fgColor="F4B183")
        ws["E2"].font = Font(name="Aptos", size=14, bold=True, italic=True, color="7030A0")
        ws["E2"].border = Border(bottom=Side(style="thick", color="00AA00"))
        ws["E2"].alignment = Alignment(horizontal="right", vertical="center")
        ws["E2"].number_format = '0 "points"'
        ws.row_dimensions[2].height = 37
        ws.column_dimensions["E"].width = 27
        ws.freeze_panes = "C3"
        ws.auto_filter.ref = "A1:X2"
        ws.auto_filter.add_filter_column(4, ["88", "90"])
        ws.auto_filter.add_sort_condition("E2:E2", descending=True)
        ws["Y1"] = "User Formula"
        ws["Y2"] = "=E2*2"
        validation = DataValidation(type="list", formula1='"Yes,No"')
        validation.add("T2")
        ws.add_data_validation(validation)
        ws.add_table(Table(displayName="ScoutHistory", ref="A1:X2"))
        ws.conditional_formatting.add(
            "A2:X2", FormulaRule(formula=['$E2>80'], fill=PatternFill("solid", fgColor="ABCDEF")),
        )
        workbook.save(self.path)

        sync_job_scout([self.job(2)], self.path)
        ws = load_workbook(self.path, data_only=False)[SHEET_NAME]

        self.assertEqual(ws["E2"].fill.fgColor.rgb, "00F4B183")
        self.assertEqual(ws["E2"].font.color.rgb, "007030A0")
        self.assertEqual(ws["E2"].border.bottom.style, "thick")
        self.assertEqual(ws["E2"].number_format, '0 "points"')
        self.assertEqual(ws["E2"].alignment.horizontal, "right")
        self.assertEqual(ws.row_dimensions[2].height, 37)
        self.assertEqual(ws.row_dimensions[3].height, 37)
        self.assertEqual(ws.column_dimensions["E"].width, 27)
        self.assertEqual(ws.freeze_panes, "C3")
        self.assertEqual(ws.auto_filter.ref, "A1:X3")
        self.assertEqual(len(ws.auto_filter.filterColumn), 1)
        self.assertIsNotNone(ws.auto_filter.sortState)
        self.assertTrue(ws.auto_filter.sortState.sortCondition[0].descending)
        self.assertEqual(ws.tables["ScoutHistory"].ref, "A1:X3")
        self.assertIn("T3", str(ws.data_validations.dataValidation[0].sqref))
        self.assertEqual(ws["Y2"].value, "=E2*2")
        self.assertEqual(ws["Y3"].value, "=E3*2")
        self.assertEqual(len(ws.conditional_formatting), 1)

    def test_old_scout_rows_are_not_deleted(self):
        sync_job_scout([self.job(1)], self.path)
        sync_job_scout([self.job(2)], self.path)
        ws = load_workbook(self.path)[SHEET_NAME]
        self.assertEqual(ws.max_row, 3)
        self.assertEqual({ws.cell(row, 1).value for row in (2, 3)}, {1, 2})

    def test_selected_status_is_written(self):
        job = self.job(1, status="selected")
        sync_job_scout([job], self.path)
        ws = load_workbook(self.path)[SHEET_NAME]
        self.assertEqual(self.cell(ws, "Gecko Status").value, "Selected")

    def test_resume_completion_updates_scout_row_without_manual_columns(self):
        sync_job_scout([self.job(1, status="selected")], self.path)
        workbook = load_workbook(self.path)
        scout = workbook[SHEET_NAME]
        self.cell(scout, "Applied").value = "X"
        self.cell(scout, "Contacted").value = "manual"
        record = JobRecord(
            company="Company 1", job_title="Paid Search Manager 1", pay="",
            job_number="1", match_score="90/100", job_link="",
            resume_path=Path("resume.docx"), date_created=None, scout_id=1,
        )
        self.assertTrue(mark_scout_resume_created(workbook, record))
        self.assertEqual(self.cell(scout, "Gecko Status").value, "Resume Created")
        self.assertEqual(self.cell(scout, "Resume Created").value, "X")
        self.assertEqual(self.cell(scout, "Applied").value, "X")
        self.assertTrue(str(self.cell(scout, "Resume Link").value).startswith("file:///"))
        self.assertEqual(self.cell(scout, "Contacted").value, "manual")

    def test_existing_tracker_loader_accepts_job_scout_tab(self):
        workbook = Workbook()
        tracker = workbook.active
        tracker.title = "Job Tracker"
        tracker.append(TRACKER_HEADERS)
        workbook.create_sheet(SHEET_NAME)
        workbook.save(self.path)
        loaded = load_tracker(self.path)
        self.assertEqual(loaded.sheetnames, ["Job Tracker", SHEET_NAME])

    def test_application_upsert_updates_owned_fields_and_preserves_manual_fields(self):
        workbook = Workbook()
        tracker = workbook.active
        tracker.title = "Job Tracker"
        tracker.append(TRACKER_HEADERS)
        tracker.append([1, "Old Company", "Old Title", None, "job-1", "80/100", None,
                        "resumes/old.docx", None, "X", "manual", None, None, "resume-created"])
        record = JobRecord(
            company="New Company", job_title="New Title", pay="$100,000",
            job_number="job-1", match_score="95/100", job_link="https://example.test/job-1",
            resume_path=Path("resumes/new.docx"), date_created=None, status="resume-created",
        )
        added, changed = upsert_record(tracker, self.path, record)
        self.assertFalse(added)
        self.assertTrue(changed)
        self.assertEqual(tracker.max_row, 2)
        self.assertEqual(tracker["B2"].value, "New Company")
        self.assertEqual(tracker["F2"].value, "95/100")
        self.assertEqual(tracker["J2"].value, "X")
        self.assertEqual(tracker["K2"].value, "manual")

    def test_application_append_copies_prior_format_and_expands_table_without_restyling(self):
        workbook = Workbook()
        tracker = workbook.active
        tracker.title = "Job Tracker"
        tracker.append(TRACKER_HEADERS)
        tracker.append([1, "Existing", "Existing Role", None, "old", "80/100", None,
                        "resumes/old.docx", None, None, "=A2+1", None, None, "resume-created"])
        tracker["C2"].fill = PatternFill("solid", fgColor="92D050")
        tracker["C2"].font = Font(size=15, italic=True, color="C00000")
        tracker.row_dimensions[2].height = 31
        tracker.column_dimensions["C"].width = 55
        tracker.freeze_panes = "D4"
        tracker.auto_filter.ref = "A1:N2"
        tracker.add_table(Table(displayName="ApplicationHistory", ref="A1:N2"))
        validation = DataValidation(type="list", formula1='"Yes,No"')
        validation.add("J2")
        tracker.add_data_validation(validation)
        workbook.save(self.path)

        record = JobRecord(
            company="New Company", job_title="New Role", pay="$100,000",
            job_number="new", match_score="90/100", job_link="https://example.test/new",
            resume_path=Path("resumes/new.docx"), date_created=None,
        )
        workbook = load_workbook(self.path)
        tracker = workbook["Job Tracker"]
        self.assertTrue(append_record(tracker, self.path, record))
        workbook.save(self.path)

        tracker = load_workbook(self.path, data_only=False)["Job Tracker"]
        self.assertEqual(tracker["C2"].fill.fgColor.rgb, "0092D050")
        self.assertEqual(tracker["C2"].font.color.rgb, "00C00000")
        self.assertEqual(tracker["K2"].value, "=A2+1")
        self.assertEqual(tracker["C3"].fill.fgColor.rgb, "0092D050")
        self.assertEqual(tracker.row_dimensions[3].height, 31)
        self.assertEqual(tracker.column_dimensions["C"].width, 55)
        self.assertEqual(tracker.freeze_panes, "D4")
        self.assertEqual(tracker.auto_filter.ref, "A1:N3")
        self.assertEqual(tracker.tables["ApplicationHistory"].ref, "A1:N3")
        self.assertIn("J3", str(tracker.data_validations.dataValidation[0].sqref))


if __name__ == "__main__":
    unittest.main()
