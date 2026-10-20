"""Word-source fidelity and no-fallback checks for the canonical reader."""

from pathlib import Path
import tempfile
import unittest

from docx import Document
from resume_source import read_master_text


class WordMasterReaderTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "master-resume-v4.docx"

    def test_tables_stay_between_surrounding_paragraphs(self):
        doc = Document()
        doc.add_paragraph("Before table")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Verified project"
        table.cell(0, 1).text = "Supported result"
        doc.add_paragraph("After table")
        doc.save(self.path)
        self.assertEqual(read_master_text(self.path).splitlines(),
                         ["Before table", "Verified project", "Supported result", "After table"])

    def test_word_list_styles_and_literal_bullets_are_preserved(self):
        doc = Document()
        doc.add_paragraph("Supported accomplishment", style="List Bullet")
        doc.add_paragraph("\u2022 Existing literal bullet", style="List Bullet")
        doc.save(self.path)
        self.assertEqual(read_master_text(self.path).splitlines(),
                         ["\u2022 Supported accomplishment", "\u2022 Existing literal bullet"])

    def test_reads_saved_changes_without_a_text_export(self):
        doc = Document()
        doc.add_paragraph("Original evidence")
        doc.save(self.path)
        self.assertNotIn("Updated evidence", read_master_text(self.path))
        doc.add_paragraph("Updated evidence")
        doc.save(self.path)
        self.assertIn("Updated evidence", read_master_text(self.path))

    def test_corrupt_docx_does_not_fall_back_to_old_text(self):
        self.path.with_suffix(".txt").write_text("Old evidence", encoding="utf-8")
        self.path.write_bytes(b"not a Word package")
        with self.assertRaisesRegex(OSError, "missing or unreadable"):
            read_master_text(self.path)

    def test_empty_word_document_is_rejected(self):
        Document().save(self.path)
        with self.assertRaisesRegex(ValueError, "is empty"):
            read_master_text(self.path)
