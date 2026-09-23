"""The scouting reader must use the current master DOCX for evidence."""

import tempfile
import unittest
import sys
from pathlib import Path

from docx import Document
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from resume_evidence import extract_resume_text


class ResumeEvidenceTests(unittest.TestCase):
    def test_reads_updated_docx_paragraphs_and_tables(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "Dave-Call-resume-9-23-26.docx"
            doc = Document()
            doc.add_paragraph("First approved fact")
            doc.add_table(rows=1, cols=1).cell(0, 0).text = "Approved tool"
            doc.save(path)
            self.assertIn("Approved tool", extract_resume_text(path))
            doc.add_paragraph("Updated accomplishment")
            doc.save(path)
            self.assertIn("Updated accomplishment", extract_resume_text(path))

    def test_rejects_old_pdf_as_evidence(self):
        with self.assertRaisesRegex(ValueError, "must come from"):
            extract_resume_text("input/master-resume/dcall-resume-3-15-26.pdf")


if __name__ == "__main__":
    unittest.main()
