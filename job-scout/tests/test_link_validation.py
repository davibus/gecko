"""Availability checks and daily cleanup never treat transient errors as expiry."""

from __future__ import annotations

import io
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch
from urllib.error import HTTPError, URLError

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.formatting.rule import FormulaRule

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from link_validation import LinkResult, validate_url
from models import RawListing
from normalize import normalize
from scout import build_parser, daily
from service import discover
from sources.base import SearchRequest
from storage import JobStore
from tracker_sync import HEADERS, sync_job_scout


URL = "https://jobs.example.test/posting/123"


class FakeResponse:
    def __init__(self, body=b"<h1>Paid Search Manager</h1>", final_url=URL):
        self.body, self.final_url, self.status = body, final_url, 200

    def __enter__(self):
        return self

    def __exit__(self, *_):
        return False

    def geturl(self):
        return self.final_url

    def read(self, _limit):
        return self.body


class FakeOpener:
    def __init__(self, result):
        self.result = result

    def open(self, *_args, **_kwargs):
        if isinstance(self.result, Exception):
            raise self.result
        return self.result


class LinkValidationTests(unittest.TestCase):
    def check_response(self, response):
        with patch("link_validation.build_opener", return_value=FakeOpener(response)):
            return validate_url(URL)

    def test_200_valid_listing(self):
        self.assertEqual(self.check_response(FakeResponse()).status, "valid")

    def test_redirect_to_valid_final_listing(self):
        final = "https://company.example.test/careers/123"
        result = self.check_response(FakeResponse(final_url=final))
        self.assertEqual((result.status, result.final_url), ("valid", final))

    def test_404_and_410_are_dead(self):
        for status in (404, 410):
            with self.subTest(status=status):
                result = self.check_response(HTTPError(URL, status, "gone", {}, None))
                self.assertEqual((result.status, result.http_status), ("dead", status))

    def test_explicit_no_longer_available_page_is_dead(self):
        response = FakeResponse(body=b"<html><h1>This job is no longer available</h1></html>")
        self.assertEqual(self.check_response(response).status, "dead")

    def test_403_429_500_and_timeout_are_temporary(self):
        for status in (403, 429, 500):
            with self.subTest(status=status):
                result = self.check_response(HTTPError(URL, status, "temporary", {}, None))
                self.assertEqual(result.status, "temporary_failure")
        self.assertEqual(self.check_response(TimeoutError()).status, "temporary_failure")
        self.assertEqual(self.check_response(URLError("DNS unavailable")).status, "temporary_failure")
        self.assertEqual(self.check_response(FakeResponse(body=b"<h1>Cloudflare: checking your browser</h1>")).status,
                         "temporary_failure")

    def test_official_ats_api_absence_is_dead(self):
        lever = "https://jobs.lever.co/example/abc123"
        with patch("link_validation._request", return_value=LinkResult("dead", lever, http_status=404)) as fetch:
            result = validate_url(lever)
        self.assertEqual(result.status, "dead")
        self.assertIn("Official ATS API", result.reason)
        self.assertEqual(fetch.call_count, 2)


def job(source_id: str, company: str):
    return normalize(RawListing(source="test", source_job_id=source_id,
                                url=f"https://jobs.example.test/{source_id}",
                                title="Paid Search Manager", company=company,
                                description="Manage Google Ads campaigns and budgets."))


class DailyCleanupTests(unittest.TestCase):
    def test_sheet_only_active_row_is_validated_and_cleared(self):
        with tempfile.TemporaryDirectory() as tmp:
            tracker = Path(tmp) / "tracker.xlsx"
            workbook = Workbook()
            workbook.active.title = "Job Tracker"
            workbook["Job Tracker"].append(["Company", "Job Title", "Job Number"])
            scout = workbook.create_sheet("Job Scout")
            scout.append(HEADERS)
            row = {"Scout ID": 999, "Company": "Sheet Only Co", "Job Title": "Paid Search Manager",
                   "Job URL": "https://jobs.example.test/sheet-only", "Gecko Status": "New"}
            scout.append([row.get(header) for header in HEADERS])
            workbook.save(tracker)
            with JobStore(Path(tmp) / "jobs.sqlite3") as store:
                args = build_parser().parse_args(["--tracker", str(tracker), "daily"])
                def no_search(search_args, *_):
                    search_args.run_result = {"new_job_ids": []}
                    return 0
                with patch("link_validation.validate_job", return_value=LinkResult("dead", row["Job URL"], reason="HTTP 410", http_status=410)), \
                     patch("scout.search", side_effect=no_search), redirect_stdout(io.StringIO()):
                    self.assertEqual(daily(args, store, {}), 0)
            workbook = load_workbook(tracker)
            self.assertIsNone(workbook["Job Scout"]["A2"].value)

    def test_new_dead_listing_is_checked_before_scoring_or_storage(self):
        class Provider:
            def search(self, _request):
                return [RawListing(source="test", source_job_id="new-dead", url=URL,
                                   title="Paid Search Manager", company="Dead Co",
                                   description="Manage Google Ads campaigns.")]

        with tempfile.TemporaryDirectory() as tmp, JobStore(Path(tmp) / "jobs.sqlite3") as store:
            from link_validation import DailyLinkValidator
            validator = DailyLinkValidator()
            with patch("link_validation.validate_job", return_value=LinkResult("dead", URL, reason="HTTP 404", http_status=404)), \
                 patch("service.score_job", side_effect=AssertionError("scored before validation")):
                summary = discover(Provider(), [SearchRequest("paid search", "Remote")], store,
                                   {"minimum_score": 80}, "master resume", link_validator=validator)
            self.assertEqual(summary.scored, 0)
            self.assertEqual(store.all(), [])
            self.assertEqual(validator.report()["removed_dead"], 1)

    def test_dead_unprotected_is_cleared_without_losing_history_or_formatting(self):
        with tempfile.TemporaryDirectory() as tmp:
            tracker = Path(tmp) / "tracker.xlsx"
            workbook = Workbook()
            app = workbook.active
            app.title = "Job Tracker"
            app.append(["Company", "Job Title", "Job Number", "Job Link", "Applied", "Contacted"])
            app.append(["Protected Co", "Paid Search Manager", "protected", "https://jobs.example.test/protected", "X", "Called"])
            workbook.save(tracker)
            with JobStore(Path(tmp) / "jobs.sqlite3") as store:
                dead = job("dead", "Dead Co")
                protected = job("protected", "Protected Co")
                resume_created = job("resume-created", "Resume Co")
                resume_created.status = "resume-created"
                applied = job("applied", "Applied Co")
                applied.status = "applied"
                temporary = job("temporary", "Temporary Co")
                dead.id = store.save(dead)
                protected.id = store.save(protected)
                resume_created.id = store.save(resume_created)
                applied.id = store.save(applied)
                temporary.id = store.save(temporary)
                sync_job_scout(store.all(), tracker)
                workbook = load_workbook(tracker)
                sheet = workbook["Job Scout"]
                columns = {cell.value: cell.column for cell in sheet[1]}
                by_id = {sheet.cell(row, columns["Scout ID"]).value: row for row in range(2, sheet.max_row + 1)}
                dead_row = by_id[dead.id]
                protected_row = by_id[protected.id]
                sheet.cell(dead_row, columns["Company"]).fill = PatternFill("solid", fgColor="FFCC00")
                sheet.cell(protected_row, columns["Applied"]).value = "X"
                sheet.cell(protected_row, columns["Contacted"]).value = "Called"
                sheet.column_dimensions["B"].width = 47
                sheet.freeze_panes = "C3"
                sheet.auto_filter.ref = f"A1:Y{sheet.max_row}"
                sheet.conditional_formatting.add(f"A2:Y{sheet.max_row}", FormulaRule(formula=['$D2>80'], fill=PatternFill("solid", fgColor="00FF00")))
                workbook.save(tracker)

                def result(candidate):
                    if candidate.source_job_id == "temporary":
                        return LinkResult("temporary_failure", candidate.url, reason="HTTP 429", http_status=429)
                    return LinkResult("dead", candidate.url, reason="HTTP 404 at final URL", http_status=404)

                def no_search(search_args, *_):
                    search_args.run_result = {"new_job_ids": []}
                    return 0

                args = build_parser().parse_args(["--tracker", str(tracker), "daily"])
                output = io.StringIO()
                with patch("link_validation.validate_job", side_effect=result), patch("scout.search", side_effect=no_search), redirect_stdout(output):
                    self.assertEqual(daily(args, store, {}), 0)
                self.assertIsNone(store.get(dead.id))
                self.assertIsNotNone(store.get(protected.id))
                self.assertEqual(store.get(resume_created.id).status, "resume-created")
                self.assertEqual(store.get(applied.id).status, "applied")
                self.assertIsNotNone(store.get(temporary.id))
                self.assertIn("removed_dead: 1", output.getvalue())
                self.assertIn("temporary_failure: 1", output.getvalue())
                self.assertIn("Dead Co", output.getvalue())

            workbook = load_workbook(tracker)
            sheet = workbook["Job Scout"]
            self.assertIsNone(sheet.cell(dead_row, columns["Scout ID"]).value)
            self.assertEqual(sheet.cell(dead_row, columns["Company"]).fill.fgColor.rgb[-6:], "FFCC00")
            self.assertEqual(sheet.cell(protected_row, columns["Applied"]).value, "X")
            self.assertEqual(sheet.cell(protected_row, columns["Contacted"]).value, "Called")
            self.assertIn(resume_created.id, [sheet.cell(row, columns["Scout ID"]).value for row in range(2, sheet.max_row + 1)])
            self.assertIn(applied.id, [sheet.cell(row, columns["Scout ID"]).value for row in range(2, sheet.max_row + 1)])
            self.assertEqual(sheet.column_dimensions["B"].width, 47)
            self.assertEqual(sheet.freeze_panes, "C3")
            self.assertEqual(sheet.auto_filter.ref, f"A1:Y{sheet.max_row}")
            self.assertEqual(len(sheet.conditional_formatting), 2)
            self.assertEqual(workbook["Job Tracker"]["E2"].value, "X")
            self.assertEqual(workbook["Job Tracker"]["F2"].value, "Called")


if __name__ == "__main__":
    unittest.main()
