"""Synchronize Gecko's two tracker worksheets with the Google Sheets API."""

from __future__ import annotations

import os
import sys
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any, Iterable
from urllib.parse import quote, unquote, urlparse

from openpyxl import load_workbook

from normalize import canonicalize_url
from excel_preservation import append_preserving_format
from tracker_sync import (
    APPLICATION_SHEET,
    HEADERS as SCOUT_HEADERS,
    PRE_APPLY_HEADERS as PRE_APPLY_SCOUT_HEADERS,
    LEGACY_HEADERS as LEGACY_SCOUT_HEADERS,
    PREVIOUS_HEADERS as PREVIOUS_SCOUT_HEADERS,
    RESUME_LINK_HEADERS as RESUME_LINK_SCOUT_HEADERS,
    SOURCE_B_HEADERS as SOURCE_B_SCOUT_HEADERS,
    URL_STATUSLESS_HEADERS as URL_STATUSLESS_SCOUT_HEADERS,
    SHEET_NAME as SCOUT_SHEET,
    STATUS_RANK,
    _sort_key as scout_sort_key,
)


APPLICATION_HEADERS = [
    "Resume #", "Company", "Job Title", "Pay", "Job Number", "Match Score",
    "Job Link", "Resume Link", "Date Created", "Applied", "Contacted", "Source",
    "Date Found", "Status",
]
MANUAL_COLUMNS = ("Apply?", "Applied", "Contacted")
SHEETS_SCOPE = "https://www.googleapis.com/auth/spreadsheets"
VALID_BACKENDS = {"xlsx", "google-sheets"}


@dataclass(frozen=True)
class GoogleSheetsConfig:
    spreadsheet_id: str
    credentials_file: Path

    @classmethod
    def from_environment(cls) -> "GoogleSheetsConfig":
        spreadsheet_id = os.getenv("GOOGLE_SHEETS_SPREADSHEET_ID", "").strip()
        credentials_value = os.getenv("GOOGLE_SHEETS_CREDENTIALS_FILE", "").strip()
        missing = []
        if not spreadsheet_id:
            missing.append("GOOGLE_SHEETS_SPREADSHEET_ID")
        if not credentials_value:
            missing.append("GOOGLE_SHEETS_CREDENTIALS_FILE")
        if missing:
            raise ValueError("Missing Google Sheets configuration: " + ", ".join(missing))
        credentials_file = Path(os.path.expandvars(credentials_value)).expanduser()
        if not credentials_file.is_file():
            raise ValueError(f"Google Sheets credentials file does not exist: {credentials_file}")
        return cls(spreadsheet_id=spreadsheet_id, credentials_file=credentials_file)


@dataclass(frozen=True)
class GoogleSyncSummary:
    application_rows: int
    scout_rows: int
    spreadsheet_id: str


def load_environment_files(project_root: Path) -> None:
    """Load ignored local environment files without replacing process variables."""
    for path in (project_root / ".env.local", project_root / ".env.google-sheets.local"):
        if not path.is_file():
            continue
        for number, raw_line in enumerate(path.read_text(encoding="utf-8-sig").splitlines(), 1):
            line = raw_line.strip()
            if not line or line.startswith("#"):
                continue
            if line.startswith("export "):
                line = line[7:].strip()
            if "=" not in line:
                raise ValueError(f"Invalid environment entry at {path.name}:{number}")
            name, value = line.split("=", 1)
            name = name.strip()
            if not name or not name.replace("_", "").isalnum() or name[0].isdigit():
                raise ValueError(f"Invalid environment variable name at {path.name}:{number}")
            value = value.strip()
            if len(value) >= 2 and value[0] == value[-1] and value[0] in "\"'":
                value = value[1:-1]
            os.environ.setdefault(name, value)


def tracker_backend() -> str:
    backend = os.getenv("TRACKER_BACKEND", "xlsx").strip().lower() or "xlsx"
    if backend not in VALID_BACKENDS:
        raise ValueError(f"TRACKER_BACKEND must be one of: {', '.join(sorted(VALID_BACKENDS))}")
    return backend


def build_sheets_service(config: GoogleSheetsConfig):
    """Create the official Google Sheets API client with narrow spreadsheet scope."""
    try:
        from google.oauth2.service_account import Credentials
        from googleapiclient.discovery import build
    except ImportError as error:
        raise RuntimeError(
            "Google Sheets dependencies are missing; install job-scout/requirements.txt"
        ) from error
    credentials = Credentials.from_service_account_file(
        str(config.credentials_file), scopes=[SHEETS_SCOPE]
    )
    return build("sheets", "v4", credentials=credentials, cache_discovery=False)


def _sheet_range(title: str, columns: str) -> str:
    return f"'{title.replace(chr(39), chr(39) * 2)}'!A1:{columns}"


def _normalize_cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%m/%d/%Y")
    if isinstance(value, date):
        return value.strftime("%m/%d/%Y")
    return value


def _records_from_values(values: list[list[Any]], headers: list[str], title: str) -> list[dict]:
    if not values:
        return []
    actual = list(values[0])
    while actual and actual[-1] in (None, ""):
        actual.pop()
    if len(set(actual)) != len(actual):
        raise ValueError(f"Google Sheet {title!r} contains duplicate column names")
    if title == SCOUT_SHEET:
        supported = (
            headers, PRE_APPLY_SCOUT_HEADERS, LEGACY_SCOUT_HEADERS, SOURCE_B_SCOUT_HEADERS,
            RESUME_LINK_SCOUT_HEADERS, PREVIOUS_SCOUT_HEADERS, URL_STATUSLESS_SCOUT_HEADERS,
        )
        if not any(set(candidate).issubset(actual) for candidate in supported):
            raise ValueError(f"Google Sheet {title!r} columns do not match Gecko's required schema")
    elif not set(headers).issubset(actual):
        raise ValueError(f"Google Sheet {title!r} columns do not match Gecko's required schema")
    records = []
    for source_row in values[1:]:
        row = list(source_row) + [""] * (len(actual) - len(source_row))
        record = {header: row[index] for index, header in enumerate(actual)}
        if any(value not in (None, "") for value in record.values()):
            records.append(record)
    return records


def _output_headers(values: list[list[Any]], required: list[str]) -> list[str]:
    """Keep the user's column order and append newly managed columns if needed."""
    if not values:
        return list(required)
    actual = list(values[0])
    while actual and actual[-1] in (None, ""):
        actual.pop()
    return actual + [header for header in required if header not in actual]


def _records_from_xlsx(ws, headers: list[str]) -> list[dict]:
    actual = [ws.cell(1, column).value for column in range(1, ws.max_column + 1)]
    while actual and actual[-1] in (None, ""):
        actual.pop()
    source_headers = actual
    if ws.title == SCOUT_SHEET and actual in (
        PRE_APPLY_SCOUT_HEADERS, LEGACY_SCOUT_HEADERS, SOURCE_B_SCOUT_HEADERS, RESUME_LINK_SCOUT_HEADERS,
        PREVIOUS_SCOUT_HEADERS, URL_STATUSLESS_SCOUT_HEADERS,
    ):
        source_headers = actual
    elif not set(headers).issubset(actual):
        raise ValueError(f"XLSX worksheet {ws.title!r} columns do not match Gecko's required schema")
    return [
        {header: ws.cell(row, column).value for column, header in enumerate(source_headers, 1)}
        for row in range(2, ws.max_row + 1)
        if any(ws.cell(row, column).value not in (None, "") for column in range(1, len(source_headers) + 1))
    ]


def _canonical_identity(value: Any) -> str:
    return canonicalize_url(str(value or "")).lower()


def _copy_manual(target: dict, source: dict) -> None:
    for column in MANUAL_COLUMNS:
        if source.get(column) not in (None, ""):
            target[column] = source[column]


def _copy_custom(target: dict, source: dict, managed_headers: list[str]) -> None:
    """Carry user-owned columns with their row when sorted or upserted."""
    managed = set(managed_headers)
    for column, value in source.items():
        if column not in managed:
            target[column] = value


def _preserve_remote_url_verification(target: dict, source: dict) -> None:
    """Do not erase an existing verified URL with an older local backup."""
    local_status = str(target.get("URL Status") or "")
    remote_status = str(source.get("URL Status") or "")
    if remote_status and local_status in {"", "Not checked"}:
        target["URL Status"] = source.get("URL Status")
        target["Authoritative URL"] = source.get("Authoritative URL")


def _merge_scout_lifecycle(target: dict, source: dict) -> None:
    _copy_manual(target, source)
    candidates = [str(target.get("Gecko Status") or "New"), str(source.get("Gecko Status") or "New")]
    if target.get("Resume Created") or source.get("Resume Created"):
        candidates.append("Resume Created")
        target["Resume Created"] = "X"
    if target.get("Applied"):
        candidates.append("Applied")
    if target.get("Contacted"):
        candidates.append("Contacted")
    target["Gecko Status"] = max(candidates, key=lambda value: STATUS_RANK.get(value, 0))


def merge_application_rows(local_rows: list[dict], remote_rows: list[dict]) -> list[dict]:
    """Merge by Job Number while preserving manual Google Sheet fields and history."""
    merged = [dict(row) for row in local_rows]
    by_job = {
        str(row.get("Job Number") or "").strip(): row
        for row in merged if row.get("Job Number") not in (None, "")
    }
    for remote in remote_rows:
        key = str(remote.get("Job Number") or "").strip()
        existing = by_job.get(key) if key else None
        if existing is None:
            clone = dict(remote)
            merged.append(clone)
            if key:
                by_job[key] = clone
        else:
            _copy_custom(existing, remote, APPLICATION_HEADERS)
            _merge_scout_lifecycle(existing, remote)

    def rank(record: dict):
        value = record.get("Resume #")
        try:
            return (0, int(value))
        except (TypeError, ValueError):
            return (1, str(value or ""))

    merged.sort(key=rank)
    return merged


def merge_scout_rows(local_rows: list[dict], remote_rows: list[dict]) -> list[dict]:
    """Upsert by Scout ID or canonical URL without deleting historical records."""
    merged = [dict(row) for row in local_rows]
    by_id = {
        str(row.get("Scout ID")): row
        for row in merged if row.get("Scout ID") not in (None, "")
    }
    by_url: dict[str, dict] = {}
    for row in merged:
        for column in ("Job URL", "Enrichment URL"):
            identity = _canonical_identity(row.get(column))
            if identity:
                by_url[identity] = row
    for remote in remote_rows:
        existing = None
        scout_id = remote.get("Scout ID")
        if scout_id not in (None, ""):
            existing = by_id.get(str(scout_id))
        if existing is None:
            for column in ("Job URL", "Enrichment URL"):
                existing = by_url.get(_canonical_identity(remote.get(column)))
                if existing is not None:
                    break
        if existing is None:
            clone = dict(remote)
            merged.append(clone)
            if scout_id not in (None, ""):
                by_id[str(scout_id)] = clone
            for column in ("Job URL", "Enrichment URL"):
                identity = _canonical_identity(clone.get(column))
                if identity:
                    by_url[identity] = clone
        else:
            _copy_custom(existing, remote, SCOUT_HEADERS)
            _preserve_remote_url_verification(existing, remote)
            _copy_manual(existing, remote)
    merged.sort(key=scout_sort_key)
    return merged


def _text_key(value: Any) -> str:
    return " ".join(str(value or "").lower().split())


def _resume_uri(value: Any, tracker_path: Path) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    if "://" in text:
        return text
    path = Path(text)
    if not path.is_absolute():
        path = tracker_path.parent / path
    return path.resolve().as_uri()


def attach_resume_links(application_rows: list[dict], scout_rows: list[dict], tracker_path: Path) -> None:
    """Backfill Scout resume links from matching application rows."""
    by_url = {}
    by_name = {}
    for application in application_rows:
        resume_link = _resume_uri(application.get("Resume Link"), tracker_path)
        if not resume_link:
            continue
        identity = _canonical_identity(application.get("Job Link"))
        if identity:
            by_url[identity] = resume_link
        name_key = (_text_key(application.get("Company")), _text_key(application.get("Job Title")))
        if all(name_key):
            by_name[name_key] = resume_link
    for scout in scout_rows:
        if scout.get("Resume Link"):
            scout["Resume Link"] = _resume_uri(scout["Resume Link"], tracker_path)
            continue
        resume_link = ""
        for column in ("Job URL", "Enrichment URL"):
            resume_link = by_url.get(_canonical_identity(scout.get(column)), "")
            if resume_link:
                break
        if not resume_link:
            resume_link = by_name.get(
                (_text_key(scout.get("Company")), _text_key(scout.get("Job Title"))), ""
            )
        if resume_link:
            scout["Resume Link"] = resume_link


def _values_from_records(headers: list[str], records: list[dict]) -> list[list[Any]]:
    return [headers] + [
        [_normalize_cell(record.get(header)) for header in headers]
        for record in records
    ]


def _preserve_remote_row_order(records: list[dict], remote_rows: list[dict], title: str) -> list[dict]:
    """Keep existing Google rows in place so their user formatting stays attached."""
    by_key: dict[str, dict] = {}
    for record in records:
        for key in _xlsx_record_keys(record, title):
            by_key[key] = record
    ordered: list[dict] = []
    seen: set[int] = set()
    for remote in remote_rows:
        record = next(
            (by_key[key] for key in _xlsx_record_keys(remote, title) if key in by_key),
            None,
        )
        if record is not None and id(record) not in seen:
            ordered.append(record)
            seen.add(id(record))
    ordered.extend(record for record in records if id(record) not in seen)
    return ordered


def _scout_values_preserving_positions(headers: list[str], records: list[dict],
                                      remote_values: list[list[Any]], removed_ids: set[int]) -> list[list[Any]]:
    """Leave cleared/empty Scout rows in place so formatting stays with survivors."""
    if not remote_values:
        return _values_from_records(headers, records)
    by_key = {key: record for record in records for key in _xlsx_record_keys(record, SCOUT_SHEET)}
    existing_headers = list(remote_values[0])
    id_column = existing_headers.index("Scout ID") if "Scout ID" in existing_headers else -1
    output = [headers]
    seen: set[int] = set()
    for source in remote_values[1:]:
        scout_id = str(source[id_column]) if id_column >= 0 and id_column < len(source) else ""
        if scout_id and scout_id in {str(value) for value in removed_ids}:
            output.append([""] * len(headers))
            continue
        row = dict(zip(existing_headers, source))
        record = next((by_key[key] for key in _xlsx_record_keys(row, SCOUT_SHEET) if key in by_key), None)
        if record is None or id(record) in seen:
            output.append([""] * len(headers))
            continue
        output.append([_normalize_cell(record.get(header)) for header in headers])
        seen.add(id(record))
    output.extend([_normalize_cell(record.get(header)) for header in headers]
                  for record in records if id(record) not in seen)
    return output


def _metadata(api, spreadsheet_id: str) -> dict:
    return api.get(
        spreadsheetId=spreadsheet_id,
        fields=(
            "sheets(properties(sheetId,title,index,gridProperties),"
            "basicFilter,conditionalFormats)"
        ),
    ).execute()


def _ensure_required_sheets(api, values_api, spreadsheet_id: str) -> dict[str, dict]:
    metadata = _metadata(api, spreadsheet_id)
    sheets = metadata.get("sheets", [])
    by_title = {sheet["properties"]["title"]: sheet for sheet in sheets}
    requests = []

    if APPLICATION_SHEET not in by_title and SCOUT_SHEET not in by_title and len(sheets) == 1:
        only = sheets[0]
        only_title = only["properties"]["title"]
        current = values_api.get(
            spreadsheetId=spreadsheet_id,
            range=_sheet_range(only_title, "A"),
            valueRenderOption="UNFORMATTED_VALUE",
        ).execute().get("values", [])
        if not current:
            requests.append({
                "updateSheetProperties": {
                    "properties": {"sheetId": only["properties"]["sheetId"], "title": APPLICATION_SHEET},
                    "fields": "title",
                }
            })
            by_title[APPLICATION_SHEET] = only

    for title in (APPLICATION_SHEET, SCOUT_SHEET):
        if title not in by_title:
            requests.append({"addSheet": {"properties": {"title": title}}})
    if requests:
        api.batchUpdate(spreadsheetId=spreadsheet_id, body={"requests": requests}).execute()
        metadata = _metadata(api, spreadsheet_id)
    return {sheet["properties"]["title"]: sheet for sheet in metadata.get("sheets", [])}


def _column_letter(index: int) -> str:
    """Return an A1 column name for a zero-based index."""
    value = index + 1
    result = ""
    while value:
        value, remainder = divmod(value - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _new_row_format_requests(
    sheet: dict, title: str, previous_row_count: int, row_count: int,
    headers: list[str],
) -> list[dict]:
    """Apply only essential number formats to cells in newly appended rows."""
    if row_count <= previous_row_count:
        return []
    sheet_id = sheet["properties"]["sheetId"]
    requests: list[dict] = []
    date_headers = (
        ("Date Created", "Date Found")
        if title == APPLICATION_SHEET
        else ("Date Posted", "Date Found", "Last Seen")
    )
    formats = [(header, {"type": "DATE", "pattern": "mm/dd/yyyy"}) for header in date_headers]
    if title == SCOUT_SHEET:
        formats.append(("Evidence Confidence", {"type": "PERCENT", "pattern": "0%"}))
    for header, number_format in formats:
        if header not in headers:
            continue
        column = headers.index(header)
        requests.append({
            "repeatCell": {
                "range": {
                    "sheetId": sheet_id,
                    "startRowIndex": max(previous_row_count, 1),
                    "endRowIndex": row_count,
                    "startColumnIndex": column,
                    "endColumnIndex": column + 1,
                },
                "cell": {"userEnteredFormat": {"numberFormat": number_format}},
                "fields": "userEnteredFormat.numberFormat",
            }
        })
    return requests


def _resume_link_format_requests(sheet_id: int, scout_rows: list[dict], headers: list[str] = SCOUT_HEADERS) -> list[dict]:
    """Apply hyperlinks that can launch local DOCX files from a browser sheet."""
    requests = []
    resume_column = headers.index("Resume Link")
    for row_index, record in enumerate(scout_rows, 1):
        link = str(record.get("Resume Link") or "").strip()
        if not link:
            continue
        # Google Sheets only activates ordinary web protocols. Route local files
        # through Gecko's localhost-only opener, which then hands the DOCX to Word.
        if link.lower().startswith("file:///"):
            filename = Path(unquote(urlparse(link).path)).name
            click_target = f"http://127.0.0.1:8765/open/{quote(filename)}"
        else:
            click_target = link
        requests.append({
            "updateCells": {
                "range": {
                    "sheetId": sheet_id,
                    "startRowIndex": row_index,
                    "endRowIndex": row_index + 1,
                    "startColumnIndex": resume_column,
                    "endColumnIndex": resume_column + 1,
                },
                "rows": [{
                    "values": [{
                        "userEnteredFormat": {
                            "textFormat": {
                                "link": {"uri": click_target},
                            }
                        }
                    }]
                }],
                "fields": "userEnteredFormat.textFormat.link",
            }
        })
    return requests


def _write_sheet(
    values_api, spreadsheet_id: str, title: str,
    existing_values: list[list[Any]], values: list[list[Any]],
) -> None:
    """Write only changed row spans; never clear or rewrite the complete sheet."""
    escaped = title.replace("'", "''")
    for row_index, desired_row in enumerate(values, 1):
        existing_row = existing_values[row_index - 1] if row_index <= len(existing_values) else []
        width = len(desired_row)
        current = list(existing_row) + [""] * max(0, width - len(existing_row))
        changed = [index for index in range(width) if _normalize_cell(current[index]) != desired_row[index]]
        if not changed:
            continue
        start, end = min(changed), max(changed)
        start_column = _column_letter(start)
        end_column = _column_letter(end)
        values_api.update(
            spreadsheetId=spreadsheet_id,
            range=f"'{escaped}'!{start_column}{row_index}:{end_column}{row_index}",
            valueInputOption="USER_ENTERED",
            body={"values": [desired_row[start:end + 1]]},
        ).execute()


def _run_optional_requests(api, spreadsheet_id: str, requests: list[dict], label: str) -> None:
    """Formatting enhancements must never make a successful data sync fail."""
    if not requests:
        return
    try:
        api.batchUpdate(
            spreadsheetId=spreadsheet_id, body={"requests": requests},
        ).execute()
    except Exception as error:
        print(f"Warning: skipped optional Google Sheets {label}: {error}", file=sys.stderr)


def _xlsx_record_keys(record: dict, title: str) -> list[str]:
    if title == APPLICATION_SHEET:
        value = str(record.get("Job Number") or "").strip()
        return [f"job:{value}"] if value else []
    keys = []
    scout_id = record.get("Scout ID")
    if scout_id not in (None, ""):
        keys.append(f"id:{scout_id}")
    for header in ("Job URL", "Enrichment URL"):
        identity = _canonical_identity(record.get(header))
        if identity:
            keys.append(f"url:{identity}")
    return keys


def _upsert_xlsx_rows_in_place(ws, managed_headers: list[str], records: list[dict]) -> bool:
    """Merge values by identity without deleting rows or touching existing styles."""
    actual_headers = [ws.cell(1, column).value for column in range(1, ws.max_column + 1)]
    while actual_headers and actual_headers[-1] in (None, ""):
        actual_headers.pop()
    columns = {
        str(header): index for index, header in enumerate(actual_headers, 1)
        if header not in (None, "")
    }
    by_key: dict[str, int] = {}
    changed = False
    for row in range(2, ws.max_row + 1):
        record = {header: ws.cell(row, column).value for header, column in columns.items()}
        for key in _xlsx_record_keys(record, ws.title):
            by_key[key] = row

    hyperlink_headers = (
        {"Job Link", "Resume Link"}
        if ws.title == APPLICATION_SHEET
        else {"Authoritative URL", "Job URL", "Enrichment URL", "Resume Link"}
    )
    for record in records:
        row = next((by_key[key] for key in _xlsx_record_keys(record, ws.title) if key in by_key), None)
        if row is None:
            row = append_preserving_format(ws, set(managed_headers))
            changed = True
        for header, column in columns.items():
            if header not in record:
                continue
            value = record.get(header)
            normalized = None if value == "" else value
            cell = ws.cell(row, column)
            if cell.value != normalized:
                cell.value = normalized
                changed = True
            if header in hyperlink_headers:
                target = str(normalized) if normalized else None
                current = cell.hyperlink.target if cell.hyperlink else None
                if current != target:
                    cell.hyperlink = target
                    changed = True
        for key in _xlsx_record_keys(record, ws.title):
            by_key[key] = row
    return changed


def _save_xlsx_atomic(workbook, path: Path) -> None:
    with NamedTemporaryFile(prefix="google-sheets-backup-", suffix=".xlsx", dir=path.parent, delete=False) as temp:
        temp_path = Path(temp.name)
    try:
        workbook.save(temp_path)
        os.replace(temp_path, path)
    finally:
        if temp_path.exists():
            temp_path.unlink()


def _sync_workbook_to_google(
    tracker_path: str | Path,
    *,
    config: GoogleSheetsConfig | None = None,
    service=None,
    remove_scout_ids: set[int] | None = None,
) -> GoogleSyncSummary:
    """Merge the XLSX backup with Google, preserving manual fields in both copies."""
    tracker_path = Path(tracker_path)
    if not tracker_path.is_file():
        raise ValueError(f"Tracker workbook does not exist: {tracker_path}")
    config = config or GoogleSheetsConfig.from_environment()
    service = service or build_sheets_service(config)
    api = service.spreadsheets()
    values_api = api.values()
    sheets = _ensure_required_sheets(api, values_api, config.spreadsheet_id)
    missing = [name for name in (APPLICATION_SHEET, SCOUT_SHEET) if name not in sheets]
    if missing:
        raise RuntimeError("Google Sheets API did not create required worksheets: " + ", ".join(missing))

    workbook = load_workbook(tracker_path)
    if APPLICATION_SHEET not in workbook.sheetnames or SCOUT_SHEET not in workbook.sheetnames:
        raise ValueError("XLSX tracker must contain both 'Job Tracker' and 'Job Scout' worksheets")
    local_application = _records_from_xlsx(workbook[APPLICATION_SHEET], APPLICATION_HEADERS)
    local_scout = _records_from_xlsx(workbook[SCOUT_SHEET], SCOUT_HEADERS)

    remote_application_values = values_api.get(
        spreadsheetId=config.spreadsheet_id,
        # Read the complete user-visible layout so reordered and custom columns
        # can be preserved during the merge.
        range=_sheet_range(APPLICATION_SHEET, "ZZ"),
        # Return formulas verbatim so targeted value updates preserve formulas
        # in user-owned custom columns rather than freezing their results.
        valueRenderOption="FORMULA",
        dateTimeRenderOption="FORMATTED_STRING",
    ).execute().get("values", [])
    remote_scout_values = values_api.get(
        spreadsheetId=config.spreadsheet_id,
        # Read beyond Gecko's managed columns so user-added columns can be
        # detected and ignored without being overwritten.
        range=_sheet_range(SCOUT_SHEET, "ZZ"),
        valueRenderOption="FORMULA",
        dateTimeRenderOption="FORMATTED_STRING",
    ).execute().get("values", [])
    remote_application = _records_from_values(remote_application_values, APPLICATION_HEADERS, APPLICATION_SHEET)
    remote_scout = _records_from_values(remote_scout_values, SCOUT_HEADERS, SCOUT_SHEET)
    removed = remove_scout_ids or set()
    if removed:
        for row in remote_scout:
            if str(row.get("Scout ID")) in {str(value) for value in removed} and (
                any(row.get(field) not in (None, "") for field in ("Apply?", "Applied", "Contacted", "Resume Created"))
                or str(row.get("Gecko Status") or "").lower() not in {"", "new", "reviewing"}
            ):
                raise ValueError("A proposed dead Scout row has protected Google Sheets history")
        remote_scout = [row for row in remote_scout if str(row.get("Scout ID")) not in {str(value) for value in removed}]
        local_scout = [row for row in local_scout if str(row.get("Scout ID")) not in {str(value) for value in removed}]
    application_headers = _output_headers(remote_application_values, APPLICATION_HEADERS)
    scout_headers = _output_headers(remote_scout_values, SCOUT_HEADERS)

    merged_application = merge_application_rows(local_application, remote_application)
    merged_scout = merge_scout_rows(local_scout, remote_scout)
    attach_resume_links(merged_application, merged_scout, tracker_path)
    merged_application = _preserve_remote_row_order(
        merged_application, remote_application, APPLICATION_SHEET,
    )
    merged_scout = _preserve_remote_row_order(merged_scout, remote_scout, SCOUT_SHEET)
    application_values = _values_from_records(application_headers, merged_application)
    scout_values = _scout_values_preserving_positions(scout_headers, merged_scout, remote_scout_values, removed)

    _write_sheet(
        values_api, config.spreadsheet_id, APPLICATION_SHEET,
        remote_application_values, application_values,
    )
    _write_sheet(
        values_api, config.spreadsheet_id, SCOUT_SHEET,
        remote_scout_values, scout_values,
    )

    format_requests = _new_row_format_requests(
        sheets[APPLICATION_SHEET], APPLICATION_SHEET,
        len(remote_application_values), len(application_values), application_headers,
    )
    format_requests.extend(_new_row_format_requests(
        sheets[SCOUT_SHEET], SCOUT_SHEET,
        len(remote_scout_values), len(scout_values), scout_headers,
    ))
    _run_optional_requests(
        api, config.spreadsheet_id, format_requests, "new-row number formatting",
    )
    link_requests = _resume_link_format_requests(
        sheets[SCOUT_SHEET]["properties"]["sheetId"], merged_scout, scout_headers,
    )
    _run_optional_requests(
        api, config.spreadsheet_id, link_requests, "resume-link formatting",
    )

    application_changed = _upsert_xlsx_rows_in_place(
        workbook[APPLICATION_SHEET], APPLICATION_HEADERS, merged_application,
    )
    scout_changed = _upsert_xlsx_rows_in_place(
        workbook[SCOUT_SHEET], SCOUT_HEADERS, merged_scout,
    )
    if application_changed or scout_changed:
        _save_xlsx_atomic(workbook, tracker_path)

    return GoogleSyncSummary(
        application_rows=len(merged_application),
        scout_rows=len(merged_scout),
        spreadsheet_id=config.spreadsheet_id,
    )


def sync_workbook_to_google(
    tracker_path: str | Path,
    *,
    config: GoogleSheetsConfig | None = None,
    service=None,
    remove_scout_ids: set[int] | None = None,
) -> GoogleSyncSummary:
    """Run synchronization and translate Google API failures into actionable errors."""
    try:
        return _sync_workbook_to_google(tracker_path, config=config, service=service,
                                        remove_scout_ids=remove_scout_ids)
    except Exception as error:
        if error.__class__.__module__.startswith("googleapiclient"):
            status = getattr(getattr(error, "resp", None), "status", None)
            if status == 404:
                raise RuntimeError(
                    "Google Sheet was not found or is not shared with the configured "
                    "service-account email; verify the Spreadsheet ID and Editor share"
                ) from error
            raise RuntimeError(f"Google Sheets API request failed: {error}") from error
        raise
