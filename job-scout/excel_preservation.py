"""Helpers for modifying persistent Excel workbooks without restyling them."""

from __future__ import annotations

from copy import copy

from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter, range_boundaries


def copy_row_format(ws, source_row: int, target_row: int, max_column: int | None = None) -> None:
    """Copy cell and row formatting, but not values, into a newly appended row."""
    if source_row < 1 or target_row < 1:
        return
    max_column = max_column or ws.max_column
    for column in range(1, max_column + 1):
        source = ws.cell(source_row, column)
        target = ws.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
    source_dimension = ws.row_dimensions[source_row]
    target_dimension = ws.row_dimensions[target_row]
    for attribute in ("height", "hidden", "outlineLevel", "collapsed"):
        setattr(target_dimension, attribute, getattr(source_dimension, attribute))


def copy_unmanaged_formulas(
    ws, source_row: int, target_row: int, managed_headers: set[str],
) -> None:
    """Fill down formulas in user-owned columns when a new data row is appended."""
    for column in range(1, ws.max_column + 1):
        header = str(ws.cell(1, column).value or "")
        source = ws.cell(source_row, column)
        if header in managed_headers or not isinstance(source.value, str) or not source.value.startswith("="):
            continue
        try:
            ws.cell(target_row, column).value = Translator(
                source.value, origin=source.coordinate,
            ).translate_formula(ws.cell(target_row, column).coordinate)
        except (TypeError, ValueError):
            ws.cell(target_row, column).value = source.value


def expand_row_structures(ws, previous_row: int, new_row: int) -> None:
    """Extend existing filters, tables, and row-scoped validations by one row."""
    if new_row != previous_row + 1:
        return

    if ws.auto_filter.ref:
        min_col, min_row, max_col, max_row = range_boundaries(ws.auto_filter.ref)
        if min_row <= previous_row == max_row:
            ws.auto_filter.ref = (
                f"{get_column_letter(min_col)}{min_row}:"
                f"{get_column_letter(max_col)}{new_row}"
            )

    for table in ws.tables.values():
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        if min_row <= previous_row == max_row:
            expanded = (
                f"{get_column_letter(min_col)}{min_row}:"
                f"{get_column_letter(max_col)}{new_row}"
            )
            table.ref = expanded
            if table.autoFilter is not None:
                table.autoFilter.ref = expanded

    for validation in ws.data_validations.dataValidation:
        additions = []
        for cell_range in list(validation.ranges.ranges):
            if cell_range.min_row <= previous_row == cell_range.max_row:
                additions.append(
                    f"{get_column_letter(cell_range.min_col)}{new_row}:"
                    f"{get_column_letter(cell_range.max_col)}{new_row}"
                )
        for addition in additions:
            validation.add(addition)


def append_preserving_format(ws, managed_headers: set[str]) -> int:
    """Create a blank row inheriting the previous data row's presentation."""
    previous_row = 1
    for row in range(ws.max_row, 1, -1):
        if any(ws.cell(row, column).value not in (None, "") for column in range(1, ws.max_column + 1)):
            previous_row = row
            break
    new_row = previous_row + 1
    if previous_row >= 2:
        copy_row_format(ws, previous_row, new_row, ws.max_column)
        copy_unmanaged_formulas(ws, previous_row, new_row, managed_headers)
        expand_row_structures(ws, previous_row, new_row)
    return new_row
