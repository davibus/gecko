"""Read the current canonical Gecko DOCX; no claims are introduced here."""

from __future__ import annotations

from pathlib import Path


def extract_resume_text(path: str | Path) -> str:
    path = Path(path)
    if path.name != "Dave-Call-resume-9-23-26.docx":
        raise ValueError("Gecko evidence must come from Dave-Call-resume-9-23-26.docx")
    from docx import Document
    document = Document(path)
    lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    lines.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells if cell.text.strip())
    return "\n".join(lines)
